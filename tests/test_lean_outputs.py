import copy
import importlib.util
import io
import json
import shutil
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest

from study2 import adapters, descriptives, descriptive_reports as reports, storage, store
from study2.common import PipelineError, dumps, utcnow
from study2.descriptive_spec import REPORT_FILES
from test_identity_and_time import rec
from test_store_and_review import database, source
from test_pipeline_integration import project, execute


@pytest.mark.parametrize('headers', [
    ['Email', '建立時間', 'value'], ['same', 'same', None],
    [10, '010', 'with space '], ['日期', None, '最後'], [],
])
@pytest.mark.parametrize('sheet', [0, 'Selected'])
def test_header_only_matches_full_pandas_headers(tmp_path, headers, sheet):
    path = tmp_path/'headers.xlsx'
    book = openpyxl.Workbook(); page = book.active; page.title = 'Selected'
    if headers:
        page.append(headers); page.append(['001']*len(headers))
    book.save(path)
    entry = {'kind': 'merged', 'sheet': sheet}
    full = adapters.read_source(path, entry, None)
    header = adapters.read_source(path, entry, None, headers_only=True)
    assert list(header.columns) == list(full.columns)
    # A trailing blank header is ambiguous and conservatively uses the full reader.
    if headers and headers[-1] is not None:
        assert header.empty


def test_manifest_uses_header_only_but_retains_schema_guard(tmp_path):
    path = tmp_path/'source.xlsx'
    pd.DataFrame({'A': [1], 'B': [2]}).to_excel(path, index=False)
    cfg = {'project_root': str(tmp_path), 'source_folder': '.', 'sources': [
        {'source_id': 'one', 'pattern': 'source.xlsx', 'kind': 'merged', 'columns': ['A', 'B']}]}
    with patch.object(adapters, 'modules', return_value=(None, None, None)), patch.object(
            pd.ExcelFile, 'parse', autospec=True, side_effect=pd.ExcelFile.parse) as read:
        manifest = adapters.make_manifest(cfg)
        assert read.call_args.kwargs['nrows'] == 0
        assert len(manifest['sources']) == 1
        cfg['sources'][0]['columns'] = ['unexpected']
        with pytest.raises(PipelineError, match='Unregistered columns'):
            adapters.make_manifest(cfg)


def test_qualtrics_keeps_existing_multirow_reader():
    online = SimpleNamespace(read_qualtrics_csv=lambda path: pd.DataFrame({'Q1': ['answer']}))
    assert not adapters.read_source(Path('synthetic.csv'), {'kind': 'qualtrics'}, online,
                                    headers_only=True).empty


def test_existing_versions_skip_payload_serialization_and_changed_versions_do_not(tmp_path):
    db = database(tmp_path)
    row = source(rec('a'))
    with db:
        store.ingest(db, 'run', [row])
    original = tuple(db.execute('SELECT * FROM submission_versions').fetchone())
    with db:
        db.execute('INSERT INTO runs VALUES(?,?,?,?,?,?,?)', ('second', 'fp', utcnow(), 'running', '{}', None, None))
        with patch.object(store, 'dumps', side_effect=AssertionError('unchanged source serialization')):
            store.ingest(db, 'second', [copy.deepcopy(row)])
    assert tuple(db.execute('SELECT * FROM submission_versions').fetchone()) == original
    changed = copy.deepcopy(row)
    changed.update(version_id='changed-version', raw_hash='changed-hash', raw={'answer': 'changed'})
    with db:
        db.execute('INSERT INTO runs VALUES(?,?,?,?,?,?,?)', ('third', 'fp', utcnow(), 'running', '{}', None, None))
        with patch.object(store, 'dumps', wraps=dumps) as serial:
            store.ingest(db, 'third', [changed])
            assert serial.call_count == 3
    assert db.execute('SELECT count(*) FROM submission_versions').fetchone()[0] == 2
    assert db.execute('SELECT version_id FROM run_records WHERE run_id=?', ('third',)).fetchone()[0] == 'changed-version'
    db.close()


def test_version_binding_mismatch_does_not_commit(tmp_path):
    db = database(tmp_path); row = source(rec('a'))
    with db: store.ingest(db, 'run', [row])
    with pytest.raises(PipelineError, match='version binding'):
        with db:
            db.execute('INSERT INTO runs VALUES(?,?,?,?,?,?,?)', ('bad', 'fp', utcnow(), 'running', '{}', None, None))
            store.ingest(db, 'bad', [{**row, 'raw_hash': 'wrong'}])
    assert db.execute("SELECT count(*) FROM runs WHERE run_id='bad'").fetchone()[0] == 0
    db.close()


def test_value_cache_reuses_only_verified_content_within_transaction(tmp_path):
    db = database(tmp_path); queries=[]; db.set_trace_callback(queries.append)
    value = {'number': '001', 'missing': None}
    with db:
        db.execute('BEGIN')
        with storage.record_value_cache(db) as cache:
            first=storage._value(db,value)
            assert storage._value(db,copy.deepcopy(value)) == first
            value['number']='002'
            assert storage._value(db,value) != first
            assert len(cache['values']) == 2
    reads = [q for q in queries if q.startswith('SELECT payload_json FROM record_value_json')]
    assert len(reads) == 2
    assert storage._VALUE_CACHE.get() is None and not cache['values']
    queries.clear()
    storage._value(db, {'number': '001', 'missing': None})
    assert any(q.startswith('SELECT payload_json FROM record_value_json') for q in queries)
    db.close()


def test_cache_failure_clears_entries_before_rollback_and_is_bounded(tmp_path):
    db=database(tmp_path)
    with pytest.raises(RuntimeError):
        with db:
            db.execute('BEGIN')
            with storage.record_value_cache(db, max_bytes=40) as cache:
                for i in range(12): storage._value(db, {'i':i})
                assert cache['bytes'] <= 40
                raise RuntimeError('synthetic failure')
    assert storage._VALUE_CACHE.get() is None
    assert not cache['values']
    assert db.execute('SELECT count(*) FROM record_values').fetchone()[0] == 0
    with db:
        db.execute('BEGIN')
        with storage.record_value_cache(db): storage._value(db, {'i':0})
    assert db.execute('SELECT count(*) FROM record_values').fetchone()[0] == 1
    db.close()


def test_cache_never_crosses_connections(tmp_path):
    one=database(tmp_path/'one'); two=database(tmp_path/'two')
    with one, two:
        one.execute('BEGIN'); two.execute('BEGIN')
        with storage.record_value_cache(one):
            assert storage._value(one, {'x':1}) == storage._value(two, {'x':1})
    assert two.execute('SELECT count(*) FROM record_values').fetchone()[0] == 1
    one.close();two.close()


@pytest.fixture(scope='module')
def publication(tmp_path_factory):
    root=tmp_path_factory.mktemp('lean-output-synthetic')
    setup=project.__wrapped__(root)
    cfg=setup[0];cfg['output_policy']={'formatted_descriptive_reports':'on_demand'}
    with patch.object(descriptives, 'cohort', wraps=descriptives.cohort) as cohorts, patch.object(
            reports, 'workbook_bytes', side_effect=AssertionError('workbooks must be deferred')):
        result=execute(setup)
        assert cohorts.call_count == 8
    folder=Path(cfg['release_dir'])/result['release_id']
    assert all(result[k] == 1 for k in ['A','B','C','D','AC','BD','E','F','EF'])
    assert execute(setup)['unchanged'] is True
    return root,folder,result,setup


def test_deferred_outputs_keep_numeric_contract_and_release_verification(publication):
    _,folder,result,_=publication
    assert not (folder/'4.1_descriptive').exists()
    assert result['descriptive']['reports'] == {}
    assert result['descriptive']['formatted_reports']['status'] == 'deferred'
    reports.verify_reports(folder,result['descriptive'])
    from study2.runner import verify_release
    db=store.connect(Path(publication[3][0]['state_dir']))
    verify_release(db,result['release_id']);db.close()


def cells(data):
    book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
    values={s.title:list(s.values) for s in book}
    book.close();return values


def test_prepared_cohorts_and_on_demand_reports_equal_prior_eager_results(publication):
    root,folder,result,_=publication
    # Compare shared prepared cohorts with the retained unshared eager path.
    old=reports
    target=root/'baseline'/folder.name;target.mkdir(parents=True)
    measurements=pd.read_parquet(folder/'descriptive_measurements.parquet')
    frame=pd.read_parquet(folder/'analysis.parquet')
    parents=pd.read_parquet(folder/'parent_longitudinal.parquet') if (folder/'parent_longitudinal.parquet').exists() else None
    old.write_reports(target,measurements,frame,parents,folder.name)
    for name in ('descriptive_statistics.parquet','descriptive_changes.parquet'):
        pd.testing.assert_frame_equal(pd.read_parquet(folder/name),pd.read_parquet(target/name))
    before=(folder/'descriptive_manifest.json').read_bytes()
    output=root/'on-demand'
    exported=reports.export_reference_reports(folder,output)
    assert set(exported['files']) == set(REPORT_FILES.values())
    assert exported['membership'] == 'canonical_OFF_reference'
    for name in REPORT_FILES.values():
        assert cells((output/name).read_bytes()) == cells((target/'4.1_descriptive'/name).read_bytes())
    assert (folder/'descriptive_manifest.json').read_bytes() == before
    assert not (folder/'4.1_descriptive').exists()
    with pytest.raises(PipelineError,match='new directory'):
        reports.export_reference_reports(folder,output)


def test_preparation_is_fresh_and_consumers_do_not_modify_shared_cohorts(publication):
    _,folder,_,_=publication
    frame=pd.read_parquet(folder/'analysis.parquet')
    parents=pd.read_parquet(folder/'parent_longitudinal.parquet') if (folder/'parent_longitudinal.parquet').exists() else None
    measurements=pd.read_parquet(folder/'descriptive_measurements.parquet')
    prepared=descriptives.prepare_report_cohorts(frame,parents,measurements)
    saved={key:tuple(v.copy(deep=True) for v in pair) for key,pair in prepared.items()}
    from study2.descriptive_exploration import explore
    for key,(_,selected) in prepared.items():
        descriptives.statistics(selected,key[1],['all'])
        descriptives.paired_changes(selected,key[1],['all'])
        explore(selected,key[1],['all'])
        for actual,expected in zip(prepared[key],saved[key]):pd.testing.assert_frame_equal(actual,expected)
    # A changed measurement is consumed by new preparation, never a global cache.
    again=descriptives.prepare_report_cohorts(frame,parents,measurements)
    assert all(again[k][0] is not prepared[k][0] for k in prepared)
    changed=frame.copy();changed['mainstream']=False
    refreshed=descriptives.prepare_report_cohorts(changed,parents,measurements)
    assert len(refreshed['abcd','student'][1]) == 0
    assert len(prepared['abcd','student'][1]) > 0


def test_numeric_tampering_blocks_on_demand_export(publication,tmp_path):
    _,folder,_,_=publication
    copied=tmp_path/folder.name;copied.mkdir()
    for name in ('descriptive_manifest.json','descriptive_statistics.parquet','descriptive_changes.parquet'):
        shutil.copy2(folder/name,copied/name)
    data=pd.read_parquet(copied/'descriptive_changes.parquet');data['mean_change']=99
    data.to_parquet(copied/'descriptive_changes.parquet',index=False)
    with pytest.raises(PipelineError,match='fingerprint'):
        reports.export_reference_reports(copied,tmp_path/'must-not-exist')
    assert not (tmp_path/'must-not-exist').exists()


def test_unknown_output_policy_stops_before_database_work(tmp_path):
    from study2.runner import run
    with pytest.raises(PipelineError,match='Unknown formatted'):
        run.__wrapped__({'output_policy':{'formatted_descriptive_reports':'typo'}}, {})


def test_formatting_only_extra_columns_preserve_full_header_result(tmp_path):
    path=tmp_path/'formatting.xlsx'
    book=openpyxl.Workbook();sheet=book.active
    sheet.append(['A','B']);sheet.append([1,2]);sheet.cell(50,7).number_format='0.00'
    book.save(path)
    full=adapters.read_source(path,{'kind':'merged'},None)
    header=adapters.read_source(path,{'kind':'merged'},None,headers_only=True)
    assert list(header.columns)==list(full.columns)==['A','B']


def test_deferred_ui_is_explicit_on_demand_and_never_substitutes_off_for_on():
    summary={'descriptive':{'formatted_reports':{'status':'deferred'}}}
    ui=SimpleNamespace(info=lambda text:None, button=lambda *a,**k:False,
                       download_button=lambda *a,**k:None)
    with patch.object(reports,'reference_workbook',return_value=b'generated') as render:
        reports.render_reference_downloads(ui,Path('does-not-exist'),summary)
        render.assert_not_called()
        ui.button=lambda *a,**k:k['key']=='ds_prepare_acbd'
        downloads=[];ui.download_button=lambda *a,**k:downloads.append(a)
        reports.render_reference_downloads(ui,Path('does-not-exist'),summary,exclusions_on=True)
        render.assert_not_called();assert not downloads
        reports.render_reference_downloads(ui,Path('does-not-exist'),summary)
        render.assert_called_once_with(Path('does-not-exist'),'acbd')
        assert len(downloads)==1 and downloads[0][2].startswith('OFF-reference-')


def test_exploration_tables_equal_previous_unshared_cohorts(publication,tmp_path):
    _,folder,_,_=publication
    from study2 import descriptive_exploration_outputs as old
    target=tmp_path/folder.name;target.mkdir()
    panel=folder/'independent_parent_analysis.parquet'
    if panel.exists():shutil.copy2(panel,target/panel.name)
    measurements=pd.read_parquet(folder/'descriptive_measurements.parquet')
    frame=pd.read_parquet(folder/'analysis.parquet')
    parents=pd.read_parquet(folder/'parent_longitudinal.parquet') if (folder/'parent_longitudinal.parquet').exists() else None
    old.write_exploration(target,measurements,frame,parents,folder.name)
    for name in old.FILES.values():
        pd.testing.assert_frame_equal(pd.read_parquet(folder/name),pd.read_parquet(target/name))
