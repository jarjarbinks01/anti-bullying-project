"""Invented release, pure projection and real in-process Streamlit reruns."""
from copy import deepcopy
import ast
import gzip
import hashlib
import io
import json
import os
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from study2.analysis_eligibility import manifest_revision
from study2.descriptive_spec import VERSION,REPORT_FILES
from study2.streamlit_eligibility import (ActiveEligibilityUnavailable,EligibilityState,
    ENVIRONMENT_KEY,SESSION_KEY,load_eligibility,prepare_eligibility,stamp_frame)


def inputs():
    root=Path(os.environ['MANUAL_UI_CORE_FIXTURE_DIR'])
    data=json.loads((root/'SYNTHETIC_INPUT_ORACLE.json').read_text())
    public=json.loads((root/'SYNTHETIC_FIXTURE.json').read_text())
    panel={k:pd.DataFrame(v) for k,v in data['panel'].items()}
    if panel['candidates'].empty:panel['candidates']=pd.DataFrame(columns=['one_digit_candidate'])
    if panel['sensitivity'].empty:panel['sensitivity']=pd.DataFrame(columns=['child_id','t0','exact_T1','experiment_T1','experiment_rule','comparison','experimental_only'])
    detail=dict(parent_submissions=len(panel['audit']),baselines=int((panel['analysis'].panel_branch.eq('child_report')&panel['analysis'].panel_wave.eq('T0')).sum()),
        child_report_pairs=int(panel['pairs'].panel_branch.eq('child_report').sum()),same_adult_pairs=int(panel['pairs'].panel_branch.eq('same_adult').sum()),
        minimum_followup_days=90,one_digit_candidates=0,same_adult_alternative_T1=1,
        experiment_pairs=0,experiment_overlap_children=0,experiment_added_children=0,experiment_held_exact_children=0,experiment_changed_T1=0,audit_status_counts={})
    summary=dict(release_id=data['accepted_manifest']['release_id'],descriptive={'version':VERSION},independent_parent=detail,
                 **public['snapshot']['counts'])
    return data,pd.DataFrame(data['analysis']),pd.DataFrame(data['parents']),pd.DataFrame(data['measurements']),panel,summary


def accepted(records,original,ids):
    manifest=deepcopy(original)
    # Honor the accepted all-identical-copies policy for the invented inventory.
    groups={(r['role'],r['raw_hash']) for r in records if r['submission_id'] in ids}
    manifest['entries']=[dict(submission_id=r['submission_id'],version_id=r['version_id']) for r in records if (r['role'],r['raw_hash']) in groups]
    manifest['revision']=manifest_revision(manifest)
    return manifest


def release_fixture(folder,*,baseline_flag=False):
    data,analysis,parents,measurements,panel,summary=inputs();folder=Path(folder);folder.mkdir(parents=True)
    manifest=data['accepted_manifest']
    if baseline_flag:
        sid=panel['pairs'].loc[panel['pairs'].panel_branch.eq('child_report'),'t0'].iloc[0]
        manifest=accepted(data['records'],manifest,{sid})
    (folder/'summary.json').write_text(json.dumps(summary))
    analysis.to_parquet(folder/'analysis.parquet',index=False);parents.to_parquet(folder/'parent_longitudinal.parquet',index=False)
    measurements.to_parquet(folder/'descriptive_measurements.parquet',index=False)
    for key,frame in panel.items():frame.to_parquet(folder/f'independent_parent_{key}.parquet',index=False)
    with gzip.open(folder/'records.jsonl.gz','wt',encoding='utf-8') as stream:
        for row in data['records']:stream.write(json.dumps(row)+'\n')
    target=folder/'accepted-synthetic-manifest.json';target.write_text(json.dumps(manifest))
    from openpyxl import Workbook
    reports=folder/'4.1_descriptive';reports.mkdir()
    for name in REPORT_FILES.values():
        book=Workbook();book.active.append(['Synthetic original OFF/reference']);book.save(reports/name)
    return target,data,analysis,parents,measurements,panel,summary


def test_pure_prepared_members_anchors_and_immutability():
    data,analysis,parents,measurements,panel,summary=inputs()
    original=[x.copy(deep=True) for x in [analysis,parents,measurements,*panel.values()]]
    state=prepare_eligibility(data['records'],analysis,parents,measurements,summary,data['accepted_manifest'],panel=panel)
    on=state.select(True)
    for code,members in data['retained_members'].items():
        rows=on.sample_rows(analysis,parents,code)
        assert len(rows)==len(members)
    assert len(on.impact(summary)['samples'])==9
    assert state.select(False).metadata()['eligibility_mode']=='off'
    before=state.cohort('acbd','parent');after=on.cohort('acbd','parent')
    expected={sid for pair in data['retained_members']['BD'] for sid in pair}
    assert set(after.submission_id)==expected
    for frame in [before,after]:assert frame['adult_continuity_confirmed'].all()
    sid=panel['pairs'].loc[panel['pairs'].panel_branch.eq('child_report'),'t0'].iloc[0]
    missing_base=prepare_eligibility(data['records'],analysis,parents,measurements,summary,
        accepted(data['records'],data['accepted_manifest'],{sid}),panel=panel).select(True)
    all_rows=missing_base.cohort('independent_all','parent')
    assert set(all_rows.wave)=={'T1'} and sid not in set(all_rows.submission_id)
    canonical=state.cohort('independent_all','parent')
    first=canonical.loc[canonical.wave.eq('T0')].iloc[0]
    assert all_rows.anchor_source_type.eq(first.anchor_source_type).all()
    assert all_rows.source_type.ne(all_rows.anchor_source_type).all()
    assert missing_base.cohort('independent_child','parent').empty
    assert missing_base.cohort('independent_adult','parent').empty
    # Returned frames are owned copies, never cached input aliases.
    all_rows.loc[:,'anchor_source_type']='changed-in-test'
    assert missing_base.cohort('independent_all','parent').anchor_source_type.eq(first.anchor_source_type).all()
    for left,right in zip(original,[analysis,parents,measurements,*panel.values()]):pd.testing.assert_frame_equal(left,right)


def test_off_without_explicit_manifest_does_not_read_inputs(monkeypatch):
    monkeypatch.delenv(ENVIRONMENT_KEY,raising=False)
    with patch.object(Path,'read_bytes',side_effect=AssertionError('must not read')),patch.object(pd,'read_parquet',side_effect=AssertionError('must not read')):
        state=load_eligibility('unused',{'release_id':'invented'},pd.DataFrame(),None)
    assert not state.enabled and state.allowed and not state.ready
    requested=state.select(True);assert not requested.allowed
    with pytest.raises(ActiveEligibilityUnavailable):requested.cohort('abcd','student')
    assert requested.select(False).allowed


def test_local_loader_strict_manifest_and_safe_errors(tmp_path,monkeypatch):
    path,data,analysis,parents,measurements,panel,summary=release_fixture(tmp_path/'release')
    state=load_eligibility(path.parent,summary,analysis,parents,manifest_path=path)
    assert state.ready
    bad=deepcopy(data['accepted_manifest']);bad['entries'][0]['version_id']='synthetic-wrong-version';bad['revision']=manifest_revision(bad)
    path.write_text(json.dumps(bad));failed=load_eligibility(path.parent,summary,analysis,parents,manifest_path=path)
    assert not failed.ready and failed.reason=='manifest_version_mismatch'
    assert str(path) not in failed.reason and 'synthetic-wrong-version' not in failed.reason
    missing=load_eligibility(path.parent,summary,analysis,parents,manifest_path=tmp_path/'private-contact-person@example.com')
    assert not missing.ready and missing.reason=='manual_exclusion_inputs_unavailable'
    before=analysis.copy(deep=True);failed.select(False).analytical_frame(analysis);pd.testing.assert_frame_equal(before,analysis)


def harness():
    """Actual checkbox, file adapter and renderers; no official dashboard/DB."""
    import streamlit as raw
    from study2.english_ui import st
    from study2.streamlit_eligibility import render_selector,render_impact
    from study2 import descriptive_dashboard as descriptive,parent_panel_dashboard as parent
    from study2.descriptives import wave_counts
    root=Path(os.environ['MANUAL_STREAMLIT_TEST_ROOT'])
    summary=json.loads((root/'summary.json').read_text())
    analysis=pd.read_parquet(root/'analysis.parquet');parents=pd.read_parquet(root/'parent_longitudinal.parquet')
    prepared=load_eligibility(root,summary,analysis,parents)
    raw.session_state['_captured_downloads']={}
    original_download=raw.download_button
    def download(label,data=None,file_name=None,mime=None,**kwargs):
        raw.session_state['_captured_downloads'][kwargs.get('key',label)]=dict(name=file_name,data=data)
        return original_download(label,data,file_name,mime,**kwargs)
    workbench=__import__('study2.descriptive_workbench',fromlist=['render_overview','render_workbench'])
    trend=__import__('study2.descriptive_trend_dashboard',fromlist=['render_trends'])
    def capture_rows(renderer):
        def render(rows,*args,**kwargs):
            raw.session_state['_rendered_rows']=rows.copy(deep=True)
            raw.session_state['_renderer_metadata']=kwargs.get('metadata')
            return renderer(rows,*args,**kwargs)
        return render
    def canonical_audit(data,details,**kwargs):
        raw.session_state['_audit_ids']=tuple(data['audit'].submission_id)
        st.caption('Synthetic canonical OFF/reference administrative audit seam')
    with patch.object(raw,'download_button',side_effect=download):
        state=render_selector(st,prepared)
        raw.session_state['_eligibility_mode']=state.metadata()['eligibility_mode']
        raw.session_state['_active_allowed']=state.allowed
        if state.allowed:raw.session_state['_impact']=state.impact(summary)
        else:raw.session_state.pop('_impact',None)
        render_impact(st,state,summary)
        surface=st.selectbox('Synthetic test surface',['descriptive','independent'],key='_surface')
        with patch.object(workbench,'render_overview',side_effect=capture_rows(workbench.render_overview)),patch.object(workbench,'render_workbench',side_effect=capture_rows(workbench.render_workbench)),patch.object(trend,'render_trends',side_effect=capture_rows(trend.render_trends)),patch.object(descriptive,'_capability',return_value=True):
            if surface=='descriptive':
                with patch.object(descriptive,'cohort',side_effect=AssertionError('renderer reconstructed canonical anchors')) if prepared.ready else patch.object(descriptive,'_capability',return_value=True):
                    descriptive.render_descriptive(root,summary,analysis,parents,{},eligibility=state)
            else:
                with patch.object(parent,'_render_panel_audit',side_effect=canonical_audit),patch('study2.descriptives.independent_cohort',side_effect=AssertionError('renderer reconstructed independent anchors')) if prepared.ready else patch.object(parent,'_render_panel_audit',side_effect=canonical_audit):
                    parent.render_parent_panel(root,summary,eligibility=state)


def widget(app,kind,key):
    return next(x for x in getattr(app,kind) if x.key==key)


def test_actual_apptest_off_on_off_filters_downloads_and_invalid_recovery(tmp_path,monkeypatch):
    manifest,*_=release_fixture(tmp_path/'release')
    monkeypatch.setenv('MANUAL_STREAMLIT_TEST_ROOT',str(manifest.parent));monkeypatch.setenv(ENVIRONMENT_KEY,str(manifest))
    app=AppTest.from_string('from test_streamlit_eligibility import harness\nharness()').run(timeout=60)
    assert not app.exception
    assert widget(app,'checkbox',SESSION_KEY).value is False
    widget(app,'selectbox','ds_mainstream').select(False).run(timeout=60)
    widget(app,'multiselect','ds_sources').set_value(['qualtrics']).run(timeout=60)
    assert not app.exception
    off_rows=app.session_state['_rendered_rows'].copy(deep=True)
    off_impact=deepcopy(app.session_state['_impact'])
    widget(app,'checkbox',SESSION_KEY).check().run(timeout=60)
    assert not app.exception and app.session_state['_eligibility_mode']=='on'
    assert app.session_state['ds_sources']==['qualtrics'] and app.session_state['ds_mainstream'] is False
    assert app.session_state['_impact']['samples']!=off_impact['samples']
    assert app.session_state['_renderer_metadata']['eligibility_mode']=='on'
    captures=app.session_state['_captured_downloads']
    assert not any(key.startswith('ds_report_') for key in captures),'Canonical workbooks cannot masquerade as ON'
    table=next(v for k,v in captures.items() if k.endswith('_statistics_csv'))
    csv=pd.read_csv(io.BytesIO(table['data']))
    assert csv.eligibility_mode.eq('on').all() and csv.eligibility_revision.ne('not_available').all()
    assert 'eligibility-on' in table['name']
    widget(app,'checkbox',SESSION_KEY).uncheck().run(timeout=60)
    assert not app.exception and app.session_state['ds_sources']==['qualtrics']
    pd.testing.assert_frame_equal(app.session_state['_rendered_rows'],off_rows)
    assert sum(k.startswith('ds_report_') for k in app.session_state['_captured_downloads'])==4
    widget(app,'checkbox',SESSION_KEY).check().run(timeout=60)
    raw=json.loads(manifest.read_text());raw['status']='pending';manifest.write_text(json.dumps(raw))
    app.run(timeout=60);assert not app.exception
    assert widget(app,'checkbox',SESSION_KEY).value is True and not app.session_state['_active_allowed']
    assert not any(k.endswith('_statistics_csv') for k in app.session_state['_captured_downloads'])
    assert str(manifest) not in '\n'.join(x.value for x in app.error)
    widget(app,'checkbox',SESSION_KEY).uncheck().run(timeout=60)
    assert not app.exception and app.session_state['_active_allowed'] and app.session_state['ds_sources']==['qualtrics']
    pd.testing.assert_frame_equal(app.session_state['_rendered_rows'],off_rows)


def test_actual_apptest_independent_singleton_keeps_original_anchor(tmp_path,monkeypatch):
    manifest,*_=release_fixture(tmp_path/'release',baseline_flag=True)
    monkeypatch.setenv('MANUAL_STREAMLIT_TEST_ROOT',str(manifest.parent));monkeypatch.setenv(ENVIRONMENT_KEY,str(manifest))
    app=AppTest.from_string('from test_streamlit_eligibility import harness\nharness()').run(timeout=60)
    widget(app,'selectbox','_surface').select('independent').run(timeout=60)
    assert not app.exception
    widget(app,'selectbox','ipp_cohort_child_report').select('全部獨立基線及追蹤').run(timeout=60)
    widget(app,'selectbox','ipp_wave').select('T1').run(timeout=60)
    original=app.session_state['_rendered_rows'].copy(deep=True)
    widget(app,'multiselect','ipp_sources').set_value(['qualtrics']).run(timeout=60)
    widget(app,'checkbox',SESSION_KEY).check().run(timeout=60)
    assert not app.exception and app.session_state['ipp_sources']==['qualtrics']
    current=app.session_state['_rendered_rows']
    assert len(current)==1 and current.wave.eq('T1').all()
    assert current.anchor_source_type.eq('qualtrics').all() and current.source_type.ne(current.anchor_source_type).all()
    captures=app.session_state['_captured_downloads']
    selected=next(v for v in captures.values() if v['name'].startswith('study2-independent-parent-responses'))
    table=pd.read_csv(io.BytesIO(selected['data']));assert table.eligibility_mode.eq('on').all()
    assert tuple(app.session_state['_audit_ids'])
    widget(app,'checkbox',SESSION_KEY).uncheck().run(timeout=60)
    assert not app.exception and app.session_state['ipp_sources']==['qualtrics']
    assert len(app.session_state['_rendered_rows'])==len(original)


def test_actual_apptest_bd_and_calendar_download_membership_and_provenance(tmp_path,monkeypatch):
    manifest,data,analysis,parents,measurements,panel,summary=release_fixture(tmp_path/'release')
    monkeypatch.setenv('MANUAL_STREAMLIT_TEST_ROOT',str(manifest.parent));monkeypatch.setenv(ENVIRONMENT_KEY,str(manifest))
    app=AppTest.from_string('from test_streamlit_eligibility import harness\nharness()').run(timeout=60)
    widget(app,'radio','ds_reading').set_value('跨期比較').run(timeout=60)
    widget(app,'selectbox','dx_role').select('parent').run(timeout=60)
    widget(app,'selectbox','dx_mainstream').select(False).run(timeout=60)
    widget(app,'checkbox',SESSION_KEY).check().run(timeout=60)
    assert not app.exception
    expected={sid for members in data['retained_members']['BD'] for sid in members}
    current=app.session_state['_rendered_rows']
    assert set(current.submission_id)==expected and current.adult_continuity_confirmed.all()
    captures=app.session_state['_captured_downloads']
    pairs=pd.read_csv(io.BytesIO(captures['dx_pair_summary_csv']['data']))
    assert pairs.pairs_n.eq(len(data['retained_members']['BD'])).all()
    assert pairs.eligibility_mode.eq('on').all()
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(captures['dx_xlsx']['data']),read_only=True)
    metadata={r[0]:json.loads(r[1]) for r in book['篩選及版本'].iter_rows(min_row=2,values_only=True)}
    assert metadata['eligibility_mode']=='on' and metadata['eligibility_revision']==data['accepted_manifest']['revision']
    assert metadata['eligibility_release_id']==summary['release_id'];book.close()
    widget(app,'radio','ds_reading').set_value('時段趨勢').run(timeout=60)
    widget(app,'selectbox','ds_mainstream').select(False).run(timeout=60)
    assert not app.exception
    captures=app.session_state['_captured_downloads']
    settings=json.loads(captures['ds_trend_settings']['data'])
    assert settings['eligibility_mode']=='on' and settings['eligibility_revision']==data['accepted_manifest']['revision']
    counts=pd.read_csv(io.BytesIO(captures['ds_trend_csv_counts']['data']))
    assert counts.eligibility_mode.eq('on').all()
    expected={sid for code in ('A','C') for member in data['retained_members'][code] for sid in member}
    assert set(app.session_state['_rendered_rows'].submission_id)==expected
    excluded=pd.read_csv(io.BytesIO(captures['ds_trend_csv_exclusions']['data']))
    assert counts.sample_n.sum()+excluded.excluded_n.sum()==len(expected)
    dates=app.session_state['ds_trend_dates']
    widget(app,'checkbox',SESSION_KEY).uncheck().run(timeout=60)
    assert not app.exception and app.session_state['ds_trend_dates']==dates
    assert json.loads(app.session_state['_captured_downloads']['ds_trend_settings']['data'])['eligibility_mode']=='off'


def test_dashboard_root_wires_actual_adapter_and_both_renderers():
    source=(Path(__file__).parents[1]/'dashboard.py').read_text();tree=ast.parse(source)
    calls=[n for n in ast.walk(tree) if isinstance(n,ast.Call)]
    assert any(isinstance(n.func,ast.Name) and n.func.id=='load_eligibility' for n in calls)
    assert any(isinstance(n.func,ast.Name) and n.func.id=='render_selector' for n in calls)
    for function in ['render_descriptive','render_parent_panel']:
        call=next(n for n in calls if isinstance(n.func,ast.Name) and n.func.id==function)
        assert any(k.arg=='eligibility' and isinstance(k.value,ast.Name) and k.value.id=='eligibility' for k in call.keywords)
    assert 'eligibility.sample_rows(sample_frame,parent_frame,sample)' in source
    assert 'Published OFF/reference' in source


def legacy_release_fixture(folder,*,provenance='checksums'):
    path,data,analysis,parents,measurements,panel,summary=release_fixture(folder)
    frames=[]
    for name,frame in [('analysis.parquet',analysis),('parent_longitudinal.parquet',parents)]:
        legacy=frame.drop(columns=['version_id','observed_at','claimed_wave'],errors='ignore')
        legacy.to_parquet(path.parent/name,index=False);frames.append(legacy)
    if provenance in {'checksums','logical_gzip','artifact'}:
        objects={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in path.parent.iterdir()
                 if p.suffix=='.parquet' or p.name=='records.jsonl.gz'}
        (path.parent/'artifact_manifest.json').write_text(json.dumps(dict(format_version=1,objects=objects)))
    if provenance in {'checksums','logical_gzip'}:pin_legacy_files(path.parent,logical_gzip=provenance=='logical_gzip')
    return path,data,*frames,measurements,panel,summary


def pin_legacy_files(folder,*,logical_gzip=False):
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.iterdir() if p.is_file() and p.name!='checksums.json'}
    if logical_gzip:
        hashes['records.jsonl']=hashlib.sha256(gzip.decompress((folder/'records.jsonl.gz').read_bytes())).hexdigest()
        del hashes['records.jsonl.gz']
    (folder/'checksums.json').write_text(json.dumps(hashes))


@pytest.mark.parametrize('provenance',['none','checksums','logical_gzip','artifact'])
def test_legacy_loader_binds_exact_published_frames_and_inventory_versions(tmp_path,provenance):
    path,data,analysis,parents,measurements,panel,summary=legacy_release_fixture(tmp_path/'release',provenance=provenance)
    by_id={r['submission_id']:r for r in data['records']}
    # Same two fields added by dashboard.sample_frames, without any database.
    enriched=analysis.assign(observed_at=[by_id[sid].get('observed_at') for sid in analysis.submission_id],
                             claimed_wave=[by_id[sid].get('claimed_wave') for sid in analysis.submission_id])
    before=[f.copy(deep=True) for f in (enriched,parents)]
    fingerprint=lambda:{p.name:(hashlib.sha256(p.read_bytes()).hexdigest(),p.stat().st_mtime_ns) for p in path.parent.iterdir() if p.is_file()}
    files=fingerprint()
    state=load_eligibility(path.parent,summary,enriched.iloc[::-1],parents,manifest_path=path)
    assert state.ready,state.reason
    for mode,role in [('abcd','student'),('acbd','parent'),('independent_adult','parent')]:
        rows=state.select(True).cohort(mode,role)
        assert all(r.version_id==by_id[r.submission_id]['version_id'] for r in rows.itertuples())
    assert set(state.select(True).cohort('acbd','parent').submission_id)=={sid for pair in data['retained_members']['BD'] for sid in pair}
    assert 'version_id' not in enriched and 'version_id' not in parents
    for current,original in zip((enriched,parents),before):pd.testing.assert_frame_equal(current,original)
    assert fingerprint()==files


@pytest.mark.parametrize('fault,reason',[
    ('altered_metric','legacy_published_frame_mismatch'),
    ('mixed_scientific_row','legacy_published_frame_mismatch'),
    ('parent_changed','legacy_published_frame_mismatch'),
    ('missing_column','legacy_published_frame_schema_mismatch'),
    ('unexpected_extra','legacy_published_frame_schema_mismatch'),
    ('duplicate_row','legacy_published_frame_identity_mismatch'),
    ('unknown_id','legacy_published_frame_identity_mismatch'),
    ('partial_membership','legacy_published_frame_identity_mismatch'),
    ('wrong_observed_at','legacy_published_frame_mismatch'),
    ('wrong_claimed_wave','legacy_published_frame_mismatch'),
    ('explicit_version_conflict','analysis_version_mismatch'),
    ('duplicate_inventory','invalid_inventory_identity'),
    ('unversioned_inventory','invalid_inventory_version_or_role'),
    ('wrong_release_summary','legacy_release_summary_mismatch'),
    ('missing_summary','legacy_release_artifact_unavailable'),
    ('missing_parquet','legacy_release_artifact_unavailable'),
    ('bad_checksum','legacy_release_artifact_hash_mismatch'),
    ('unlisted_parquet','legacy_release_artifact_unpinned'),
    ('duplicate_canonical','legacy_published_frame_identity_mismatch'),
    ('changed_canonical_bytes','legacy_release_artifact_hash_mismatch'),
])
def test_legacy_loader_rejects_altered_mixed_unbound_and_conflicting_evidence(tmp_path,fault,reason):
    path,data,analysis,parents,measurements,panel,summary=legacy_release_fixture(tmp_path/'release')
    if fault=='altered_metric':analysis.loc[0,'PHQ_total']=999
    if fault=='mixed_scientific_row':analysis.loc[0,'school']='Different invented release school'
    if fault=='parent_changed':parents.loc[0,'PHQ_total']=999
    if fault=='missing_column':analysis=analysis.drop(columns='school')
    if fault=='unexpected_extra':analysis['unverified_provenance']='invented'
    if fault=='duplicate_row':analysis=pd.concat([analysis,analysis.iloc[:1]],ignore_index=True)
    if fault=='unknown_id':analysis.loc[0,'submission_id']='unknown-invented-submission'
    if fault=='partial_membership':analysis=analysis.iloc[1:]
    if fault=='wrong_observed_at':analysis['observed_at']='1999-01-01T00:00:00Z'
    if fault=='wrong_claimed_wave':analysis['claimed_wave']='invented-wrong-wave'
    if fault=='explicit_version_conflict':analysis['version_id']='invented-wrong-version'
    if fault in {'duplicate_inventory','unversioned_inventory'}:
        records=deepcopy(data['records'])
        if fault=='duplicate_inventory':records.append(deepcopy(records[0]))
        else:records[0].pop('version_id')
        with gzip.open(path.parent/'records.jsonl.gz','wt',encoding='utf-8') as stream:
            for row in records:stream.write(json.dumps(row)+'\n')
        # No hash manifests for this case: isolate inventory identity/version proof.
        (path.parent/'checksums.json').unlink();(path.parent/'artifact_manifest.json').unlink()
    if fault=='wrong_release_summary':
        changed=deepcopy(summary);changed['release_id']='invented-other-release'
        (path.parent/'summary.json').write_text(json.dumps(changed));pin_legacy_files(path.parent)
    if fault=='missing_summary':(path.parent/'summary.json').unlink()
    if fault=='missing_parquet':(path.parent/'analysis.parquet').unlink()
    if fault in {'bad_checksum','unlisted_parquet'}:
        checks=json.loads((path.parent/'checksums.json').read_text())
        if fault=='bad_checksum':checks['analysis.parquet']='0'*64
        else:checks.pop('analysis.parquet')
        (path.parent/'checksums.json').write_text(json.dumps(checks))
    if fault=='duplicate_canonical':
        analysis=pd.concat([analysis,analysis.iloc[:1]],ignore_index=True);analysis.to_parquet(path.parent/'analysis.parquet',index=False)
        (path.parent/'artifact_manifest.json').unlink();pin_legacy_files(path.parent)
    if fault=='changed_canonical_bytes':
        analysis.loc[0,'PHQ_total']=999;analysis.to_parquet(path.parent/'analysis.parquet',index=False)
    state=load_eligibility(path.parent,summary,analysis,parents,manifest_path=path)
    assert not state.ready and state.reason==reason,(fault,state.reason)
    assert not state.select(True).allowed and state.select(False).allowed
    assert str(path) not in state.reason and '@' not in state.reason


def test_pure_prepare_still_rejects_unbound_unversioned_frames():
    data,analysis,parents,measurements,panel,summary=inputs()
    from study2.analysis_eligibility import EligibilityError
    with pytest.raises(EligibilityError,match='^analysis_version_evidence_required$'):
        prepare_eligibility(data['records'],analysis.drop(columns='version_id'),parents,measurements,summary,data['accepted_manifest'],panel=panel)


def test_actual_apptest_legacy_publication_off_on_off(tmp_path,monkeypatch):
    path,data,analysis,parents,measurements,panel,summary=legacy_release_fixture(tmp_path/'release')
    monkeypatch.setenv('MANUAL_STREAMLIT_TEST_ROOT',str(path.parent));monkeypatch.setenv(ENVIRONMENT_KEY,str(path))
    app=AppTest.from_string('from test_streamlit_eligibility import harness\nharness()').run(timeout=60)
    assert not app.exception and widget(app,'checkbox',SESSION_KEY).value is False
    widget(app,'selectbox','ds_mainstream').select(False).run(timeout=60)
    original=app.session_state['_rendered_rows'].copy(deep=True)
    widget(app,'checkbox',SESSION_KEY).check().run(timeout=60)
    assert not app.exception and app.session_state['_active_allowed'] and app.session_state['_eligibility_mode']=='on'
    expected={sid for code in ('A','C') for member in data['retained_members'][code] for sid in member}
    assert set(app.session_state['_rendered_rows'].submission_id)==expected
    assert app.session_state['_renderer_metadata']['eligibility_revision']==data['accepted_manifest']['revision']
    widget(app,'checkbox',SESSION_KEY).uncheck().run(timeout=60)
    assert not app.exception and app.session_state['_eligibility_mode']=='off'
    pd.testing.assert_frame_equal(app.session_state['_rendered_rows'],original)
    assert 'version_id' not in pd.read_parquet(path.parent/'analysis.parquet')
