import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from study2 import review_core as core
from study2.common import PipelineError, dumps, utcnow
from study2.matching import candidate
from study2.review import validation_report
from study2.review_excel import import_package, read_package, comparison
from study2.store import connect, writer_lock, ingest
from test_identity_and_time import rec
from test_store_and_review import database, published, source


@pytest.fixture
def project(tmp_path):
    cfg={"state_dir":str(tmp_path),"minimum_followup_days":90,"duplicate_days":30}
    db=database(tmp_path)
    rows=[rec('a',child_id='child'),rec('b',120,child_id='child'),
          rec('p',role='parent',child_id='child'),rec('q',120,role='parent',child_id='child')]
    for r in rows[2:]:
        r['answer'].update(adult_demographics_05='1980',adult_demographics_06='1',contact_02='61234567')
    c=published(db,cfg,rows)
    db.execute("UPDATE release_candidates SET payload_json=?",(dumps({**c,'status':'pending'}),))
    family=candidate(rows[0],rows[2],'family','family_exact','confirmed','synthetic')
    adult=candidate(rows[2],rows[3],'respondent','adult_rule','pending','synthetic')
    for item in [family,adult]:
        db.execute('INSERT INTO release_candidates VALUES(?,?,?)',('release',item['case_id'],dumps(item)))
    db.commit()
    db.close()
    return cfg,[c['case_id'],family['case_id'],adult['case_id']]


def package(project,reviewer='A',ids=None,path=None,independent=False):
    cfg,cases=project
    return core.create_package(cfg,reviewer,path,case_ids=ids or cases,independent=independent)


def answers(pkg,verdict='確認相同',reason='已核對合成收集紀錄第 2 行'):
    return [{'task_id':t['task_id'],'verdict':verdict,'evidence_source':'既有收集紀錄','reason':reason} for t in pkg['tasks']]


def submit(cfg,pkg,verdict='確認相同'):
    return core.submit_reviews(cfg,pkg['package_id'],answers(pkg,verdict))


def approved(cfg,pkg,verdict='確認相同'):
    submit(cfg,pkg,verdict)
    batch=core.preview_batch(cfg)
    assert not batch['conflicts']
    core.approve_batch(cfg,batch['batch_id'],reason='核對無誤')
    return batch


def test_staging_and_approval_do_not_touch_main(project):
    cfg,_=project
    pkg=package(project)
    batch=approved(cfg,pkg)
    main=connect(Path(cfg['state_dir']))
    assert main.execute('SELECT COUNT(*) FROM decisions').fetchone()[0]==0
    assert main.execute('SELECT COUNT(*) FROM labels').fetchone()[0]==0
    assert core.apply_batch(cfg,batch['batch_id'])['corrections']==2
    assert main.execute('SELECT COUNT(*) FROM labels').fetchone()[0]==3
    assert {json.loads(r[0])['action'] for r in main.execute('SELECT payload_json FROM decisions')}=={'LINK','LINK_ADULT'}
    assert core.apply_batch(cfg,batch['batch_id'])['recovered_or_repeated']
    assert main.execute('SELECT COUNT(*) FROM decisions').fetchone()[0]==2
    main.close()


def test_grade_hint_card_keeps_inputs_usable_and_rejects_hint_tampering(project):
    cfg, cases = project
    main = connect(Path(cfg['state_dir']))
    records = []
    for sid in ('a', 'b'):
        r = json.loads(main.execute('SELECT payload_json FROM normalized_records WHERE submission_id=?', (sid,)).fetchone()[0])
        r['answer']['school_type'] = 'primary'
        r['grade_raw'] = str(r['grade'])
        if sid == 'b':
            r['grade'] = 4
            r['grade_raw'] = '4'
        main.execute('UPDATE normalized_records SET payload_json=? WHERE submission_id=?', (dumps(r), sid))
        records.append(r)
    updated = candidate(*records, 'identity', 'exact_school_contact_child_demo', 'confirmed', 'Synthetic grade hint')
    main.execute('UPDATE release_candidates SET payload_json=? WHERE case_id=?', (dumps(updated), cases[0]))
    main.commit(); main.close()
    path = Path(cfg['state_dir'])/'hint_card.xlsx'
    pkg = package(project, ids=[cases[0]], path=path)
    wb = load_workbook(path)
    sheet = wb['審核卡']
    assert '年級上升 1 級' in sheet['A33'].value
    assert sheet['A33'].protection.locked
    positions = pkg['contract']['inputs'][pkg['tasks'][0]['task_id']]
    assert not sheet[positions['verdict']].protection.locked
    sheet[positions['verdict']] = '確認相同'
    sheet[positions['evidence_source']] = '問卷資料比較'
    sheet[positions['reason']] = '提示不代表拒配，已核實合成身份證據'
    sheet['A33'] = 'changed evidence'
    wb.save(path); wb.close()
    with pytest.raises(PipelineError):
        import_package(cfg, path)


def test_idempotent_and_revision_history(project):
    cfg,_=project
    pkg=package(project)
    assert submit(cfg,pkg)['saved']==3
    assert submit(cfg,pkg)['unchanged']==3
    assert submit(cfg,pkg,'確認不同')['saved']==3
    db=core.workspace(cfg)
    assert db.execute('SELECT COUNT(*) FROM reviews').fetchone()[0]==6
    assert len(core.latest_reviews(db))==3
    db.close()


@pytest.mark.parametrize('field,value',[('reason',''),('verdict','非法'),('evidence_source','未知')])
def test_invalid_partial_all_or_nothing(project,field,value):
    cfg,_=project
    pkg=package(project)
    items=answers(pkg)
    items[1][field]=value
    with pytest.raises(PipelineError):
        core.submit_reviews(cfg,pkg['package_id'],items)
    db=core.workspace(cfg)
    assert db.execute('SELECT COUNT(*) FROM reviews').fetchone()[0]==0
    db.close()


def test_blank_card_and_unresolved_preserve_pair(project):
    cfg,cases=project
    pkg=package(project,ids=[cases[1]])
    assert core.submit_reviews(cfg,pkg['package_id'],[{'task_id':pkg['tasks'][0]['task_id']}])['saved']==0
    batch=approved(cfg,pkg,'資料不足')
    assert core.apply_batch(cfg,batch['batch_id'])['corrections']==0
    main=connect(Path(cfg['state_dir']))
    assert main.execute('SELECT COUNT(*) FROM decisions').fetchone()[0]==0
    assert json.loads(main.execute('SELECT payload_json FROM release_candidates WHERE case_id=?',(cases[1],)).fetchone()[0])['status']=='confirmed'
    assert core.review_report(cfg)['followups_pending']==1
    main.close()


def test_assignments_and_independent_review(project):
    cfg,cases=project
    p1=package(project,ids=[cases[0]])
    with pytest.raises(PipelineError,match='有效任務'):
        package(project,ids=[cases[0]])
    with pytest.raises(PipelineError,match='另一位'):
        package(project,ids=[cases[0]],independent=True)
    p2=package(project,'B',[cases[0]],independent=True)
    submit(cfg,p1)
    submit(cfg,p2,'確認不同')
    batch=core.preview_batch(cfg,[p1['package_id']])
    assert not batch['rows'] and len(batch['conflicts'])==1
    assert core.review_report(cfg)['independent_review']['disagreement']==1
    chosen=batch['conflicts'][0]['reviews'][0]['review_id']
    batch2=core.preview_batch(cfg,adjudications={cases[0]:{'review_id':chosen,'reason':'Jarvis 查核合成原始紀錄'}})
    assert len(batch2['rows'])==1 and not batch2['conflicts']
    assert 'alternatives' in batch2['rows'][0]['adjudication']


def test_approved_invalid_after_reviewer_revision(project):
    cfg,_=project
    pkg=package(project)
    batch=approved(cfg,pkg)
    submit(cfg,pkg,'資料不足')
    with pytest.raises(PipelineError,match='失效'):
        core.apply_batch(cfg,batch['batch_id'])


def test_new_independent_answer_invalidates_preview(project):
    cfg,cases=project
    p1=package(project,ids=[cases[0]])
    p2=package(project,'B',[cases[0]],independent=True)
    batch=approved(cfg,p1)
    submit(cfg,p2,'確認不同')
    with pytest.raises(PipelineError,match='失效'):
        core.apply_batch(cfg,batch['batch_id'])


def new_release(cfg,change=None):
    db=connect(Path(cfg['state_dir']))
    db.execute("INSERT INTO runs VALUES('next','fp',?,'published','{}','{}',NULL)",(utcnow(),))
    for row in db.execute("SELECT submission_id,payload_json FROM normalized_records WHERE run_id='run'").fetchall():
        payload=json.loads(row[1])
        if change and payload['submission_id']=='a':
            payload.update(change)
        db.execute('INSERT INTO normalized_records VALUES(?,?,?)',('next',row[0],dumps(payload)))
    db.execute("INSERT INTO releases VALUES('next_release','next',?,?, '{}')",(utcnow(),cfg['state_dir']))
    db.execute("INSERT INTO release_candidates SELECT 'next_release',case_id,payload_json FROM release_candidates WHERE release_id='release'")
    db.execute("UPDATE publication SET release_id='next_release'")
    db.commit()
    db.close()


def test_unchanged_release_can_carry_but_approval_version_still_frozen(project):
    cfg,_=project
    pkg=package(project)
    new_release(cfg)
    batch=approved(cfg,pkg)
    assert len(batch['carried_tasks'])==3
    assert core.apply_batch(cfg,batch['batch_id'])['state']=='applied'
    db=core.workspace(cfg)
    assert db.execute('SELECT COUNT(*) FROM carries').fetchone()[0]==3
    db.close()


@pytest.mark.parametrize('change',[{'dob':'2011-01-01'},{'child_id':'different'},{'selected_wave':'T1'}])
def test_changed_evidence_requires_fresh_task(project,change):
    cfg,_=project
    pkg=package(project)
    new_release(cfg,change)
    with pytest.raises(PipelineError,match='重新審核'):
        submit(cfg,pkg)


def test_publish_change_after_approval_blocks_apply(project):
    cfg,_=project
    batch=approved(cfg,package(project))
    new_release(cfg)
    with pytest.raises(PipelineError,match='失效'):
        core.apply_batch(cfg,batch['batch_id'])


@pytest.mark.parametrize('change_kind', ['evidence', 'new_member'])
def test_competing_candidate_group_changes_require_new_review(project, change_kind):
    cfg, cases = project
    db = connect(Path(cfg['state_dir']))
    x, y = rec('x', child_id='other_child'), rec('y', child_id='other_child')
    ingest(db, 'run', [source(x), source(y)])
    for row in [x, y]:
        db.execute('INSERT INTO normalized_records VALUES(?,?,?)', ('run', row['submission_id'], dumps(row)))
    a = json.loads(db.execute("SELECT payload_json FROM normalized_records WHERE submission_id='a'").fetchone()[0])
    rival = candidate(a, x, 'identity', 'competing_candidate', 'pending', 'synthetic')
    db.execute('INSERT INTO release_candidates VALUES(?,?,?)', ('release', rival['case_id'], dumps(rival)))
    db.commit(); db.close()
    pkg = package(project, ids=[cases[0]])
    assert {'x', 'y'} <= pkg['tasks'][0]['snapshot']['records'].keys()
    new_release(cfg)
    db = connect(Path(cfg['state_dir']))
    if change_kind == 'evidence':
        y['dob'] = '2011-01-01'
        db.execute("UPDATE normalized_records SET payload_json=? WHERE run_id='next' AND submission_id='y'", (dumps(y),))
    else:
        z = rec('z', child_id='other_child')
        ingest(db, 'next', [source(z)])
        db.execute('INSERT INTO normalized_records VALUES(?,?,?)', ('next', 'z', dumps(z)))
    db.commit(); db.close()
    with pytest.raises(PipelineError, match='重新審核'):
        submit(cfg, pkg)


def test_pipeline_busy_can_stage_but_cannot_apply(project):
    cfg,_=project
    pkg=package(project)
    with writer_lock(Path(cfg['state_dir'])):
        batch=approved(cfg,pkg)
        assert core.apply_batch(cfg,batch['batch_id'])['state']=='approved'
    assert core.apply_batch(cfg,batch['batch_id'])['state']=='applied'


def test_recovery_after_main_commit(project,monkeypatch):
    cfg,_=project
    batch=approved(cfg,package(project))
    original=core._finish_apply
    def crash(*args):
        raise RuntimeError('synthetic process crash after canonical commit')
    monkeypatch.setattr(core,'_finish_apply',crash)
    with pytest.raises(RuntimeError):
        core.apply_batch(cfg,batch['batch_id'])
    monkeypatch.setattr(core,'_finish_apply',original)
    assert core.apply_batch(cfg,batch['batch_id'])['recovered_or_repeated']
    db=connect(Path(cfg['state_dir']))
    assert db.execute('SELECT COUNT(*) FROM review_applications').fetchone()[0]==1
    assert db.execute('SELECT COUNT(*) FROM decisions').fetchone()[0]==2
    db.close()


def test_backup_restores_unpublished_reviews(project):
    cfg,_=project
    submit(cfg,package(project))
    path=core.backup_workspace(cfg,'test')
    db=sqlite3.connect(path)
    assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert db.execute('SELECT COUNT(*) FROM reviews').fetchone()[0]==3
    assert db.execute("SELECT COUNT(*) FROM batches WHERE state='applied'").fetchone()[0]==0
    db.close()


def fill(path,pkg,partial=False):
    wb=load_workbook(path)
    for coordinates in pkg['contract']['inputs'].values():
        for k,v in {'verdict':'確認相同','evidence_source':'問卷資料比較','reason':'' if partial else '已核對合成來源第 2 行'}.items():
            wb['審核卡'][coordinates[k]]=v
    wb.save(path)
    wb.close()


def test_native_workbook_structure_and_roundtrip(project,tmp_path):
    cfg,_=project
    path=tmp_path/'cards.xlsx'
    pkg=package(project,path=path)
    wb=load_workbook(path)
    assert wb.sheetnames==['審核目錄','審核卡','群組參考','系統資料']
    assert wb['系統資料'].sheet_state=='veryHidden'
    assert all(s.protection.sheet for s in wb)
    assert len(wb['審核卡'].data_validations.dataValidation)==6
    assert wb['審核卡']['A3'].hyperlink.location=="'審核目錄'!A12"
    # family card at row 41; always student on left despite canonical pair sort.
    assert wb['審核卡']['B45'].value=='學生'
    assert wb['審核卡']['C45'].value=='家長'
    assert not wb['審核卡']['B28'].protection.locked
    assert len(list(wb['群組參考'].values))>3
    wb.close()
    assert import_package(cfg,path)['saved']==0
    fill(path,pkg)
    assert import_package(cfg,path)['saved']==3
    assert import_package(cfg,path)['unchanged']==3


@pytest.mark.parametrize('mutation',['evidence','formula','id','merge','delete','extra','lock','option'])
def test_tampered_workbook_rejected(project,tmp_path,mutation):
    cfg,_=project
    path=tmp_path/'cards.xlsx'
    pkg=package(project,path=path)
    wb=load_workbook(path)
    s=wb['審核卡']
    if mutation=='evidence':s['B6']='modified school'
    if mutation=='formula':s['B32']='已填寫'
    if mutation=='id':wb['系統資料']['B2']='modified package'
    if mutation=='merge':s.unmerge_cells('A1:D1')
    if mutation=='delete':s.delete_rows(41,40)
    if mutation=='extra':s['A999']='added data'
    if mutation=='lock':s.protection.sheet=False
    if mutation=='option':s.data_validations.dataValidation[0].formula1='"wrong"'
    wb.save(path)
    wb.close()
    with pytest.raises(PipelineError):import_package(cfg,path)


def test_incomplete_excel_returns_exact_fields(project,tmp_path):
    cfg,_=project
    path=tmp_path/'cards.xlsx'
    pkg=package(project,path=path)
    fill(path,pkg,partial=True)
    with pytest.raises(PipelineError,match='補充理由'):import_package(cfg,path)


def test_missing_not_equal():
    assert comparison(None,None)=='兩方缺失'
    assert comparison('',None)=='兩方缺失'
    assert comparison('0',None)=='一方缺失'
    assert comparison('0','0')=='相同'


def test_original_prediction_preserved_after_manual_correction(project):
    cfg,cases=project
    pkg=package(project,ids=[cases[0]])
    batch=approved(cfg,pkg)
    core.apply_batch(cfg,batch['batch_id'])
    db=connect(Path(cfg['state_dir']))
    corrected={**pkg['tasks'][0]['snapshot']['case'],'status':'confirmed','rule':'manual_link'}
    result=validation_report(db,[corrected])
    assert result[0]['rule']=='exact' and result[0]['missed_match']==1
    db.close()


def test_known_nonmatch_publish_guard(project):
    cfg,cases=project
    pkg=package(project,ids=[cases[1]])
    batch=approved(cfg,pkg,'確認不同')
    core.apply_batch(cfg,batch['batch_id'])
    db=connect(Path(cfg['state_dir']))
    with pytest.raises(PipelineError,match='NONMATCH'):
        validation_report(db,[pkg['tasks'][0]['snapshot']['case']])
    db.close()


def test_adult_requires_confirmed_child_and_consistent_demographics(project):
    cfg,cases=project
    db=connect(Path(cfg['state_dir']))
    row=json.loads(db.execute("SELECT payload_json FROM normalized_records WHERE submission_id='q'").fetchone()[0])
    row['answer']['adult_demographics_05']='1981'
    db.execute("UPDATE normalized_records SET payload_json=? WHERE submission_id='q'",(dumps(row),))
    db.commit();db.close()
    pkg=package(project,ids=[cases[2]])
    submit(cfg,pkg)
    batch=core.preview_batch(cfg)
    assert not batch['rows'] and '成人身份資料矛盾' in batch['conflicts'][0]['message']


def test_manual_opposite_requires_adjudication(project):
    cfg,cases=project
    pkg=package(project,ids=[cases[0]])
    batch=approved(cfg,pkg)
    core.apply_batch(cfg,batch['batch_id'])
    submit(cfg,pkg,'確認不同')
    preview=core.preview_batch(cfg)
    assert '既有人工決定相反' in preview['conflicts'][0]['message']


def test_package_reassign_before_submission(project):
    cfg,cases=project
    p=package(project,ids=[cases[0]])
    replacement=core.create_package(cfg,'B',replaces=p['package_id'])
    assert replacement['reviewer']=='B'
    with pytest.raises(PipelineError):submit(cfg,p)


def test_rules_change_requires_reissue(project):
    cfg,_=project
    p=package(project)
    cfg['duplicate_days']=29
    with pytest.raises(PipelineError,match='規則已改變'):submit(cfg,p)


def test_old_apply_retry_does_not_hide_later_revision(project):
    cfg,cases=project
    p=package(project,ids=[cases[0]])
    batch=approved(cfg,p)
    core.apply_batch(cfg,batch['batch_id'])
    submit(cfg,p,'確認不同')
    core.apply_batch(cfg,batch['batch_id'])
    next_batch=core.preview_batch(cfg)
    assert len(next_batch['conflicts'])==1 and '相反' in next_batch['conflicts'][0]['message']


def test_stratification_covers_types_and_statuses():
    rows=[]
    for kind in core.KINDS:
        for status in ['confirmed','pending','rejected']:
            for rule in range(12):
                rows.append({'case_id':f'{kind}_{status}_{rule}','kind':kind,'status':status,'rule':str(rule)})
    selected=core.stratified_cases(rows,'A',20)
    assert len(selected)==20
    assert {r['kind'] for r in selected}==set(core.KINDS)
    assert {r['status'] for r in selected}=={'confirmed','pending','rejected'}
    assert core.stratified_cases(rows,'A',20)==selected


def test_non_identity_source_update_and_reorder_can_carry(project):
    cfg,_=project
    p=package(project)
    new_release(cfg,{'version_id':'new_version','location':{'file':'new_export.xlsx','row':999,'sha256':'new_hash'}})
    assert submit(cfg,p)['saved']==3
    preview=core.preview_batch(cfg)
    assert len(preview['rows'])==3 and not preview['conflicts']


def test_known_manual_decision_change_invalidates_preview(project):
    cfg,cases=project
    p=package(project,ids=[cases[0]])
    batch=approved(cfg,p)
    db=connect(Path(cfg['state_dir']))
    case=candidate(rec('p',role='parent'),rec('q',role='parent'),'respondent','external','pending','synthetic')
    db.execute('INSERT INTO labels VALUES(?,?,?,?,?,?)',(case['case_id'],'UNRESOLVED','Jarvis','外部工具新决定',utcnow(),dumps(case)))
    db.commit();db.close()
    with pytest.raises(PipelineError,match='人工決定'):
        core.apply_batch(cfg,batch['batch_id'])


def test_transitive_unlink_conflicts_isolated_from_adult_case(project):
    cfg,cases=project
    db=connect(Path(cfg['state_dir']))
    records={r[0]:json.loads(r[1]) for r in db.execute('SELECT submission_id,payload_json FROM normalized_records')}
    extra=candidate(records['b'],records['p'],'family','new_family','pending','synthetic')
    db.execute('INSERT INTO release_candidates VALUES(?,?,?)',('release',extra['case_id'],dumps(extra)))
    db.commit();db.close()
    p=package(project,ids=[cases[0],cases[1],extra['case_id'],cases[2]])
    items=answers(p)
    items[1]['verdict']='確認不同'
    core.submit_reviews(cfg,p['package_id'],items)
    preview=core.preview_batch(cfg)
    assert len(preview['conflicts'])==3
    assert len(preview['rows'])==1 and preview['rows'][0]['kind']=='respondent'
    core.approve_batch(cfg,preview['batch_id'],reason='只批准無衝突成人個案')
    assert core.apply_batch(cfg,preview['batch_id'])['corrections']==1


def test_new_candidate_reopens_old_task_on_next_export(project):
    cfg,cases=project
    p=package(project,ids=[cases[0]])
    new_release(cfg,{'dob':'2011-01-01'})
    replacement=package(project,ids=[cases[0]])
    assert replacement['tasks'][0]['task_id']!=p['tasks'][0]['task_id']
    db=core.workspace(cfg)
    assert db.execute('SELECT state FROM tasks WHERE id=?',(p['tasks'][0]['task_id'],)).fetchone()[0]=='stale'
    db.close()


def adult_difference(project, field='adult_demographics_05', values=('1980','1981')):
    cfg,cases=project
    db=connect(Path(cfg['state_dir']))
    for sid,value in zip(('p','q'),values):
        row=json.loads(db.execute('SELECT payload_json FROM normalized_records WHERE submission_id=?',(sid,)).fetchone()[0])
        row['answer'][field]=value
        if field=='adult_demographics_06.1':
            row['answer']['adult_demographics_06']='2'
        db.execute('UPDATE normalized_records SET payload_json=? WHERE submission_id=?',(dumps(row),sid))
    db.commit();db.close()
    p=package(project,ids=[cases[2]])
    submit(cfg,p)
    work=core.workspace(cfg)
    rid=core.latest_reviews(work)[0]['id'];work.close()
    return cfg,p,cases[2],rid


def adult_resolution(rid,kind='accepted_typo_unknown_year'):
    return {'review_id':rid,'reason':'Jarvis 根據主管已確認的同一成人裁定',
            'adult_identity_exception':{'type':kind,'authority':'研究主管，由 Jarvis 記錄',
                                        'reason':'主管接受出生年份錯字，正確年份未確認'}}


@pytest.mark.parametrize('field,values,kind',[
    ('adult_demographics_05',('1980','1981'),'accepted_typo_unknown_year'),
    ('adult_demographics_06.1',('姑侄','姑姐'),'equivalent_relationship_wording')])
def test_operator_adult_exception_retains_evidence(project,field,values,kind):
    cfg,p,cid,rid=adult_difference(project,field,values)
    assert len(core.preview_batch(cfg,[p['package_id']])['conflicts'])==1
    preview=core.preview_batch(cfg,[p['package_id']],{cid:adult_resolution(rid,kind)})
    assert not preview['conflicts'] and len(preview['rows'])==1
    core.approve_batch(cfg,preview['batch_id'],reason='只批准這宗已裁定的身份例外')
    assert core.apply_batch(cfg,preview['batch_id'])['corrections']==1
    db=connect(Path(cfg['state_dir']))
    for sid,value in zip(('p','q'),values):
        row=json.loads(db.execute('SELECT payload_json FROM normalized_records WHERE submission_id=?',(sid,)).fetchone()[0])
        assert row['answer'][field]==value
    decision=json.loads(db.execute('SELECT payload_json FROM decisions').fetchone()[0])
    assert decision['adult_identity_exception']['type']==kind
    if kind=='accepted_typo_unknown_year':
        assert decision['adult_identity_exception']['correct_birth_year'] is None
    db.close()


def test_adult_exception_cannot_waive_changed_parent_relationship(project):
    cfg,p,cid,rid=adult_difference(project,'adult_demographics_06',('0','1'))
    preview=core.preview_batch(cfg,[p['package_id']],{cid:adult_resolution(rid)})
    assert not preview['rows'] and len(preview['conflicts'])==1


def test_exception_does_not_waive_transitive_adult_conflict(project):
    from copy import deepcopy
    cfg,p,cid,rid=adult_difference(project)
    db=connect(Path(cfg['state_dir']),readonly=True);ctx=core.Context(db,cfg)
    r=deepcopy(ctx.records['q']);r['submission_id']='r';r['answer']['adult_demographics_05']='1982';ctx.records['r']=r
    rows=[]
    for a,b in [('p','q'),('q','r')]:
        row={'record_a':a,'record_b':b,'kind':'respondent','action':'LINK_ADULT'}
        row['adult_identity_exception']=core._adult_exception(ctx,row,adult_resolution(rid)['adult_identity_exception'])
        rows.append(row)
    with pytest.raises(PipelineError,match='連鎖'):
        core._check_group_conflicts(ctx,rows)
    db.close()


def test_adult_exception_reopens_after_child_assignment_change(project):
    from study2.respondents import adult_followups
    cfg,p,cid,rid=adult_difference(project)
    preview=core.preview_batch(cfg,[p['package_id']],{cid:adult_resolution(rid)})
    core.approve_batch(cfg,preview['batch_id'],reason='已裁定')
    core.apply_batch(cfg,preview['batch_id'])
    db=connect(Path(cfg['state_dir']),readonly=True)
    decisions=core.active_decisions(db)
    rows=[json.loads(r[0]) for r in db.execute("SELECT payload_json FROM normalized_records WHERE submission_id IN ('p','q')")]
    next(r for r in rows if r['submission_id']=='p')['selected_wave']='T0'
    pairs,cases,events=adult_followups(rows,decisions)
    assert len(pairs)==1 and all(r['adult_birth_year_status']=='unconfirmed' for r in rows)
    from study2.outputs import export_row
    baseline=next(r for r in rows if r['submission_id']=='p')
    baseline['child_id']='CHILD_test'
    assert export_row(baseline)['adult_birth_year_status']=='unconfirmed'
    for r in rows:r['child_id']='different_child'
    pairs,cases,events=adult_followups(rows,decisions)
    assert not pairs and events[0]['status']=='reopened'
    db.close()
