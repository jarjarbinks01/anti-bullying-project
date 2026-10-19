"""Hand-calculated denominators, bins and paired populations, not implementation mirrors."""
import io
import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from study2.common import PipelineError
from study2.descriptive_dashboard import filtered_workbook
from study2.descriptive_exploration import (comparison_variables, explore, numeric_summary, shared_bins)
from study2.descriptive_exploration_outputs import available, verify_exploration, write_exploration
from study2.descriptives import cohort
from test_descriptives import inputs
from test_pipeline_integration import project


def panel(a,b,mode='acbd',variable='PHQ_total'):
    rows=[]
    for i,(x,y) in enumerate(zip(a,b)):
        for wave,value in [('T0',x),('T1',y)]:
            rows.append(dict(child_id=str(i),submission_id=str(i)+wave,wave=wave,report_mode=mode,
                grouping_anchor='學生T0',anchor_gender='0' if i%2 else '1',anchor_school_stage='primary',
                **{variable:value}))
    return pd.DataFrame(rows)


def test_paired_valid_means_are_not_wave_valid_means_and_four_missing_states():
    rows=panel([0,2,10,None,None],[2,4,None,8,None])
    paired=explore(rows,'student',release_id='s',paired=True,selected_variables=['PHQ_total'])
    r=paired['pair_summary'].iloc[0]
    assert [r[k] for k in ['pairs_n','paired_valid_n','only_t0_valid_n','only_t1_valid_n','both_missing_n']]==[5,2,1,1,1]
    assert r['t0_mean']==1 and r['t1_mean']==3 and r['mean_change']==2 and r['sd_change']==0
    assert (r['up_n'],r['same_n'],r['down_n'])==(2,0,0)
    waves=explore(rows,'student',selected_variables=['PHQ_total'])['distribution_summary']
    assert waves.iloc[0]['mean']==4 and waves.iloc[1]['mean']==14/3
    assert all(paired['distribution_summary']['valid_n']==2)
    for row in paired['distribution_summary'].to_dict('records'):
        assert sum(json.loads(row['missing_detail']).values())==row['missing_n']
    assert json.loads(paired['distribution_summary'].iloc[0]['missing_detail'])=={'missing':2,'other_wave_missing_or_invalid':1}
    assert paired['histograms'].groupby('wave')['count'].sum().to_dict()=={'T0':2,'T1':2,'delta':2}


def test_linear_quantiles_whiskers_bins_empty_single_equal_and_unexpected():
    r=numeric_summary([0,1,2,3,100])
    assert (r['q1'],r['median'],r['q3'],r['lower_whisker'],r['upper_whisker'],r['outside_whisker_n'])==(1,2,3,0,3,1)
    assert r['mean']==21.2 and r['maximum']==100
    assert numeric_summary([])['mean'] is None
    assert numeric_summary([7])['sd'] is None
    assert numeric_summary([4,4,4])['q1']==4
    assert list(shared_bins([[0,2],[1,3]]))==[-.5,.5,1.5,2.5,3.5]
    edges=shared_bins([[.15,.2],[.3,.35]])
    assert edges[0]==.15 and edges[-1]==.35 and len(edges)-1<=20
    assert np.histogram([.15,.2,.3,.35],edges)[0].sum()==4
    result=explore(panel([0,100],[1,100]),'student',paired=True,selected_variables=['PHQ_total'])
    assert result['distribution_summary'].query('wave == "T0"').iloc[0]['unexpected_display_range_n']==1
    assert result['histograms'].query('wave == "T0"')['count'].sum()==2
    assert set(result['histograms'].query('wave != "delta"')['bin_left'])==set(np.arange(-.5,100))


def test_zero_change_tolerance_does_not_round_original_mean():
    result=explore(panel([1.,1.,2.],[1.+1e-10,1.-1e-10,1.]),'student',paired=True,selected_variables=['PHQ_total'])
    r=result['pair_summary'].iloc[0]
    assert (r['same_n'],r['up_n'],r['down_n'])==(2,0,1)
    assert r['mean_change']==pytest.approx(-1/3)


def test_binary_four_cells_and_percentage_points():
    result=explore(panel([0,1,0,1,1],[0,0,1,1,None],variable='TBV_prevalence2'),
                   'student',paired=True,selected_variables=['TBV_prevalence2'])
    assert result['pair_transitions']['count'].tolist()==[1,1,1,1]
    r=result['pair_summary'].iloc[0]
    assert r['paired_valid_n']==4 and r['t0_proportion']==.5 and r['t1_proportion']==.5
    assert r['change_percentage_points']==0
    assert 'mean_change' not in result['pair_summary']


def test_category_bands_no_numeric_delta_and_adult_phq_guard():
    result=explore(panel([0,4,7],[1,6,12]),'student',paired=True,selected_variables=['phq_band'])
    assert result['pair_summary']['paired_valid_n'].iloc[0]==3
    assert result['pair_transitions'].empty and result['distribution_summary'].empty
    for mode in ['efmatched','independent_child_report']:
        assert not any(v.key.startswith('PHQ_') or v.key=='phq_band' for v in comparison_variables(mode,'parent'))
        with pytest.raises(PipelineError,match='unavailable'):
            explore(panel([1],[2],mode),'parent',paired=True,selected_variables=['PHQ_total'])
    for mode in ['acbd','independent_same_adult']:
        rows = panel([0], [12], mode)
        assert not any(v.key=='PHQ_total' for v in comparison_variables(mode,'parent',rows))
        rows['adult_continuity_confirmed'] = True
        assert any(v.key=='PHQ_total' for v in comparison_variables(mode,'parent',rows))
        assert explore(rows,'parent',paired=True,selected_variables=['PHQ_total'])['pair_summary'].iloc[0]['mean_change']==12


@pytest.mark.parametrize('corrupt',['duplicate','missing','same_submission','null','unknown_wave','wrong_mode'])
def test_bad_formal_structure_is_never_silently_repaired(corrupt):
    rows=panel([0,1],[2,3])
    if corrupt=='duplicate':rows=pd.concat([rows,rows.iloc[:1]])
    if corrupt=='missing':rows=rows.iloc[:-1]
    if corrupt=='same_submission':rows.loc[1,'submission_id']=rows.loc[0,'submission_id']
    if corrupt=='null':rows.loc[0,'child_id']=None
    if corrupt=='unknown_wave':rows.loc[0,'wave']='T2'
    if corrupt=='wrong_mode':rows['report_mode']='abcd'
    with pytest.raises(PipelineError):explore(rows,'student',paired=True,selected_variables=['PHQ_total'])


def test_shared_subgroup_bins_and_unknown_stage_kept_in_all():
    rows=panel([0,2,8],[1,3,9]);rows.loc[rows['child_id'].eq('2'),'anchor_school_stage']=None
    result=explore(rows,'student',('all','male','female','primary','secondary'),paired=True,selected_variables=['PHQ_total'])
    counts=result['pair_summary'].set_index('group')['pairs_n'].to_dict()
    assert counts=={'all':3,'male':1,'female':2,'primary':2,'secondary':0}
    assert result['histograms'].query('wave != "delta"').groupby(['group','wave'])['bin_left'].apply(list).map(tuple).nunique()==1
    assert result['histograms'].query('group == "secondary"')['count'].sum()==0


def test_both_missing_pair_and_one_pair():
    rows=panel([None],[None]);rows['PHQ_total']=np.nan
    result=explore(rows,'student',paired=True,selected_variables=['PHQ_total'])
    assert result['histograms'].empty and result['pair_summary'].iloc[0]['both_missing_n']==1
    one=explore(panel([1],[2]),'student',paired=True,selected_variables=['PHQ_total'])
    assert one['distribution_summary']['sd'].isna().all()
    assert one['pair_summary']['sd_change'].isna().all()


def test_chart_keeps_t0_t1_at_distinct_named_positions_and_nine_groups_readable():
    from study2.descriptive_charts import distribution_chart,category_chart
    from study2.descriptive_spec import GROUPS,variables
    rows=panel([0,1],[1,2])
    result=explore(rows,'student',paired=True,selected_variables=['PHQ_total'])
    var=next(v for v in variables('student') if v.key=='PHQ_total')
    figure=distribution_chart(result['distribution_summary'].query('wave != "delta"'),var)
    assert {trace.x[0] for trace in figure.data}=={'全部 T0','全部 T1'}
    assert all(trace.q1 is not None for trace in figure.data)
    result=explore(rows,'student',tuple(GROUPS),paired=True,selected_variables=['phq_band'])
    band=next(v for v in variables('student') if v.key=='phq_band')
    assert category_chart(result['categories'],band).layout.height>=1500


def test_histogram_separates_waves_with_common_axes_and_preserves_counts():
    from study2.descriptive_charts import distribution_chart
    from study2.descriptive_spec import variables
    result=explore(panel([0,0,3],[1,2,2]),'student',paired=True,selected_variables=['PHQ_total'])
    histogram=result['histograms'].query('wave != "delta"').copy()
    before=histogram.copy(deep=True)
    var=next(v for v in variables('student') if v.key=='PHQ_total')
    figure=distribution_chart(result['distribution_summary'],var,histogram=histogram)
    assert len(figure.data)==2
    first,last=figure.data
    assert (first.xaxis,last.xaxis)==('x','x2')
    assert (first.yaxis,last.yaxis)==('y','y2')
    assert list(first.x)==list(last.x)==[0,1,2,3]
    assert list(first.y)==[2,0,0,1] and list(last.y)==[0,1,2,0]
    assert figure.layout.xaxis.range==figure.layout.xaxis2.range==(-.5,3.5)
    assert figure.layout.yaxis.range==figure.layout.yaxis2.range
    assert figure.layout.xaxis.matches=='x2' and figure.layout.yaxis.matches=='y2'
    assert figure.layout.yaxis.dtick==figure.layout.yaxis2.dtick==1
    assert figure.layout.yaxis.range[0]==0 and figure.layout.yaxis.range[1]>2
    assert [label.text for label in figure.layout.annotations]==['全部 · T0（有效 N=3）','全部 · T1（有效 N=3）']
    pd.testing.assert_frame_equal(histogram,before)


def test_single_delta_histogram_keeps_one_panel_and_zero_reference():
    from study2.descriptive_charts import distribution_chart
    from study2.descriptive_spec import variables
    result=explore(panel([1,3],[3,2]),'student',paired=True,selected_variables=['PHQ_total'])
    histogram=result['histograms'].query('wave == "delta"')
    var=next(v for v in variables('student') if v.key=='PHQ_total')
    figure=distribution_chart(result['distribution_summary'],var,histogram=histogram)
    figure.add_vline(x=0)
    assert len(figure.data)==1 and sum(figure.data[0].y)==2
    assert 'T1 減 T0' in figure.layout.xaxis.title.text
    assert len(figure.layout.shapes)==1 and figure.layout.shapes[0].x0==0


def test_aggregate_workbook_contains_bases_and_no_person_rows():
    result=explore(panel([1,2],[2,3]),'student',paired=True,selected_variables=['PHQ_total'],metadata={'schools':['Example']})
    data=filtered_workbook(None,pd.DataFrame(),result['metadata'],exploration=result)
    book=load_workbook(io.BytesIO(data),data_only=False)
    assert set(book.sheetnames)=={'分布摘要','共同分箱','完整配對摘要','篩選及版本'}
    for sheet in book:
        assert sheet.freeze_panes=='C2' and sheet.auto_filter.ref
        assert not any(c.data_type=='f' for row in sheet for c in row)
        assert not any(c.value in {'child_id','submission_id','contact','dob'} for c in sheet[1])
    values=dict(list(book['篩選及版本'].values)[1:])
    assert json.loads(values['basis'])=='paired_valid'


def test_single_variable_statistics_download_preserves_filters_and_validity_basis():
    from contextlib import nullcontext
    from study2.descriptive_workbench import downloads
    from study2.descriptives import statistics
    class DownloadSurface:
        def __init__(self):self.files={}
        def expander(self,*args):return nullcontext()
        def dataframe(self,*args,**kwargs):pass
        def download_button(self,label,data,filename,*args,**kwargs):self.files[label]=data
    rows=panel([1,2],[2,3])
    result=explore(rows,'student',selected_variables=['PHQ_total'],metadata={'schools':['Example'],'mainstream':False})
    raw=statistics(rows,'student',('all',),'test','scales').query('variable_id == "PHQ_total"')
    surface=DownloadSurface();downloads(surface,result,'detail',stats=raw)
    csv=pd.read_csv(io.BytesIO(surface.files['下載每輪各自有效統計 CSV']))
    assert csv['basis'].eq('wave_valid').all()
    assert all(json.loads(value)==result['metadata'] for value in csv['filters_json'])
    assert 'basis' not in raw and 'filters_json' not in raw
    book=load_workbook(io.BytesIO(surface.files['下載目前分析 Excel']),data_only=True)
    columns=[c.value for c in book['統計明細'][1]]
    assert 'basis' in columns and 'filters_json' not in columns


def test_release_exploration_gates_and_old_capability(inputs,tmp_path):
    f,p,m,_=inputs
    folder=tmp_path/'synthetic';folder.mkdir()
    assert not available(folder,{'release_id':'synthetic'})
    result=write_exploration(folder,m,f,p,'synthetic')
    assert available(folder,{'descriptive':{'exploration':result}})
    assert len(result['coverage'])==12
    artifact=folder/result['files']['pair_summary']['file']
    artifact.unlink()
    with pytest.raises(FileNotFoundError):verify_exploration(folder,result)


def test_workbench_controls_all_types_and_role_protection(inputs):
    from streamlit.testing.v1 import AppTest
    f,p,m,_=inputs
    rows=cohort(f,p,m,'efmatched','parent')
    # The published sample is fixed. UI controls can change only the requested
    # aggregate, and may never expose adult PHQ on the family-child branch.
    app=AppTest.from_string(
        'import io\nimport pandas as pd\nfrom study2.descriptive_workbench import render_workbench\n'
        +'rows=pd.read_json(io.StringIO('+repr(rows.to_json(orient='table'))+'),orient="table")\n'
        +'render_workbench(rows,"parent","synthetic",prefix="test")',default_timeout=45).run()
    assert not app.exception
    assert not any('PHQ' in label for label in app.selectbox(key='test_variable_efmatched_parent').options)
    assert app.selectbox(key='test_variable_efmatched_parent').value=='FQOL_total'
    assert app.radio(key='test_view').value=='逐人變化'
    assert all(r.key!='test_distribution' for r in app.radio)
    app.radio(key='test_view').set_value('T0／T1 分布').run()
    app.radio(key='test_distribution').set_value('直方圖').run()
    assert not app.exception
    figure=json.loads(app.get('plotly_chart')[0].proto.spec)
    assert [t['xaxis'] for t in figure['data']]==['x','x2']
    assert any('T0在上、T1在下' in c.value for c in app.caption)
    app.radio(key='test_view').set_value('逐人變化').run()
    assert not app.exception
    app.selectbox(key='test_variable_efmatched_parent').set_value('TBV_prevalence2').run()
    assert not app.exception
    app.radio(key='test_view').set_value('子群比較').run()
    app.selectbox(key='test_groups').set_value('nine').run()
    assert not app.exception
    app.radio(key='test_view').set_value('T0／T1 分布').run()
    assert not app.exception
    assert all(r.key!='test_distribution' for r in app.radio)


def test_changed_pair_options_have_distinct_widget_identity_and_keep_valid_selection(inputs):
    from streamlit.testing.v1 import AppTest
    f,p,m,_=inputs;rows=cohort(f,p,m,'acbd','parent')
    app=AppTest.from_string('import io\nimport pandas as pd\nimport streamlit as st\n'
        'from study2.descriptive_workbench import render_workbench\n'
        +'rows=pd.read_json(io.StringIO('+repr(rows.to_json(orient='table'))+'),orient="table")\n'
        +'mode=st.selectbox("Mode",["acbd","efmatched"],key="mode")\n'
        +'rows["report_mode"]=mode\nrender_workbench(rows,"parent","synthetic",prefix="context")',default_timeout=45).run()
    assert not app.exception and app.selectbox(key='context_variable_acbd_parent').value=='PHQ_total'
    app.selectbox(key='mode').set_value('efmatched').run()
    assert not app.exception and app.selectbox(key='context_variable_efmatched_parent').value=='FQOL_total'
    app.selectbox(key='context_variable_efmatched_parent').set_value('SEC_total').run()
    app.selectbox(key='mode').set_value('acbd').run()
    assert not app.exception and app.selectbox(key='context_variable_acbd_parent').value=='SEC_total'
    assert not app.warning


def test_empty_filtered_overview_and_workbench_remain_readable(inputs):
    from streamlit.testing.v1 import AppTest
    f,p,m,_=inputs;rows=cohort(f,p,m,'acbd','student').iloc[:0].copy()
    rows.attrs['mode']='acbd'
    app=AppTest.from_string('import io\nimport pandas as pd\n'
        'from study2.descriptive_workbench import render_workbench,render_overview\n'
        +'rows=pd.read_json(io.StringIO('+repr(rows.to_json(orient='table'))+'),orient="table")\n'
        +'rows.attrs["mode"]="acbd"\nrender_overview(rows,"student","empty",prefix="overview")\n'
        +'render_workbench(rows,"student","empty",prefix="paired")',default_timeout=45).run()
    assert not app.exception and app.info
    app.radio(key='paired_view').set_value('逐人變化').run()
    assert not app.exception


def test_exploration_failure_preserves_publication_and_missing_file_blocks_publish(project):
    from pathlib import Path
    from unittest.mock import patch
    from test_pipeline_integration import execute
    from study2.store import connect,current
    from study2.runner import verify_release
    first=execute(project)
    cfg,_,_,_=project
    with patch('study2.descriptive_exploration_outputs.write_exploration',side_effect=RuntimeError('exploration interrupted')):
        with pytest.raises(RuntimeError,match='interrupted'):execute(project,force=True)
    db=connect(Path(cfg['state_dir']),readonly=True)
    assert current(db)['release_id']==first['release_id']
    folder=Path(cfg['release_dir'])/first['release_id']
    (folder/'descriptive_pair_summary.parquet').unlink()
    with pytest.raises((PipelineError,FileNotFoundError)):verify_release(db,first['release_id'])
    db.close()


def test_parent_fallback_cannot_authorize_phq_changes_or_fixed_workbook(inputs):
    from study2.descriptives import paired_changes, statistics
    from study2.descriptive_reports import workbook_bytes
    frame, parents, measurements, _ = inputs
    frame.loc[frame['wave'].eq('T1'), 'PHQ_total'] = 12
    rows = cohort(frame, None, measurements, 'acbd', 'parent')
    assert not rows['adult_continuity_confirmed'].any()
    for key in ['PHQ_total', 'PHQ_anxiety', 'PHQ_depression', 'phq_band']:
        with pytest.raises(PipelineError, match='unavailable'):
            explore(rows, 'parent', paired=True, selected_variables=[key])
    changes = paired_changes(rows, 'parent')
    assert not changes['variable_id'].str.startswith('PHQ_').any()
    assert 'FQOL_total' in set(changes['variable_id'])
    stats = statistics(rows, 'parent')
    assert 'PHQ_total' in set(stats['variable_id'])
    book = load_workbook(io.BytesIO(workbook_bytes(stats, changes, 'acbd', 'test')))
    # Both summary and export use the guarded shared changes table.
    change_sheet = book['配對差'] if '配對差' in book.sheetnames else None
    assert change_sheet is not None
    assert not any(str(c.value).startswith('PHQ_') for row in change_sheet for c in row)
    confirmed = cohort(frame, parents, measurements, 'acbd', 'parent')
    assert any(v.key == 'PHQ_total' for v in comparison_variables('acbd', 'parent', confirmed))
    confirmed['adult_continuity_confirmed'] = confirmed['adult_continuity_confirmed'].astype('boolean')
    confirmed.loc[confirmed.index[0], 'adult_continuity_confirmed'] = pd.NA
    assert not any(v.key == 'PHQ_total' for v in comparison_variables('acbd', 'parent', confirmed))
