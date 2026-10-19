"""Synthetic measurements and complete comparison units, no research data."""
import copy
import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from study2.common import PipelineError
from study2.descriptive_data import demographic_values, report_measurement, school_stage, sen
from study2.descriptive_dashboard import filtered_workbook
from study2.descriptive_reports import workbook_bytes
from study2.descriptive_spec import GROUPS, MODES, REPORT_RANGES, variables
from study2.descriptives import (cohort, filter_cohort, independent_cohort, paired_changes,
                                 statistics, validate_statistics, wave_counts)
from study2.outputs import metrics


def observation(sid, child='one', role='student', wave='T0', **overrides):
    answer = {'school_type':'primary', 'Survey':'example_qualtrics', 'adult_demographics_05':1980,
              'adult_demographics_06':1, 'adult_demographics_07':0, 'adult_demographics_08':3,
              'adult_demographics_09':5, 'adult_demographics_10':0, 'adult_demographics_11':0,
              'adult_demographics_12':4, 'adult_demographics_13':0, 'adult_demographics_14':'0,1',
              'student_demographics_04':150, 'student_demographics_05':45,
              'student_demographics_07':0, 'student_demographics_10':2, 'student_demographics_11':4,
              **{'student_demographics_09.'+str(i):0 for i in range(1,5)}}
    for domain, length in [('PHQ',4), ('SEC',22), ('FQOL',21), ('TBV',7), ('TBP',7), ('CBV',6), ('CBP',6)]:
        answer.update({f'{role}_{domain}_{i:02d}':1 if domain=='FQOL' else 0 for i in range(1,length+1)})
    result = dict(submission_id=sid, child_id=child, role=role, selected_wave=wave, version_id='version-'+sid,
                  observed_at='2025-09-01T00:00:00+00:00' if wave=='T0' else '2026-09-01T00:00:00+00:00',
                  time_confidence='observed', source_label='example_qualtrics', gender='0', grade=6,
                  grade_raw='P6', school='Example Primary', answer=answer)
    result.update(overrides)
    return result


@pytest.fixture
def inputs():
    records=[]
    for child, gender in [('one','0'), ('two','1')]:
        for role in ['student','parent']:
            for wave in ['T0','T1']:
                r=observation(child+'-'+role+'-'+wave, child, role, wave, gender=gender)
                if wave=='T1':
                    r['answer']['school_type']='secondary';r.update(grade=1,grade_raw='S1',school='Example Secondary')
                records.append(r)
    base=[];measurements=[]
    for r in records:
        row=dict(submission_id=r['submission_id'], child_id=r['child_id'], role=r['role'], wave=r['selected_wave'],
                 school=r['school'],school_type=r['answer']['school_type'],gender=r['gender'],grade=r['grade'],
                 source_type='qualtrics',mainstream=True,identity_rules='exact;transfer',
                 longitudinal=True,dyad=True,longitudinal_dyad=True,age_at_observation=11.5,
                 **metrics(r['answer'],r['role']))
        base.append(row);measurements.append(demographic_values(r,row))
    frame=pd.DataFrame(base)
    parents=frame[frame['role'].eq('parent')].copy()
    return frame,parents,pd.DataFrame(measurements),records


def test_question_meanings_relationship_not_gender_and_topcoded_children():
    r=observation('p',role='parent')
    value=demographic_values(r, {'age_at_observation':10})
    assert value['dem_adult_gender'] is None and value['dem_adult_gender__status']=='not_collected'
    assert value['dem_relationship']=='1'
    assert value['dem_cohabit']=='0'
    assert value['dem_children_count']=='4'
    assert next(v for v in variables('parent') if v.key=='dem_children_count').kind=='category'
    r['source_label']='legacy_example';r['answer']['adult_demographics_12']=0
    other=demographic_values(r, {'age_at_observation':10})
    assert other['dem_children_count'] is None and other['dem_children_count__status']=='unverified_encoding'


@pytest.mark.parametrize('raw,expected,status', [('0,1','0;1','valid'), ('3','3','valid'),
    ('3,4',None,'conflict'), ('',None,'missing'), ('nonsense',None,'invalid'), ('7',None,'invalid'),
    ('0.0,2.0','0;2','valid')])
def test_sen_answer_is_not_inferred_from_missing_or_conflicting_codes(raw,expected,status):
    assert sen(raw)==(expected,status)


def test_ages_fixed_to_observation_and_unknown_year_preserved(monkeypatch):
    monkeypatch.setitem(REPORT_RANGES, 'child_age', (5,20))
    monkeypatch.setitem(REPORT_RANGES, 'adult_age', (20,70))
    r=observation('p',role='parent')
    a=demographic_values(r,{'age_at_observation':10})
    assert a['dem_adult_age']==45 and a['dem_child_age']==10
    r['adult_birth_year_status']='unconfirmed'
    assert demographic_values(r,{'age_at_observation':10})['dem_adult_age__status']=='unconfirmed_year'
    r.pop('adult_birth_year_status');r['time_confidence']='accepted_creation_time'
    assert demographic_values(r,{'age_at_observation':None})['dem_adult_age__status']=='unverified_date'


def test_siblings_missing_not_zero_and_bmi_units(monkeypatch):
    monkeypatch.setitem(REPORT_RANGES, 'height', (100,200))
    monkeypatch.setitem(REPORT_RANGES, 'weight', (20,100))
    for key in ('older_brothers','older_sisters','younger_brothers','younger_sisters'):
        monkeypatch.setitem(REPORT_RANGES,key,(0,5))
    monkeypatch.setitem(REPORT_RANGES,'siblings_total',(0,10))
    r=observation('s');before=copy.deepcopy(r)
    values=demographic_values(r,{'age_at_observation':11})
    assert values['dem_siblings_total']==0
    assert values['dem_bmi']==20
    assert r==before
    r['answer']['student_demographics_09.1']=None
    assert demographic_values(r,{'age_at_observation':11})['dem_siblings_total'] is None
    r['answer']['student_demographics_04']=0
    assert demographic_values(r,{'age_at_observation':11})['dem_bmi'] is None


def test_unapproved_demographic_ranges_pause_only_affected_measures(monkeypatch):
    for key in REPORT_RANGES:monkeypatch.setitem(REPORT_RANGES,key,None)
    r=observation('s');before=copy.deepcopy(r)
    values=demographic_values(r,{'age_at_observation':11})
    for key in ('child_age','height','weight','bmi','older_brothers','older_sisters',
                'younger_brothers','younger_sisters','siblings_total'):
        assert values['dem_'+key] is None
        assert values['dem_'+key+'__status']=='range_pending'
    assert values['dem_birthplace']=='0'
    assert r==before


@pytest.mark.parametrize('key,low,high',[('child_age',5,20),('adult_age',20,70),('height',100,200),
    ('weight',20,100),('older_brothers',0,5),('older_sisters',0,5),('younger_brothers',0,5),
    ('younger_sisters',0,5),('siblings_total',0,10)])
def test_approved_report_ranges_are_inclusive_and_preserve_missing_reasons(key,low,high):
    assert REPORT_RANGES[key]==(low,high)
    assert report_measurement(key,(low,'valid'))==(low,'valid')
    assert report_measurement(key,(high,'valid'))==(high,'valid')
    assert report_measurement(key,(high+1,'valid'))==(None,'outside_report_range')
    assert report_measurement(key,(low-1,'valid'))==(None,'outside_report_range')
    assert report_measurement(key,(None,'invalid'))==(None,'invalid')
    assert report_measurement(key,(None,'missing'))==(None,'missing')


def test_outlier_is_missing_for_that_measure_only_and_is_counted(inputs):
    f,p,m,records=inputs
    r=next(r for r in records if r['submission_id']=='one-student-T0')
    r['answer']['student_demographics_04']=10000
    replacement=demographic_values(r,f[f['submission_id'].eq(r['submission_id'])].iloc[0].to_dict())
    m=pd.concat([m[~m['submission_id'].eq(r['submission_id'])],pd.DataFrame([replacement])],ignore_index=True)
    rows=cohort(f,p,m,'abcd','student')
    table=statistics(rows,'student',section='student')
    height=table.query("wave == 'T0' and variable_id == 'dem_height'").iloc[0]
    bmi=table.query("wave == 'T0' and variable_id == 'dem_bmi'").iloc[0]
    assert (height['sample_n'],height['valid_n'],height['missing_n'],height['outside_range_n'])==(2,1,1,1)
    assert height['mean']==150 and bmi['outside_range_n']==1
    assert len(rows)==4 and r['answer']['student_demographics_04']==10000
    validate_statistics(table)


def test_sibling_total_requires_valid_components_and_own_range():
    r=observation('s')
    for i in range(1,5):r['answer'][f'student_demographics_09.{i}']=5
    values=demographic_values(r,{'age_at_observation':10})
    assert values['dem_older_brothers']==5
    assert values['dem_siblings_total'] is None and values['dem_siblings_total__status']=='outside_report_range'
    r['answer']['student_demographics_09.1']=6
    values=demographic_values(r,{'age_at_observation':10})
    assert values['dem_older_brothers'] is None and values['dem_siblings_total'] is None


def test_through_train_stage_needs_existing_stage_evidence():
    r=observation('s');r['answer']['school_type']='through-train school'
    assert school_stage(r)==('primary','valid')
    r['grade_raw']='1'
    assert school_stage(r)==(None,'unknown_stage')
    r['grade_raw']='F.1'
    assert school_stage(r)==('secondary','valid')
    r['answer']['school_type']='primary'
    assert school_stage(r)==(None,'conflict')


@pytest.mark.parametrize('mode',['acbd','efmatched'])
@pytest.mark.parametrize('role',['student','parent'])
def test_longitudinal_filters_hold_complete_units_at_baseline(inputs,mode,role):
    f,p,m,_=inputs
    rows=cohort(f,p,m,mode,role)
    selected,counts=filter_cohort(rows,group='male_primary',schools=['Example Primary'])
    assert wave_counts(selected)=={'T0':1,'T1':1}
    assert set(selected['school'])=={'Example Primary','Example Secondary'}
    assert selected['anchor_school_stage'].eq('primary').all()
    assert counts.iloc[-1]['T1']==1


def test_cross_sectional_and_dyad_stage_use_that_wave_and_student_anchor(inputs):
    f,p,m,_=inputs
    rows=cohort(f,p,m,'abcd','student')
    assert wave_counts(filter_cohort(rows,group='primary')[0])=={'T0':2,'T1':0}
    # A deliberately different parent field cannot move a dyad's grouping anchor.
    f.loc[f['role'].eq('parent'),'gender']='1'
    parents=cohort(f,p,m,'ef','parent')
    male,_=filter_cohort(parents,group='male')
    assert wave_counts(male)=={'T0':1,'T1':1}


def test_bd_uses_own_selected_followup_not_latest_family_survey(inputs):
    f,p,m,_=inputs
    p.loc[p['wave'].eq('T1'),'PHQ_total']=7
    assert cohort(f,p,m,'acbd','parent').query("wave == 'T1'")['PHQ_total'].eq(7).all()
    assert cohort(f,p,m,'efmatched','parent').query("wave == 'T1'")['PHQ_total'].eq(0).all()


def test_denominators_unique_categories_multiselect_and_zero(inputs):
    f,p,m,_=inputs
    rows=cohort(f,p,m,'abcd','parent')
    stats=statistics(rows,'parent',GROUPS,'synthetic-release')
    validate_statistics(stats)
    edu=stats.query("group == 'all' and wave == 'T0' and variable_id == 'dem_education'")
    assert edu['numerator'].sum()==edu['valid_n'].iloc[0]==2
    multi=stats.query("group == 'all' and wave == 'T0' and variable_id == 'dem_sen_types'")
    assert multi['proportion'].sum()==2 and multi['denominator'].eq(2).all()
    phq=stats.query("group == 'all' and wave == 'T0' and variable_id == 'PHQ_total'").iloc[0]
    assert phq['mean']==0 and phq['valid_n']==2
    gender=stats.query("group == 'all' and variable_id == 'dem_adult_gender'")
    assert gender['valid_n'].eq(0).all() and gender['missing_n'].eq(2).all()
    assert all(json.loads(s)=={'not_collected':2} for s in gender['missing_detail'])


def test_legacy_statistics_without_outlier_column_remain_verifiable(inputs):
    f,p,m,_=inputs
    previous=statistics(cohort(f,p,m,'abcd','parent'),'parent').drop(columns='outside_range_n')
    validate_statistics(previous)


def test_missing_scores_and_complete_pair_difference(inputs):
    f,p,m,_=inputs
    f.loc[(f['child_id']=='one') & f['wave'].eq('T0'),'PHQ_total']=None
    f.loc[(f['child_id']=='one') & f['wave'].eq('T1'),'PHQ_total']=10
    f.loc[(f['child_id']=='two') & f['wave'].eq('T0'),'PHQ_total']=2
    f.loc[(f['child_id']=='two') & f['wave'].eq('T1'),'PHQ_total']=3
    rows=cohort(f,p,m,'acbd','student')
    delta=paired_changes(rows,'student').query("variable_id == 'PHQ_total'").iloc[0]
    assert delta['paired_valid_n']==1 and delta['paired_missing_n']==1 and delta['mean_change']==1
    assert pd.isna(delta['sd_change'])
    ef=cohort(f,p,m,'efmatched','parent')
    assert not paired_changes(ef,'parent')['variable_id'].str.startswith('PHQ_').any()
    assert paired_changes(cohort(f,p,m,'abcd','student'),'student').empty


def test_empty_group_and_one_observation_have_correct_missing_and_sd(inputs):
    f,p,m,_=inputs
    rows=cohort(f,p,m,'abcd','student')
    empty=filter_cohort(rows,schools=['Absent School'])[0]
    stats=statistics(empty,'student',GROUPS)
    validate_statistics(stats)
    assert stats['sample_n'].eq(0).all() and stats['mean'].isna().all()
    one=rows[rows['child_id'].eq('one')]
    assert statistics(one,'student').query("variable_id == 'PHQ_total'")['sd'].isna().all()


def test_independent_branch_and_baseline_filtering_stay_separate(inputs):
    f,p,m,_=inputs
    source=p.copy();source['panel_branch']='same_adult';source['panel_wave']=source['wave'];source['panel_longitudinal']=True
    rows=independent_cohort(source,m,'same_adult')
    selected,_=filter_cohort(rows,schools=['Example Primary'],group='male_primary')
    assert wave_counts(selected)=={'T0':1,'T1':1}
    assert selected['report_mode'].eq('independent_same_adult').all()
    assert paired_changes(selected,'parent')['variable_id'].str.startswith('PHQ_').any()


def test_shuffle_does_not_change_numeric_results(inputs):
    f,p,m,_=inputs
    before=statistics(cohort(f,p,m,'efmatched','parent'),'parent',GROUPS)
    after=statistics(cohort(f.sample(frac=1,random_state=5),p.sample(frac=1),m.sample(frac=1),'efmatched','parent'),'parent',GROUPS)
    pd.testing.assert_frame_equal(before,after)


def test_numeric_excel_and_matrix_are_projections_of_same_statistics(inputs):
    f,p,m,_=inputs
    stats=pd.concat([statistics(cohort(f,p,m,'abcd',role),role,GROUPS,'synthetic-release') for role in ['student','parent']],ignore_index=True)
    book=load_workbook(io.BytesIO(workbook_bytes(stats,pd.DataFrame(),'abcd','synthetic-release')),data_only=False)
    assert set(book.sheetnames)=={'Scales','student','parent','統計明細','口徑及版本'}
    assert book['Scales'].freeze_panes=='B7' and book['Scales'].max_column==37
    detail=list(book['統計明細'].values);headers=detail[0]
    assert len(detail)-1==len(stats)
    assert isinstance(detail[1][headers.index('valid_n')],int)
    assert isinstance(detail[1][headers.index('mean')],(float,int))
    assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
    custom=load_workbook(io.BytesIO(filtered_workbook(stats,pd.DataFrame(),{'release_id':'synthetic-release'})))
    assert set(custom.sheetnames)=={'目前篩選','統計明細','篩選及版本'}


def test_duplicate_measurement_or_slot_cannot_inflate_reports(inputs):
    f,p,m,_=inputs
    with pytest.raises(ValueError,match='unique'):
        cohort(f,p,pd.concat([m,m.iloc[:1]]),'abcd','student')
    with pytest.raises(PipelineError,match='multiple'):
        cohort(pd.concat([f,f.iloc[:1]]),p,m,'abcd','student')
