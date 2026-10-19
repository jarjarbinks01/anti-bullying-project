"""Hand-calculated calendar boundaries and denominators, using synthetic rows."""
from datetime import date
import io
import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from study2.common import PipelineError
from study2.descriptive_spec import variables
from study2.descriptive_trends import (TABLE_NAMES, calendar_trends, date_diagnostics,
    date_extent, dated_rows, download_table, periods, workbook_bytes)
from study2.descriptive_trend_charts import indicator_chart, response_chart
from test_pipeline_integration import execute, project


def response(sid, observed='2025-09-03T10:00:00Z', wave='T0', **extra):
    result = dict(submission_id=sid, version_id='version-'+sid, child_id=sid, wave=wave,
        observed_at=observed, time_confidence='observed', date_basis='system_creation',
        report_mode='abcd', grouping_anchor='該輪回應', role='student',
        school='Example Primary', source_type='qualtrics', school_stage='primary',
        anchor_gender='0', anchor_school_stage='primary', dem_child_gender='0',
        PHQ_total=3)
    result.update(extra)
    return result


def trend(rows, **kwargs):
    options = dict(start=date(2025, 9, 3), end=date(2026, 3, 30), variable_ids=['PHQ_total'])
    options.update(kwargs)
    return calendar_trends(pd.DataFrame(rows), 'student', 'synthetic-release', **options)


def test_hk_inclusive_boundaries_seven_months_no_wave_reselection_or_mutation():
    frame = pd.DataFrame([
        response('before', '2025-09-02T15:59:59.999999Z', child_id='same'),
        response('start', '2025-09-02T16:00:00Z'),
        response('end', '2026-03-30T15:59:59.999999Z', 'T1', child_id='same'),
        response('after', '2026-03-30T16:00:00Z', 'T1'),
    ])
    original = frame.copy(deep=True)
    result = calendar_trends(frame, 'student', 'r', start=date(2025,9,3), end=date(2026,3,30), variable_ids=['PHQ_total'])
    counts = result['counts']
    assert counts['period_id'].nunique() == 7
    assert counts.query('sample_n == 1')[['period_id','wave']].values.tolist() == [['2025-09','T0'], ['2026-03','T1']]
    assert counts['covered_start'].iloc[0] == '2025-09-03'
    assert counts['covered_end'].iloc[-1] == '2026-03-30'
    assert counts.groupby('period_id')['partial_period'].first().tolist() == [True,False,False,False,False,False,True]
    assert result['exclusions']['excluded_n'].sum() == 2
    pd.testing.assert_frame_equal(frame, original)


def test_week_monday_sunday_cross_year_and_leap_day():
    bins = periods(date(2024,12,29), date(2025,1,6), 'week')
    assert [b['period_id'] for b in bins] == ['2024-12-23','2024-12-30','2025-01-06']
    assert [(b['covered_start'],b['covered_end']) for b in bins] == [
        ('2024-12-29','2024-12-29'), ('2024-12-30','2025-01-05'), ('2025-01-06','2025-01-06')]
    result = trend([response('sun','2024-12-29T15:59:59Z'), response('mon','2024-12-29T16:00:00Z')],
        start=date(2024,12,29), end=date(2025,1,6), interval='week')
    assert result['counts'].query("wave == 'T0'")['sample_n'].tolist() == [1,1,0]
    leap = trend([response('leap','2024-02-29T15:59:59Z')], start=date(2024,2,29), end=date(2024,2,29))
    assert leap['counts']['sample_n'].sum() == 1


def test_date_quality_manual_precedence_batch_toggle_unknown_and_numeric_not_guessed():
    rows = [response('system'), response('batch',time_confidence='accepted_creation_time',date_basis='batch_assigned_creation'),
        response('manual',time_confidence='manual_verified',date_basis='batch_assigned_creation'),
        response('unknown',time_confidence='unknown'), response('bad','not a date'),
        response('missing',None), response('number',1700000000),
        response('mismatch',date_basis='batch_assigned_creation')]
    strict = trend(rows); inclusive = trend(rows, include_batch=True)
    assert strict['counts']['sample_n'].sum() == 2
    assert inclusive['counts']['sample_n'].sum() == 3
    assert inclusive['counts']['batch_date_n'].sum() == 1
    assert strict['counts']['manual_date_n'].sum() == 1
    assert strict['exclusions'].query("wave == 'T0'")['excluded_n'].tolist() == [0,0,5,1,0]
    quality = inclusive['date_quality'].groupby('date_quality')['input_n'].sum().to_dict()
    assert quality == {'system_creation':1,'manual_verified':1,'batch_assigned_creation':1,'unknown_basis':2,'invalid_date':3}
    assert date_extent(pd.DataFrame(rows)) == (date(2025,9,3), date(2025,9,3))
    assert date_extent(pd.DataFrame([rows[1]])) is None
    assert date_diagnostics(pd.DataFrame(rows))['input_n'].sum() == len(rows)


def test_per_variable_missingness_sd_and_whole_period_not_mean_of_months():
    result = trend([
        response('a',PHQ_total=1,FQOL_total=None), response('b',PHQ_total=3,FQOL_total=4),
        response('c','2025-10-01T01:00Z',PHQ_total=9,FQOL_total=5),
        response('d','2025-10-02T01:00Z',PHQ_total=None,FQOL_total=3)],
        end=date(2025,11,30), variable_ids=['PHQ_total','FQOL_total','phq_band','TBV_prevalence2'])
    stats = result['statistics']
    assert (stats['valid_n']+stats['missing_n']).equals(stats['sample_n'])
    phq = stats.query("variable_id == 'PHQ_total' and wave == 'T0'")
    assert phq['mean'].iloc[0] == 2 and phq['sd'].iloc[0] == pytest.approx(np.sqrt(2))
    assert phq['valid_n'].tolist() == [2,1,0]
    assert phq['missing_n'].tolist() == [0,1,0]
    assert pd.isna(phq['sd'].iloc[1]) and pd.isna(phq['mean'].iloc[2])
    all_phq = result['whole_period'].query("variable_id == 'PHQ_total' and wave == 'T0'").iloc[0]
    assert all_phq['mean'] == pytest.approx(13/3)
    assert stats.query("variable_id == 'FQOL_total' and wave == 'T0'")['valid_n'].tolist() == [1,2,0]
    bands = stats.query("variable_id == 'phq_band' and wave == 'T0' and period_id == '2025-09'")
    assert bands['mean'].isna().all() and bands['proportion'].tolist() == [.5,.5,0,0]
    assert all(json.loads(r.missing_detail) == {'missing':r.missing_n}
               for r in stats.query("variable_id == 'TBV_prevalence2' and sample_n > 0").itertuples())


def test_binary_numerator_denominator_and_composition_include_unknown():
    rows = [response('a',TBV_prevalence2=1), response('b',TBV_prevalence2=0,school=None,school_stage=None,dem_child_gender=None),
            response('c',TBV_prevalence2=None), response('d',TBV_prevalence2=5)]
    result = trend(rows, end=date(2025,9,30), variable_ids=['TBV_prevalence2'])
    r = result['statistics'].query("wave == 'T0'").iloc[0]
    assert (r.sample_n,r.valid_n,r.missing_n,r.numerator,r.denominator,r.proportion) == (4,2,2,1,2,.5)
    for dimension, block in result['composition'].groupby('dimension'):
        assert block['count'].sum() == 4 and block['denominator'].eq(4).all()
        assert block['proportion'].sum() == 1
    assert '未知／缺失' in set(result['composition']['category'])


@pytest.mark.parametrize('fault',['missing_column','duplicate','missing_version','unknown_wave','wrong_role','wrong_dictionary'])
def test_bad_or_unsupported_release_fails_closed(fault):
    rows = pd.DataFrame([response('a')])
    if fault == 'missing_column': rows = rows.drop(columns='observed_at')
    if fault == 'duplicate': rows = pd.concat([rows,rows])
    if fault == 'missing_version': rows['version_id'] = None
    if fault == 'unknown_wave': rows['wave'] = 'T2'
    if fault == 'wrong_role': rows['role'] = 'parent'
    if fault == 'wrong_dictionary': rows['descriptive_dictionary_version'] = 'old'
    with pytest.raises(PipelineError):
        calendar_trends(rows,'student','r',start=date(2025,9,1),end=date(2025,9,30))


def test_empty_cohort_empty_range_and_wave_selection_keep_schema():
    frame = pd.DataFrame([response('a')]).iloc[:0]
    assert date_extent(frame) is None
    r = calendar_trends(frame,'student','r',start=date(2025,9,1),end=date(2025,10,1),variable_ids=['PHQ_total'])
    assert len(r['counts']) == 4 and r['counts']['sample_n'].sum() == 0
    assert r['statistics']['mean'].isna().all() and not r['exclusions']['excluded_n'].any()
    r = trend([response('a'),response('b',wave='T1')],waves=['T1'])
    assert set(r['statistics']['wave']) == {'T1'} and r['counts']['sample_n'].sum() == 1
    assert r['exclusions']['excluded_n'].sum() == 1
    outside = trend([response('a','2024-01-01T00:00Z')])
    assert outside['counts']['sample_n'].sum() == 0
    assert outside['statistics']['mean'].isna().all()
    with pytest.raises(ValueError): periods(date(2025,2,1),date(2025,1,1))
    with pytest.raises(ValueError): trend([response('a')], waves=['T2'])


def test_charts_and_csv_excel_use_the_same_aggregates_no_person_rows_or_formulas():
    rows = [response('a',PHQ_total=1,contact='private contact',dob='private birthday',school='=1+1'),
            response('b','2025-11-01T00:00Z','T1',PHQ_total=4)]
    result = trend(rows,end=date(2025,11,30),variable_ids=['PHQ_total','phq_band'])
    var = next(v for v in variables('student') if v.key == 'PHQ_total')
    chart = indicator_chart(result['statistics'],var)
    assert len(chart.data) == 2 and all(not t.connectgaps for t in chart.data)
    assert chart.layout.xaxis.type == 'category'
    assert response_chart(result['counts']).layout.xaxis.type == 'category'
    assert list(chart.data[0].x) == ['2025-09＊','2025-10','2025-11']
    assert chart.data[0].y[0] == 1 and pd.isna(chart.data[0].y[1])
    assert response_chart(result['counts']).data[0].y.tolist() == [1,0,0]
    before = result['statistics'].copy(deep=True)
    csv = download_table(result,'statistics')
    pd.testing.assert_frame_equal(before,result['statistics'])
    assert json.loads(csv['settings_json'].iloc[0]) == result['metadata']
    book = load_workbook(io.BytesIO(workbook_bytes(result)),data_only=False)
    assert set(book.sheetnames) == {*TABLE_NAMES.values(),'設定及版本'}
    for key, title in TABLE_NAMES.items():
        actual = pd.DataFrame(book[title].values)
        actual.columns = actual.iloc[0];actual = actual.iloc[1:]
        expected = result[key]
        assert len(actual) == len(expected)
        for column in ['mean','sd','sample_n','valid_n','missing_n','numerator','denominator',
                       'proportion','remaining_n','excluded_n','count']:
            if column in expected:
                np.testing.assert_allclose(pd.to_numeric(actual[column]),pd.to_numeric(expected[column]),
                                           equal_nan=True,rtol=1e-14)
        assert not {'submission_id','child_id','contact','dob','version_id'} & set(actual.columns)
        for row in book[title]:
            assert not any(c.data_type == 'f' for c in row)
            assert not any(c.value in {'private contact','private birthday'} for c in row)
    settings = {k:json.loads(v) for k,v in list(book['設定及版本'].values)[1:]}
    assert settings == result['metadata'] and settings['basis'] == 'calendar_response'
    rendered = workbook_bytes(result)
    assert b'private contact' not in rendered


def test_cache_tracks_dates_policy_waves_filters_release_and_versions(monkeypatch):
    import study2.descriptive_trend_dashboard as ui
    import streamlit as st
    st.cache_data.clear()
    calls = []
    original = ui.calendar_trends
    def spy(*args,**kwargs):
        calls.append(kwargs)
        return original(*args,**kwargs)
    monkeypatch.setattr(ui,'calendar_trends',spy)
    rows = pd.DataFrame([response('a')])
    options = dict(start=date(2025,9,1),end=date(2025,9,30),variable_ids=['PHQ_total'])
    ui.cached_trends(rows,'student','r',**options); ui.cached_trends(rows,'student','r',**options)
    assert len(calls) == 1
    for change in [dict(end=date(2025,10,1)),dict(include_batch=True),dict(waves=['T0']),
                   dict(interval='week'),dict(metadata={'sources':['qualtrics']})]:
        ui.cached_trends(rows,'student','r',**(options|change))
    ui.cached_trends(rows,'student','another-release',**options)
    monkeypatch.setattr(ui,'VERSION','new-policy')
    ui.cached_trends(rows,'student','r',**options)
    monkeypatch.setattr(ui,'DICTIONARY_VERSION','new-dictionary-cache-key')
    ui.cached_trends(rows,'student','r',**options)
    assert len(calls) == 9
    st.cache_data.clear()


def test_shared_ui_controls_families_dates_empty_sample_and_batch_recovery():
    from streamlit.testing.v1 import AppTest
    rows = pd.DataFrame([response('a'),response('b','2026-03-30T00:00Z','T1'),
        response('batch','2025-08-01T00:00Z',time_confidence='accepted_creation_time',date_basis='batch_assigned_creation')])
    script = ('import io\nimport pandas as pd\nimport streamlit as st\n'
        'from study2.descriptive_trend_dashboard import render_trends\n'
        +'rows=pd.read_json(io.StringIO('+repr(rows.to_json(orient='table'))+'),orient="table")\n'
        'only_batch=st.checkbox("only batch",key="only_batch")\n'
        'if only_batch: rows=rows[rows.submission_id.eq("batch")]\n'
        'result=render_trends(rows,"student","r",prefix="trend",metadata={"mode":"abcd"})\n'
        'if result is not None: st.session_state["result"]=result\n')
    app = AppTest.from_string(script,default_timeout=45).run()
    assert not app.exception
    assert app.date_input(key='trend_dates').value == (date(2025,9,3),date(2026,3,30))
    assert [m.value for m in app.metric] == ['1','1']
    app.date_input(key='trend_dates').set_value((date(2025,10,1),date(2025,10,31))).run()
    assert not app.exception and [m.value for m in app.metric] == ['0','0']
    app.selectbox(key='trend_interval').set_value('week').run()
    app.selectbox(key='trend_family').set_value('FQOL').run()
    assert not app.exception and len(app.get('plotly_chart')) == 6
    app.selectbox(key='trend_family').set_value('bullying').run()
    assert any('回顧過去三個月' in c.value for c in app.caption)
    app.selectbox(key='trend_family').set_value('demographics').run()
    assert not app.exception
    app.selectbox(key='trend_waves').set_value('T1').run()
    assert app.metric[0].value == '未選取'
    app.selectbox(key='trend_waves').set_value('both').run()
    app.checkbox(key='only_batch').check().run()
    assert not app.exception and any('沒有可用填卷日期' in i.value for i in app.info)
    app.selectbox(key='trend_batch').set_value(True).run()
    assert not app.exception
    assert app.date_input(key='trend_dates').value == (date(2025,8,1),date(2025,8,1))
    assert app.metric[0].value == '1'


def test_unsupported_release_ui_explains_missing_fields_without_fallback():
    from streamlit.testing.v1 import AppTest
    rows = pd.DataFrame([response('a')]).drop(columns='observed_at')
    app = AppTest.from_string('import io\nimport pandas as pd\n'
        'from study2.descriptive_trend_dashboard import render_trends\n'
        +'rows=pd.read_json(io.StringIO('+repr(rows.to_json(orient='table'))+'),orient="table")\n'
        'render_trends(rows,"student","old",prefix="unsupported")',default_timeout=30).run()
    assert not app.exception
    assert any('observed_at' in i.value and '不讀取其他版本' in i.value for i in app.info)
    assert not app.date_input and not app.get('plotly_chart')


def test_both_page_entrypoints_keep_samples_filters_and_adult_guard(project):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    from study2.descriptive_exploration import comparison_variables
    from study2.descriptives import cohort
    cfg, _, records, _ = project
    for r in records:
        r['date_basis'] = 'system_creation'
        if r['role'] == 'parent':r['answer']['adult_demographics_05'] = 1980
    summary = execute(project)
    folder = Path(cfg['release_dir'])/summary['release_id']
    app = AppTest.from_string('from pathlib import Path\nimport pandas as pd\n'
        'from study2.descriptive_dashboard import render_descriptive\n'
        +'p=Path('+repr(str(folder))+')\n'
        +'render_descriptive(p,'+repr(summary)+',pd.read_parquet(p/"analysis.parquet"),'
        '(pd.read_parquet(p/"parent_longitudinal.parquet") if (p/"parent_longitudinal.parquet").exists() else None),{})',default_timeout=60).run()
    assert not app.exception
    app.multiselect(key='ds_sources').set_value(['legacy_online']).run()
    app.radio(key='ds_reading').set_value('時段趨勢').run()
    assert not app.exception
    assert app.metric[0].label == '日期篩選前 T0／T1 N'
    assert app.multiselect(key='ds_sources').value == ['legacy_online']
    assert [m.value for m in app.metric[1:]] == ['1','1']
    app.selectbox(key='ds_role').set_value('parent').run()
    assert not app.exception
    app.radio(key='ds_reading').set_value('研究總覽').run()
    assert not app.exception and app.metric[0].value == '1 / 1'
    parent = AppTest.from_string('from study2.parent_panel_dashboard import render_parent_panel\n'
        +'render_parent_panel('+repr(str(folder))+','+repr(summary)+')',default_timeout=60).run()
    assert not parent.exception
    parent.radio(key='ipp_reading').set_value('時段趨勢').run()
    assert not parent.exception and [m.value for m in parent.metric[-2:]] == ['1','1']
    parent.selectbox(key='ipp_wave').set_value('T1').run()
    assert not parent.exception and [m.value for m in parent.metric[-2:]] == ['1','1']
    parent.selectbox(key='ipp_branch').set_value('same_adult').run()
    assert not parent.exception
    # Calendar response PHQ does not add paired PHQ to the child-report branch.
    frame=pd.read_parquet(folder/'analysis.parquet');measurements=pd.read_parquet(folder/'descriptive_measurements.parquet')
    child_rows=cohort(frame,None,measurements,'efmatched','parent')
    assert not any(v.key.startswith('PHQ_') or v.key=='phq_band'
                   for v in comparison_variables('efmatched','parent',child_rows))
    assert pd.read_parquet(folder/'analysis.parquet').equals(frame)
