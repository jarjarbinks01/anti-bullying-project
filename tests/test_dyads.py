"""Synthetic acceptance tests for permanent IDs and administrative isolation."""
import copy
import json
import sqlite3
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest

from study2 import dyads, dyad_operator as operator
from study2.common import PipelineError, dumps
from study2.store import connect, current
from test_identity_and_time import rec
from test_pipeline_integration import execute, project


def published(cfg):
    db = connect(Path(cfg['state_dir']), readonly=True)
    try:
        release = current(db)
        return release, dyads.staged_rows(db, release['run_id'])
    finally:
        db.close()


def test_bootstrap_reorder_and_force_preserve_permanent_id_and_anchors(project):
    cfg, _, records, _ = project
    first = execute(project)
    _, rows = published(cfg)
    did = rows[0]['dyad_id']
    assert did.startswith('DYAD_') and rows[0]['registration_state'] == 'active'
    records.reverse()
    second = execute(project, force=True)
    _, rows2 = published(cfg)
    assert rows2[0]['dyad_id'] == did
    assert {k: first[k] for k in ('A','B','C','D','AC','BD','E','F','EF')} == {k: second[k] for k in ('A','B','C','D','AC','BD','E','F','EF')}
    assert execute(project)['unchanged']
    db = connect(Path(cfg['state_dir']))
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0] == 1
    reg = db.execute('SELECT * FROM dyad_registry').fetchone()
    assert reg['first_student_id'] == 's0' and reg['first_parent_id'] == 'p0'
    with pytest.raises(sqlite3.IntegrityError, match='immutable'):
        db.execute("UPDATE dyad_registry SET first_student_id='other'")
    db.rollback()
    with pytest.raises(sqlite3.IntegrityError, match='cannot be deleted'):
        db.execute('DELETE FROM dyad_registry')
    db.rollback()
    db.close()
    folder = Path(cfg['release_dir']) / second['release_id']
    for label in ('E_T0-student&parent-paired_final','F_T1-student&parent-paired_final','EF_T0T1-student&parent-paired_final'):
        frame = pd.read_excel(folder / '3.1.3_manual_pairing' / (label + '.xlsx'), sheet_name='paired')
        assert frame['dyad_id'].tolist() == [did]
    safe = pd.read_parquet(folder / 'dyad_tracking.parquet')
    assert not any('contact' in name or 'dob' in name for name in safe.columns)


def test_caregiver_change_does_not_change_family_id(project):
    cfg, _, records, _ = project
    records[2]['answer']['adult_demographics_06'] = 0
    records[3]['answer']['adult_demographics_06'] = 1
    execute(project)
    _, first = published(cfg)
    records[3]['answer']['adult_demographics_06'] = 2
    execute(project, force=True)
    _, second = published(cfg)
    assert second[0]['dyad_id'] == first[0]['dyad_id']
    db = connect(Path(cfg['state_dir']), readonly=True)
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0] == 1
    db.close()


def test_failed_output_reservation_is_reused_and_unpublished_not_invitable(project):
    cfg, _, _, _ = project
    with patch('study2.runner.write_outputs', side_effect=OSError('synthetic output interruption')):
        with pytest.raises(OSError):
            execute(project)
    db = connect(Path(cfg['state_dir']), readonly=True)
    reserved = db.execute('SELECT dyad_id,lifecycle FROM dyad_registry').fetchone()
    assert reserved['lifecycle'] == 'reserved' and current(db) is None
    db.close()
    with pytest.raises(PipelineError, match='尚未'):
        operator.export_invitations(cfg, Path(cfg['state_dir']) / 'list.xlsx')
    result=execute(project)
    assert published(cfg)[1][0]['dyad_id'] == reserved['dyad_id']
    db=connect(Path(cfg['state_dir']),readonly=True)
    assert db.execute('SELECT first_published_run_id FROM dyad_registry').fetchone()[0]==result['release_id']
    db.close()


@pytest.fixture
def registry(tmp_path):
    db = connect(tmp_path / 'state')
    dyads.migrate(db)
    cfg = {'minimum_followup_days': 90}
    def add_run(rid):
        with db:
            db.execute('INSERT INTO runs VALUES(?,?,?,?,?,?,?)', (rid,'fp','now','running','{}',None,None))
    add_run('r1')
    rows = [rec('s0'), rec('p0', role='parent'), rec('s1', 120), rec('p1', 120, role='parent')]
    for r in rows:
        r.update(child_id='child-a', selected_wave='T0' if r['submission_id'].endswith('0') else 'T1')
    result = dyads.stage(db, 'r1', rows, [], cfg)
    yield db, cfg, rows, result[0]['dyad_id'], add_run
    db.close()


def test_split_never_assigns_original_id_to_arbitrary_branch(registry):
    db, cfg, rows, did, add_run = registry
    changed = copy.deepcopy(rows)
    for r in changed:
        if r['submission_id'].endswith('1'):
            r.update(child_id='child-b', selected_wave='T0')
    add_run('r2')
    planned = dyads.plan(db, 'r2', changed, [], cfg)
    branches = [r for r in planned if r.get('child_id')]
    assert len(branches) == 2 and all(r['dyad_id'] is None for r in branches)
    assert all(r['registration_state'] == 'suspended' for r in branches)
    assert sum(r['dyad_id'] == did for r in planned) == 1
    assert next(r for r in planned if r['dyad_id'] == did)['child_id'] is None


def test_shared_phone_siblings_have_distinct_registrations_and_merge_blocks(registry):
    db, cfg, rows, did, add_run = registry
    siblings = [rec('sibling-s', dob='2013-02-02'), rec('sibling-p', role='parent', dob='2013-02-02')]
    for r in siblings:
        r.update(child_id='child-b', selected_wave='T0')
    add_run('r2')
    second = dyads.stage(db, 'r2', rows + siblings, [], cfg)
    ids = {r['dyad_id'] for r in second}
    assert len(ids) == 2 and did in ids
    changed = copy.deepcopy(rows + siblings)
    # Simulate a newly merged scientific group; permanent registry must not pick one.
    for r in changed:
        r['child_id'] = 'merged-child'
        if r['submission_id'].startswith('sibling'):
            r['selected_wave'] = None
    add_run('r3')
    planned = dyads.plan(db, 'r3', changed, [], cfg)
    group = next(r for r in planned if r.get('child_id'))
    assert group['dyad_id'] is None and group['registration_state'] == 'suspended'
    assert 'multiple_registered_families_merged' in group['registration_reason']


def test_changed_identity_and_contradictory_candidate_suspend(registry):
    db, cfg, rows, did, _ = registry
    changed = copy.deepcopy(rows)
    changed[0]['dob'] = '2013-03-03'
    plan = dyads.plan(db, 'r1', changed, [], cfg)
    assert plan[0]['registration_state'] == 'suspended' and plan[0]['dyad_id'] == did
    case = {'case_id':'case-synthetic','record_a':'s0','record_b':'s1','status':'pending',
            'kind':'identity','rule':'exact_school_contact_child_demo','reason':'Competing or contradictory candidate component'}
    plan = dyads.plan(db, 'r1', rows, [case], cfg)
    assert plan[0]['registration_state'] == 'suspended'


def test_excluded_original_anchor_requires_registration_adjudication(registry):
    db,cfg,rows,did,_=registry
    changed=copy.deepcopy(rows)
    changed[0].update(child_id=None,status='manual_excluded',selected_wave=None)
    changed[2]['selected_wave']='T0'
    plan=dyads.plan(db,'r1',changed,[],cfg)
    group=next(r for r in plan if r.get('child_id'))
    assert group['dyad_id']==did and group['registration_state']=='suspended'
    assert 'registered_submission_not_assigned' in group['registration_reason']


def test_new_latest_followups_reuse_id_and_keep_original_anchors(project):
    from test_store_and_review import source
    cfg,_,records,_=project
    execute(project);before=published(cfg)[1][0]
    for role,prefix,index in [('student','s',0),('parent','p',2)]:
        new=copy.deepcopy(records[index])
        new.update(source_key=prefix+'2',submission_id=prefix+'2',version_id='version_'+prefix+'2',
                   raw_hash='version_'+prefix+'2',observed_at=rec('time',240)['observed_at'],claimed_wave='T1')
        records.append(new)
    execute(project,force=True);after=published(cfg)[1][0]
    assert after['dyad_id']==before['dyad_id']
    assert after['student_t1_id']=='s2' and after['parent_t1_id']=='p2'
    assert after['registered_student_id']=='s0' and after['registered_parent_id']=='p0'


def t0_only(project):
    cfg, _, records, _ = project
    records[:] = [r for r in records if r['submission_id'].endswith('0')]
    execute(project)
    return cfg


def invitation(cfg):
    db = connect(Path(cfg['state_dir']), readonly=True)
    try:
        release, records, cases = operator._context(db, cfg)
        rows = operator.invitation_rows(db, cfg, release, records, cases)
        return release, rows
    finally:
        db.close()


def request_for(row, **changes):
    request = {'operation_id':'inv-synthetic','release_id':row['release_id'],'dyad_id':row['dyad_id'],
               'role':row['role'],'action':'sent','occurred_at':'2025-05-01T10:00:00+00:00',
               'operator':'Synthetic operator','reason':'Verified synthetic invitation',
               'method':'manual phone','result':'sent','eligibility_signature':row['eligibility_signature']}
    return {**request, **changes}


def test_invitation_is_manual_idempotent_and_does_not_change_science(project):
    cfg = t0_only(project)
    before, rows = invitation(cfg)
    assert len(rows) == 2 and all(r['sent_count'] == 0 for r in rows)
    path = Path(cfg['state_dir']) / 'controlled.xlsx'
    assert operator.export_invitations(cfg, path)['rows'] == 2
    from openpyxl import load_workbook
    wb=load_workbook(path)
    ws=wb['受控邀請清單']
    headers={c.value:c.column for c in ws[1]}
    assert ws.cell(2,headers['due_at']).data_type=='d'
    assert 'UTC' in ws.cell(2,headers['due_at']).number_format
    assert ws.column_dimensions[ws.cell(1,headers['eligibility_signature']).column_letter].hidden
    assert ws.freeze_panes=='A2' and ws.auto_filter.ref
    wb.close()
    request = request_for(rows[0])
    operator.record_invitation(cfg, request)
    assert operator.record_invitation(cfg, request)['unchanged']
    assert operator.record_invitation(cfg, {**request, 'operation_id':'another-operation'})['unchanged']
    after, current_rows = invitation(cfg)
    assert after == before
    assert next(r for r in current_rows if r['role'] == request['role'])['sent_count'] == 1
    db = connect(Path(cfg['state_dir']), readonly=True)
    assert db.execute('SELECT COUNT(*) FROM decisions').fetchone()[0] == 0
    assert db.execute('SELECT COUNT(*) FROM dyad_invitation_events').fetchone()[0] == 1
    db.close()


def test_applied_nonmatch_blocks_export_and_new_send_before_publication(project):
    from study2.review import validate_review_rows, commit_review_rows
    cfg = t0_only(project)
    _, before = invitation(cfg)
    db = connect(Path(cfg['state_dir']))
    release, records, cases = operator._context(db, cfg)
    case = next(c for c in cases if c['kind'] == 'family' and c['status'] == 'confirmed')
    row = {**case, 'release_id': release['release_id'], 'action': 'UNLINK', 'label': 'NONMATCH',
           'reviewer': 'Jarvis', 'reason': 'Verified different children in synthetic review'}
    commit_review_rows(db, validate_review_rows(db, cfg, [row]))
    db.close()
    output = Path(cfg['state_dir']) / 'blocked-invitations.xlsx'
    with pytest.raises(PipelineError, match='尚未發布'):
        operator.export_invitations(cfg, output)
    assert not output.exists()
    with pytest.raises(PipelineError, match='尚未發布'):
        operator.record_invitation(cfg, request_for(before[0]))
    db = connect(Path(cfg['state_dir']), readonly=True)
    assert current(db)['release_id'] == release['release_id']
    assert db.execute('SELECT COUNT(*) FROM dyad_invitation_events').fetchone()[0] == 0
    db.close()
    execute(project, force=True)
    assert invitation(cfg)[1] == []


def test_review_label_requires_republication_without_changing_permanent_id(project):
    from study2.review import validate_review_rows, commit_review_rows
    cfg = t0_only(project)
    _, before = invitation(cfg)
    db = connect(Path(cfg['state_dir']))
    release, _, cases = operator._context(db, cfg)
    case = next(c for c in cases if c['kind'] == 'family' and c['status'] == 'confirmed')
    row = {**case, 'release_id': release['release_id'], 'label': 'MATCH',
           'reviewer': 'Jarvis', 'reason': 'Verified same child in synthetic review'}
    commit_review_rows(db, validate_review_rows(db, cfg, [row]))
    db.close()
    with pytest.raises(PipelineError, match='尚未發布'):
        invitation(cfg)
    execute(project, force=True)
    _, after = invitation(cfg)
    assert len(after) == 2
    assert {r['dyad_id'] for r in after} == {r['dyad_id'] for r in before}


def test_applied_single_record_correction_also_blocks_stale_contact_export(project):
    from study2.review import validate_review_rows, commit_review_rows
    cfg = t0_only(project)
    db = connect(Path(cfg['state_dir']))
    release, records, _ = operator._context(db, cfg)
    record = next(r for r in records if r['role'] == 'student')
    row = {'case_id': 'NEW_RECORD', 'release_id': release['release_id'],
           'record_a': record['submission_id'], 'record_b': '', 'version_a': record['version_id'],
           'action': 'SET_FIELD', 'value': dumps({'field': 'contact_01', 'value': '62345678'}),
           'reviewer': 'Jarvis', 'reason': 'Verified corrected synthetic contact'}
    commit_review_rows(db, validate_review_rows(db, cfg, [row]))
    db.close()
    with pytest.raises(PipelineError, match='尚未發布'):
        operator.export_invitations(cfg, Path(cfg['state_dir']) / 'stale-contact.xlsx')


def test_existing_invitation_retry_survives_later_review_changes(project):
    from study2.review import validate_review_rows, commit_review_rows
    cfg = t0_only(project)
    _, rows = invitation(cfg)
    request = request_for(rows[0])
    operator.record_invitation(cfg, request)
    db = connect(Path(cfg['state_dir']))
    release, _, cases = operator._context(db, cfg)
    case = next(c for c in cases if c['kind'] == 'family' and c['status'] == 'confirmed')
    commit_review_rows(db, validate_review_rows(db, cfg, [{**case, 'release_id': release['release_id'],
                      'label': 'NONMATCH', 'reviewer': 'Jarvis', 'reason': 'Synthetic review change'}]))
    db.close()
    assert operator.record_invitation(cfg, request)['unchanged']
    db = connect(Path(cfg['state_dir']), readonly=True)
    assert db.execute('SELECT COUNT(*) FROM dyad_invitation_events').fetchone()[0] == 1
    db.close()


def test_newer_restored_applied_reviews_block_invites_until_main_is_recovered(project):
    from study2 import review_core as core
    cfg = t0_only(project)
    release, rows = invitation(cfg)
    work = core.workspace(cfg)
    with work:
        work.execute('INSERT INTO batches VALUES(?,?,?,?)', ('restored-applied-review', 'applied', 'now', '{}'))
    work.close()
    db = connect(Path(cfg['state_dir']))
    pointer = operator.backup_tracking(db, cfg, 'newer_review_workspace')
    db.close()
    destination = Path(cfg['state_dir']).parent / 'restored-newer-reviews'
    operator.restore_bundle(pointer, Path(release['path']) / 'backup.sqlite', destination)
    restored_cfg = {**cfg, 'state_dir': str(destination)}
    with pytest.raises(PipelineError, match='主資料庫缺少'):
        operator.export_invitations(restored_cfg, destination / 'blocked.xlsx')
    with pytest.raises(PipelineError, match='主資料庫缺少'):
        operator.record_invitation(restored_cfg, request_for(rows[0]))
    assert not (destination / 'blocked.xlsx').exists()


def test_unapplied_review_batches_do_not_block_published_invites(project):
    from study2 import review_core as core
    cfg = t0_only(project)
    work = core.workspace(cfg)
    with work:
        work.execute('INSERT INTO batches VALUES(?,?,?,?)', ('pending-approved-review', 'approved', 'now', '{}'))
    work.close()
    assert len(invitation(cfg)[1]) == 2


def test_stale_early_and_completed_invitation_blocked(project):
    cfg = t0_only(project)
    _, rows = invitation(cfg)
    request = request_for(rows[0])
    early={**request,'operation_id':'early-invitation','occurred_at':'2025-01-02T00:00:00+00:00'}
    operator.record_invitation(cfg,early)
    db=connect(Path(cfg['state_dir']),readonly=True)
    stored=json.loads(db.execute('SELECT payload_json FROM dyad_invitation_events WHERE operation_id=?',('early-invitation',)).fetchone()[0])
    assert stored['before_eligibility_date']
    assert not published(cfg)[1][0]['student_complete'] and not published(cfg)[1][0]['parent_complete']
    db.close()
    with pytest.raises(PipelineError, match='已改變'):
        operator.record_invitation(cfg, {**request, 'eligibility_signature':'stale'})
    execute(project, force=True)
    with pytest.raises(PipelineError, match='發布版本已改變'):
        operator.record_invitation(cfg, request)


def test_stop_resume_are_history_and_never_send(project):
    cfg = t0_only(project)
    _, rows = invitation(cfg)
    request = request_for(rows[0], action='stop')
    operator.record_invitation(cfg, request)
    assert len(invitation(cfg)[1]) == 1
    with pytest.raises(PipelineError, match='暫停'):
        operator.record_invitation(cfg, request_for(rows[0], operation_id='sent-after-stop'))
    operator.record_invitation(cfg, {**request, 'operation_id':'resume-op', 'action':'resume'})
    assert len(invitation(cfg)[1]) == 2


def test_restore_bundle_preserves_unpublished_invitation_and_review_history(project):
    cfg = t0_only(project)
    release, rows = invitation(cfg)
    operator.record_invitation(cfg, request_for(rows[0]))
    pointer = json.loads((Path(cfg['state_dir']) / 'DYAD_LEDGER.json').read_text())
    destination = Path(cfg['state_dir']).parent / 'restored-state'
    main_backup = Path(release['path']) / 'backup.sqlite'
    operator.restore_bundle(pointer['bundle'], main_backup, destination)
    restored_cfg = {**cfg, 'state_dir': str(destination)}
    db = connect(destination, readonly=True)
    operator.recovery_check(db, restored_cfg)
    assert db.execute('SELECT COUNT(*) FROM dyad_invitation_events').fetchone()[0] == 1
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0] == 1
    db.close()
    with pytest.raises(PipelineError, match='新資料夾'):
        operator.restore_bundle(pointer['bundle'], main_backup, destination)


def test_restore_rejects_older_bundle_that_would_revive_retired_registration(project):
    from study2.store import backup
    from study2.common import file_hash
    cfg = t0_only(project)
    release, rows = invitation(cfg)
    old_bundle = json.loads((Path(cfg['state_dir']) / 'DYAD_LEDGER.json').read_text())['bundle']
    did = rows[0]['dyad_id']
    operator.adjudicate(cfg, {'operation_id': 'retire-for-restore', 'release_id': release['release_id'],
                           'dyad_id': did, 'action': 'retire', 'reviewer': 'Jarvis',
                           'reason': 'Verified synthetic retirement'})
    db = connect(Path(cfg['state_dir']))
    newer_main = Path(cfg['state_dir']) / 'newer-retired-backup.sqlite'
    backup(db, newer_main)
    db.close()
    before_hash = file_hash(newer_main)
    destination = Path(cfg['state_dir']).parent / 'rejected-restore'
    with pytest.raises(PipelineError, match='停用|較舊'):
        operator.restore_bundle(old_bundle, newer_main, destination)
    assert not destination.exists()
    assert file_hash(newer_main) == before_hash
    assert invitation(cfg)[1] == []
    latest = json.loads((Path(cfg['state_dir']) / 'DYAD_LEDGER.json').read_text())['bundle']
    operator.restore_bundle(latest, Path(release['path']) / 'backup.sqlite', destination)
    restored_cfg = {**cfg, 'state_dir': str(destination)}
    restored = connect(destination, readonly=True)
    operator.recovery_check(restored, restored_cfg)
    assert restored.execute('SELECT lifecycle FROM dyad_registry WHERE dyad_id=?', (did,)).fetchone()[0] == 'retired'
    restored.close()
    assert invitation(restored_cfg)[1] == []


def test_restore_rejects_older_bundle_missing_a_newer_invitation_stop(project):
    from study2.store import backup
    cfg = t0_only(project)
    _, rows = invitation(cfg)
    old_bundle = json.loads((Path(cfg['state_dir']) / 'DYAD_LEDGER.json').read_text())['bundle']
    operator.record_invitation(cfg, request_for(rows[0], action='stop'))
    db = connect(Path(cfg['state_dir']))
    newer_main = Path(cfg['state_dir']) / 'newer-stop-backup.sqlite'
    backup(db, newer_main)
    db.close()
    destination = Path(cfg['state_dir']).parent / 'rejected-stop-restore'
    with pytest.raises(PipelineError, match='較舊|不完整'):
        operator.restore_bundle(old_bundle, newer_main, destination)
    assert not destination.exists()
    assert len(invitation(cfg)[1]) == 1


def test_restored_stale_database_cannot_issue_or_reuse_ids(project):
    cfg = t0_only(project)
    release, rows = invitation(cfg)
    operator.record_invitation(cfg, request_for(rows[0]))
    stale = sqlite3.connect(Path(release['path']) / 'backup.sqlite')
    stale.row_factory = sqlite3.Row
    try:
        with pytest.raises(PipelineError, match='遺漏'):
            operator.recovery_check(stale, cfg)
    finally:
        stale.close()


def install_scientific_version(db, rid, rows):
    """Synthetic post-review identity result; this never changes real rules."""
    with db:
        db.execute("UPDATE runs SET status='published' WHERE run_id=?", (rid,))
        for r in rows:
            db.execute('INSERT OR IGNORE INTO submissions VALUES(?,?,?,?,?)',
                       (r['submission_id'], r['submission_id'], r['role'], 'synthetic', 'now'))
            db.execute('INSERT INTO normalized_records VALUES(?,?,?)', (rid,r['submission_id'],dumps(r)))
        db.execute('INSERT INTO releases VALUES(?,?,?,?,?)', (rid,rid,'now','unused','{}'))
        db.execute('INSERT OR REPLACE INTO publication VALUES(1,?)', (rid,))


def test_explicit_split_retains_one_id_creates_another_and_is_idempotent(registry):
    db, cfg, rows, did, add_run = registry
    cfg = {**cfg, 'state_dir':str(Path(db.execute('PRAGMA database_list').fetchone()[2]).parent)}
    changed = copy.deepcopy(rows)
    for r in changed:
        if r['submission_id'].endswith('1'):
            r.update(child_id='child-b', selected_wave='T0')
    add_run('r2')
    install_scientific_version(db,'r2',changed)
    value = operator.preview(cfg)
    request = {'operation_id':'split-op','release_id':'r2','action':'split','dyad_id':did,
               'children':['child-a','child-b'],'reviewer':'Jarvis','reason':'Synthetic reviewed split',
               'signatures':{r['child_id']:r['evidence_signature'] for r in value['rows'] if r.get('child_id')}}
    result = operator.adjudicate(cfg, request)
    assert result['assignments'][0]['dyad_id'] == did
    assert result['assignments'][1]['dyad_id'] != did
    assert operator.adjudicate(cfg, request)['unchanged']
    plan = dyads.plan(db,'r2',changed,[],cfg)
    assert len(plan) == 2 and all(r['registration_state'] == 'active' for r in plan)
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0] == 2
    assert db.execute('SELECT first_student_id FROM dyad_registry WHERE dyad_id=?',(did,)).fetchone()[0] == 's0'


def test_split_can_retain_original_id_on_later_branch(registry):
    db,cfg,rows,did,add_run=registry
    cfg={**cfg,'state_dir':str(Path(db.execute('PRAGMA database_list').fetchone()[2]).parent)}
    changed=copy.deepcopy(rows)
    for r in changed:
        if r['submission_id'].endswith('1'):r.update(child_id='child-b',selected_wave='T0')
    add_run('r2');install_scientific_version(db,'r2',changed)
    value=operator.preview(cfg)
    request={'operation_id':'retain-later','release_id':'r2','action':'split','dyad_id':did,
             'children':['child-b','child-a'],'reviewer':'Jarvis','reason':'Verified historical invitation ownership',
             'signatures':{r['child_id']:r['evidence_signature'] for r in value['rows'] if r.get('child_id')}}
    result=operator.adjudicate(cfg,request)
    assert result['assignments'][0]=={'child_id':'child-b','dyad_id':did}
    plan=dyads.plan(db,'r2',changed,[],cfg)
    assert all(r['registration_state']=='active' for r in plan)
    assert next(r for r in plan if r['child_id']=='child-a')['dyad_id']!=did
    assert db.execute('SELECT first_student_id FROM dyad_registry WHERE dyad_id=?',(did,)).fetchone()[0]=='s0'


def test_chained_aliases_preserve_stop_history_and_can_be_resumed(project):
    cfg=t0_only(project);release,rows=invitation(cfg);old=rows[0]['dyad_id'];role=rows[0]['role']
    operator.record_invitation(cfg,request_for(rows[0],action='stop'))
    db=connect(Path(cfg['state_dir']))
    original=dict(db.execute('SELECT * FROM dyad_registry WHERE dyad_id=?',(old,)).fetchone())
    with db:
        for did,life in [('historical-middle','retired'),('canonical-final','active')]:
            clone={**original,'dyad_id':did,'lifecycle':life,'first_published_run_id':None}
            db.execute(f"INSERT INTO dyad_registry ({','.join(clone)}) VALUES({','.join('?' for _ in clone)})",tuple(clone.values()))
        db.execute("UPDATE dyad_registry SET lifecycle='retired' WHERE dyad_id=?",(old,))
        entries=[dict(r) for r in db.execute('SELECT * FROM dyad_lineage WHERE dyad_id=?',(old,))]
        db.execute('UPDATE dyad_lineage SET active=0 WHERE dyad_id=?',(old,))
        for entry in entries:
            db.execute('INSERT INTO dyad_lineage VALUES(?,?,?,1)',('canonical-final',entry['submission_id'],entry['evidence_hash']))
        db.execute('INSERT INTO dyad_aliases VALUES(?,?,?)',(old,'historical-middle','merge-first'))
        db.execute('INSERT INTO dyad_aliases VALUES(?,?,?)',('historical-middle','canonical-final','merge-second'))
    assert dyads.related_ids(db,'canonical-final')==sorted([old,'historical-middle','canonical-final'])
    db.close()
    execute(project,force=True)
    assert len(invitation(cfg)[1])==1
    db=connect(Path(cfg['state_dir']),readonly=True)
    current_release,recs,cases=operator._context(db,cfg)
    candidates=operator.invitation_rows(db,cfg,current_release,recs,cases,include_stopped=True)
    stopped=next(r for r in candidates if r['role']==role)
    assert stopped['dyad_id']=='canonical-final' and stopped['invitation_state']=='stop'
    db.close()
    operator.record_invitation(cfg,request_for(stopped,operation_id='resume-alias',action='resume'))
    assert len(invitation(cfg)[1])==2


def test_newer_schema_is_rejected_before_mutating_tables(registry):
    db,_,_,_,_=registry
    with db:db.execute('INSERT INTO dyad_schema VALUES(2,?)',('future',))
    before=[tuple(r) for r in db.execute('SELECT name,sql FROM sqlite_master ORDER BY name')]
    with pytest.raises(PipelineError,match='Unsupported'):
        dyads.migrate(db)
    assert [tuple(r) for r in db.execute('SELECT name,sql FROM sqlite_master ORDER BY name')]==before


def test_merge_requires_explicit_survivor_and_keeps_alias_history(registry):
    db,cfg,rows,did,add_run=registry
    cfg={**cfg,'state_dir':str(Path(db.execute('PRAGMA database_list').fetchone()[2]).parent)}
    siblings=[rec('sibling-s'),rec('sibling-p',role='parent')]
    for r in siblings:r.update(child_id='child-b',selected_wave='T0')
    add_run('r2')
    second=dyads.stage(db,'r2',rows+siblings,[],cfg)
    other=next(r['dyad_id'] for r in second if r['dyad_id']!=did)
    changed=copy.deepcopy(rows+siblings)
    for r in changed:
        r['child_id']='merged-child'
        if r['submission_id'].startswith('sibling'):r['selected_wave']=None
    add_run('r3');install_scientific_version(db,'r3',changed)
    value=operator.preview(cfg)
    group=next(r for r in value['rows'] if r.get('child_id'))
    request={'operation_id':'merge-op','release_id':'r3','action':'merge','dyad_id':did,'merge_ids':[other],
             'children':['merged-child'],'reviewer':'Jarvis','reason':'Synthetic verified duplicate registrations',
             'signatures':{'merged-child':group['evidence_signature']}}
    operator.adjudicate(cfg,request)
    result=dyads.plan(db,'r3',changed,[],cfg)
    assert len(result)==1 and result[0]['dyad_id']==did and result[0]['registration_state']=='active'
    assert db.execute('SELECT canonical_id FROM dyad_aliases WHERE alias_id=?',(other,)).fetchone()[0]==did
    assert db.execute('SELECT lifecycle FROM dyad_registry WHERE dyad_id=?',(other,)).fetchone()[0]=='retired'
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0]==2


def test_retirement_never_automatically_reissues_same_family_id(project):
    cfg=t0_only(project);release,rows=published(cfg);did=rows[0]['dyad_id']
    operator.adjudicate(cfg,{'operation_id':'retire-op','release_id':release['release_id'],'action':'retire',
                          'dyad_id':did,'reviewer':'Jarvis','reason':'Synthetic withdrawal of registration'})
    second=execute(project,force=True)
    assert second['E']==1
    db=connect(Path(cfg['state_dir']),readonly=True)
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0]==1
    assert db.execute('SELECT lifecycle FROM dyad_registry').fetchone()[0]=='retired'
    assert all(r['registration_state']!='active' for r in dyads.staged_rows(db,second['release_id']))
    db.close()
    assert invitation(cfg)[1]==[]


def test_validation_to_publication_rechecks_identity_and_administrative_changes(project):
    from study2.runner import publish_release
    cfg,_,_,_=project
    first=execute(project)
    candidate=execute(project,publish=False,force=True)
    row=published(cfg)[1][0]
    operator.adjudicate(cfg,{'operation_id':'retire-between','release_id':first['release_id'],'action':'retire',
                          'dyad_id':row['dyad_id'],'reviewer':'Jarvis','reason':'Synthetic changed decision'})
    db=connect(Path(cfg['state_dir']))
    with pytest.raises(PipelineError,match='lineage changed'):
        publish_release(db,cfg,candidate['release_id'])
    assert current(db)['release_id']==first['release_id']
    db.close()


def test_publication_backup_failure_keeps_previous_pointer_and_ids(project):
    from study2.runner import publish_release
    cfg,_,_,_=project
    first=execute(project)
    candidate=execute(project,publish=False,force=True)
    original=operator.backup_tracking
    def fail_after(db,cfg,reason='manual',**kw):
        if reason.startswith('after_publish'):raise OSError('synthetic backup interruption')
        return original(db,cfg,reason,**kw)
    db=connect(Path(cfg['state_dir']))
    with patch('study2.dyad_operator.backup_tracking',side_effect=fail_after),pytest.raises(OSError):
        publish_release(db,cfg,candidate['release_id'])
    assert current(db)['release_id']==first['release_id']
    assert json.loads((Path(cfg['release_dir'])/'CURRENT.json').read_text())['release_id']==first['release_id']
    operator.recovery_check(db,cfg)
    publish_release(db,cfg,candidate['release_id'])
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0]==1
    db.close()


def test_tracking_restore_keeps_unapproved_reviews(project):
    from study2 import review_core as core
    cfg=t0_only(project);release,rows=invitation(cfg)
    package=core.create_package(cfg,'Synthetic reviewer',limit=1,kinds=['family'])
    task=package['tasks'][0]
    core.submit_reviews(cfg,package['package_id'],[{'task_id':task['task_id'],'verdict':'資料不足',
                        'evidence_source':'沒有足夠證據','reason':'Synthetic pending answer'}])
    operator.record_invitation(cfg,request_for(rows[0]))
    pointer=json.loads((Path(cfg['state_dir'])/'DYAD_LEDGER.json').read_text())
    destination=Path(cfg['state_dir']).parent/'restored-reviews'
    operator.restore_bundle(pointer['bundle'],Path(release['path'])/'backup.sqlite',destination)
    db=sqlite3.connect(destination/'review.sqlite')
    assert db.execute('SELECT COUNT(*) FROM reviews').fetchone()[0]==1
    assert db.execute('SELECT COUNT(*) FROM tasks').fetchone()[0]==1
    db.close()


def test_research_rollback_keeps_newer_administrative_history(project):
    from study2.runner import publish_release
    cfg=t0_only(project);first,rows=invitation(cfg)
    operator.record_invitation(cfg,request_for(rows[0]))
    second=execute(project,force=True)
    db=connect(Path(cfg['state_dir']))
    publish_release(db,cfg,first['release_id'],rollback=True)
    assert current(db)['release_id']==first['release_id']
    assert db.execute('SELECT COUNT(*) FROM dyad_invitation_events').fetchone()[0]==1
    assert db.execute('SELECT COUNT(*) FROM dyad_registry').fetchone()[0]==1
    operator.recovery_check(db,cfg)
    db.close()


def test_dashboard_tracking_search_and_history_are_read_only(project,monkeypatch):
    from streamlit.testing.v1 import AppTest
    from study2.common import atomic_json
    cfg,_,_,_=project
    cfg['inbox_dir']=str(Path(cfg['project_root'])/'inbox')
    cfg['stale_after_days']=8
    execute(project)
    release,rows=published(cfg)
    config=Path(cfg['project_root'])/'config.json';atomic_json(config,cfg)
    monkeypatch.setenv('STUDY2_CONFIG',str(config))
    app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'dashboard.py')).run(timeout=60)
    assert not app.exception
    assert any(t.label=='親子追蹤' for t in app.tabs)
    app.text_input(key='dyad_search').set_value(rows[0]['dyad_id']).run(timeout=60)
    assert not app.exception
    app.selectbox(key='dyad_history').set_value(rows[0]['dyad_id']).run(timeout=60)
    assert not app.exception
    assert published(cfg)[0]==release
