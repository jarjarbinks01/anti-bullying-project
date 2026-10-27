"""Versioned 4.1 workbooks and numeric results, built from the shared service."""
import io
import json
import os
from pathlib import Path

import pandas as pd

from .common import PipelineError, atomic_json, file_hash, secure_directory
from .descriptive_data import demographic_values
from .descriptive_spec import (DICTIONARY_VERSION, GROUPS, MODES, REPORT_FILES, REPORT_RANGES, VERSION,
                               dictionary_rows, variables)
from .descriptives import (cohort, filter_cohort, formatted, paired_changes, statistics,
                           validate_statistics, wave_counts, prepare_report_cohorts)

NOTES = [
    ('計分', '重用本版共用量表分數；所需題目完整才計分；SEC 已轉為 1–4；PHQ 加總，其餘沿用共用平均。'),
    ('比例分母', '有效回答人數；每變項另列缺失。SEN 可複選，各類比例可合計超過 100%。'),
    ('分組', '九組重疊，不可相加。家長樣本按孩子性別分組。跨期固定 T0 分組，各輪樣本按該輪分組。'),
    ('日期及年齡', '沿用現有年齡可信度；孩子按觀察日期，成人只有出生年份而屬估算；批次分輪許可不放寬年齡可信度。'),
    ('成人性別', '原問卷未獨立收集；顯示填卷人與孩子關係，不自行推定成人性別。'),
    ('子女人數', '只解讀已確認的 Qualtrics 1–4 編碼，其他來源編碼待核；4 為四個或以上，不求平均。'),
    ('填卷人', 'EF 可換填卷人，成人 PHQ 跨期差只在有同一成人確認證據時提供，BD 名稱本身不代表確認；獨立家長樣本另列，不併入本四份報表。'),
    ('有效範圍', 'Jarvis於2026-10-17確認：孩子5–20歲、成人20–70歲、身高100–200cm、體重20–100kg；兄弟姊妹各項0–5、總數0–10，均含上下限。範圍外只列該指標缺失，原值及樣本保留，另列數量。'),
    ('BMI', '身高及體重均有效才計算 kg/m² 連續值；不採用成人分界為兒童分類。'),
    ('解讀', '描述數值及分布；PHQ 分數組別不代表臨床診斷，兩輪差異不自動代表研究效果。'),
]


def _safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


def _write_frame(sheet, frame):
    sheet.append([str(c) for c in frame.columns])
    for record in frame.itertuples(index=False, name=None):
        sheet.append([_safe(None if pd.isna(v) else v) for v in record])
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = 'C2'


def _matrix(sheet, stats, section, release_id, mode):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    roles = ('student', 'parent') if section == 'scales' else (section,)
    width = 1 + 18 * len(roles)
    sheet.append([MODES[mode] + '：' + {'scales': '量表及欺凌', 'student': '學生人口特徵', 'parent': '家長人口特徵'}[section]])
    sheet.append(['發布：' + release_id + '；主流學校；有效分母；' + VERSION])
    sheet.append(['平均（SD）或 n／有效分母（%）；每變項另列有效／缺失 N。跨期分組固定 T0。'])
    data = stats[stats['section'].eq(section)]
    lookup = {(r['role'], r['group'], r['wave'], r['variable_id'], r['category_code']): r
              for r in data.to_dict('records')}
    for i, role in enumerate(roles):
        left = 2 + i * 18
        for g, (group, title) in enumerate(GROUPS.items()):
            col = left + g*2
            sheet.cell(4, col, '學生' if role == 'student' else '家長')
            sheet.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col+1)
            sheet.cell(5, col, title)
            sheet.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col+1)
            for j, wave in enumerate(('T0', 'T1')):
                block = data[data['role'].eq(role) & data['group'].eq(group) & data['wave'].eq(wave)]
                n = int(block['sample_n'].iloc[0]) if len(block) else 0
                sheet.cell(6, col+j, f'{wave}（樣本 N={n:,}）')
    sheet.cell(6, 1, '指標／分類（單位）')
    representative = variables(roles[0], section)
    for var in representative:
        entries = [('', '', 'value')] if var.kind in {'continuous', 'binary'} else [(code, label, 'value') for code, label in var.categories]
        if var.kind in {'category', 'multi'}:
            sheet.append([var.label])
        fields = [('', '有效 N', 'valid_n'), ('', '缺失 N', 'missing_n')]
        if var.key.removeprefix('dem_') in REPORT_RANGES or var.key == 'dem_bmi':
            fields.append(('', '超出有效範圍 N', 'outside_range_n'))
        for code, label, field in entries + fields:
            current = sheet.max_row+1
            sheet.cell(current, 1, (var.label + (f'（{var.unit}）' if var.unit else '')) if not label else '  ' + label)
            for i, role in enumerate(roles):
                for g, group in enumerate(GROUPS):
                    for j, wave in enumerate(('T0', 'T1')):
                        key_code = code if code or var.kind in {'continuous', 'binary'} else var.categories[0][0]
                        if var.kind == 'binary':
                            key_code = '1'
                        item = lookup.get((role, group, wave, var.key, key_code))
                        if item:
                            value = formatted(item) if field == 'value' else int(item[field])
                            sheet.cell(current, 2+i*18+g*2+j, value)
        if var.key in {'dem_child_age', 'dem_adult_age'}:
            for field, label in [('minimum', '最小值'), ('maximum', '最大值')]:
                current = sheet.max_row+1; sheet.cell(current, 1, '  '+label)
                for i, role in enumerate(roles):
                    for g, group in enumerate(GROUPS):
                        for j, wave in enumerate(('T0', 'T1')):
                            item = lookup.get((role, group, wave, var.key, ''))
                            val = item[field] if item else None
                            sheet.cell(current, 2+i*18+g*2+j, None if val is None or pd.isna(val) else float(val))
    sheet.freeze_panes = 'B7'
    sheet.column_dimensions['A'].width = 53
    for col in range(2, width+1):
        sheet.column_dimensions[get_column_letter(col)].width = 23
    for row in sheet.iter_rows(min_row=4, max_row=6):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='224B57')
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 32
    sheet.row_dimensions[1].height = 28
    sheet['A1'].font = Font(name='Arial', size=14, bold=True)
    sheet.print_title_rows = '4:6'
    sheet.print_title_cols = 'A:A'
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.scale = 100
    from openpyxl.worksheet.pagebreak import Break
    for end_column in range(7, width, 6):
        sheet.col_breaks.append(Break(id=end_column))
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f'A1:{get_column_letter(width)}{sheet.max_row}'


def workbook_bytes(stats, changes, mode, release_id):
    """Pure adapter over finished numeric statistics. No scoring or sampling."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook(); book.remove(book.active)
    for name, section in [('Scales', 'scales'), ('student', 'student'), ('parent', 'parent')]:
        _matrix(book.create_sheet(name), stats, section, release_id, mode)
    _write_frame(book.create_sheet('統計明細'), stats)
    notes = [('發布版本', release_id), ('報表版本', VERSION), ('字典版本', DICTIONARY_VERSION),
             ('分析模式', MODES[mode]), ('學校範圍', '主流學校'), *NOTES]
    notes += [(v['variable_id'] + '：' + v['label'], v['source'] + '；' + v['categories'] + '；' + v['note'])
              for v in dictionary_rows()]
    ws = book.create_sheet('口徑及版本'); ws.append(['項目', '說明'])
    for row in notes:
        ws.append(list(map(_safe, row)))
    ws.column_dimensions['A'].width = 50; ws.column_dimensions['B'].width = 110
    ws.freeze_panes = 'A2'
    if not changes.empty:
        _write_frame(book.create_sheet('配對差'), changes)
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                if cell.row > 6 or sheet.title not in {'Scales', 'student', 'parent'}:
                    cell.font = Font(name='Arial', size=10)
                    cell.alignment = Alignment(vertical='center', wrap_text=sheet.title == '口徑及版本')
                    if isinstance(cell.value, float):
                        cell.number_format = '0.0000'
            if sheet.title in {'Scales', 'student', 'parent'} and row[0].row > 6:
                row[0].alignment = Alignment(vertical='center', wrap_text=True)
                if len(str(row[0].value or '')) > 23:
                    sheet.row_dimensions[row[0].row].height = 32
            if sheet.title == '口徑及版本':
                sheet.row_dimensions[row[0].row].height = 42
        if sheet.title not in {'Scales', 'student', 'parent', '口徑及版本'}:
            from openpyxl.utils import get_column_letter
            for column in range(1, sheet.max_column+1):
                name = str(sheet.cell(1, column).value)
                width = max(len(str(sheet.cell(r, column).value or '')) for r in range(1, min(sheet.max_row, 100)+1))
                sheet.column_dimensions[get_column_letter(column)].width = min(55, max(18, width+2))
                if name in {'variable', 'note', 'missing_detail'}:
                    sheet.column_dimensions[get_column_letter(column)].width = 65
                    for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                        cells[0].alignment = Alignment(vertical='center', wrap_text=True)
                        sheet.row_dimensions[cells[0].row].height = 32
            for cell in sheet[1]:
                cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='224B57')
    stream = io.BytesIO(); book.save(stream)
    return stream.getvalue()


def write_reports(folder, measurements, frame, parents, release_id, *, formatted=True, prepared=None):
    """Finish every required report before the runner may validate the release."""
    all_stats, all_changes, counts = [], [], {}
    target = Path(folder)/'4.1_descriptive'
    if formatted:
        secure_directory(target)
    for mode, filename in REPORT_FILES.items():
        mode_stats, mode_changes = [], []
        counts[mode] = {}
        for role in ('student', 'parent'):
            if prepared is None:
                unfiltered = cohort(frame, parents, measurements, mode, role)
                selected, _ = filter_cohort(unfiltered)
                selected.attrs['mode'] = mode
            else:
                unfiltered, selected = prepared[mode, role]
            counts[mode][role] = {'all_schools': wave_counts(unfiltered), 'mainstream': wave_counts(selected)}
            mode_stats.append(statistics(selected, role, GROUPS, release_id))
            mode_changes.append(paired_changes(selected, role, GROUPS, release_id))
        stats = pd.concat(mode_stats, ignore_index=True)
        changes = pd.concat(mode_changes, ignore_index=True)
        validate_statistics(stats)
        if formatted:
            data = workbook_bytes(stats, changes, mode, release_id)
            path = target/filename; temporary = target/(filename+'.partial')
            try:
                with temporary.open('wb') as handle:
                    temporary.chmod(0o600); handle.write(data); handle.flush(); os.fsync(handle.fileno())
                os.replace(temporary, path)
            finally:
                temporary.unlink(missing_ok=True)
        all_stats.append(stats); all_changes.append(changes)
    combined = pd.concat(all_stats, ignore_index=True)
    combined.to_parquet(Path(folder)/'descriptive_statistics.parquet', index=False)
    changes = pd.concat(all_changes, ignore_index=True)
    changes.to_parquet(Path(folder)/'descriptive_changes.parquet', index=False)
    result = dict(version=VERSION, dictionary_version=DICTIONARY_VERSION, statistics_rows=len(combined),
                  counts=counts, reports={name: file_hash(target/name) for name in REPORT_FILES.values()} if formatted else {},
                  formatted_reports={'status': 'written' if formatted else 'deferred',
                                     'files': list(REPORT_FILES.values())},
                  numeric_hashes={name: file_hash(Path(folder)/name) for name in
                                 ('descriptive_statistics.parquet', 'descriptive_changes.parquet')})
    atomic_json(Path(folder)/'descriptive_manifest.json', result)
    verify_reports(folder, result)
    return result


def prepare_reports(folder, by_id, scored_ids, analysis_row, caregivers=(), *, formatted=True):
    measurements = pd.DataFrame([demographic_values(by_id[sid], analysis_row(by_id[sid])) for sid in sorted(scored_ids)])
    if measurements.empty:
        measurements = pd.DataFrame(columns=['submission_id', 'descriptive_dictionary_version', 'school_stage'])
    measurements.to_parquet(Path(folder)/'descriptive_measurements.parquet', index=False)
    atomic_json(Path(folder)/'descriptive_dictionary.json',
                {'version': DICTIONARY_VERSION, 'report_version': VERSION, 'variables': dictionary_rows(), 'notes': NOTES})
    pd.DataFrame([{k: row.get(k) for k in ['child_id', 'caregiver_status', 'reported_relationship_changed']}
                  for row in caregivers], columns=['child_id', 'caregiver_status', 'reported_relationship_changed']).to_parquet(
                      Path(folder)/'descriptive_caregivers.parquet', index=False)
    frame = pd.read_parquet(Path(folder)/'analysis.parquet')
    parent_path = Path(folder)/'parent_longitudinal.parquet'
    parents = pd.read_parquet(parent_path) if parent_path.exists() else None
    prepared = prepare_report_cohorts(frame, parents, measurements)
    result = write_reports(folder, measurements, frame, parents, Path(folder).name,
                           formatted=formatted, prepared=prepared)
    from .descriptive_exploration_outputs import write_exploration
    result['exploration'] = write_exploration(folder, measurements, frame, parents, Path(folder).name,
                                             prepared=prepared)
    return result


def verify_reports(folder, details):
    folder = Path(folder)
    stats = pd.read_parquet(folder/'descriptive_statistics.parquet')
    validate_statistics(stats)
    if len(stats) != details['statistics_rows'] or set(stats['mode']) != set(MODES):
        raise PipelineError('Descriptive report coverage is incomplete')
    for name, expected in details['reports'].items():
        if file_hash(folder/'4.1_descriptive'/name) != expected:
            raise PipelineError('Descriptive workbook fingerprint mismatch')
    contract = details.get('formatted_reports')
    deferred = contract is not None and contract.get('status') == 'deferred'
    if contract is not None:
        if (set(contract) != {'status', 'files'} or contract['status'] not in {'written', 'deferred'}
                or contract['files'] != list(REPORT_FILES.values())):
            raise PipelineError('Invalid formatted report output contract')
        pins = details.get('numeric_hashes', {})
        if set(pins) != {'descriptive_statistics.parquet', 'descriptive_changes.parquet'}:
            raise PipelineError('Formatted reports require pinned numeric inputs')
        for name, expected in pins.items():
            if file_hash(folder/name) != expected:
                raise PipelineError('Formatted report numeric input fingerprint mismatch')
    if deferred and details['reports']:
        raise PipelineError('Deferred workbooks must not be declared as written')
    if not deferred and set(details['reports']) != set(REPORT_FILES.values()):
        raise PipelineError('A required 4.1 workbook is missing')
    if details.get('exploration'):
        from .descriptive_exploration_outputs import verify_exploration
        verify_exploration(folder, details['exploration'])


def _reference_inputs(folder):
    folder = Path(folder)
    details = json.loads((folder/'descriptive_manifest.json').read_text())
    if details.get('version') != VERSION or details.get('dictionary_version') != DICTIONARY_VERSION:
        raise PipelineError('Use the report renderer frozen with this publication')
    verify_reports(folder, details)
    pins = details.get('numeric_hashes')
    if not pins:
        return details, None, None
    stats = pd.read_parquet(folder/'descriptive_statistics.parquet')
    changes = pd.read_parquet(folder/'descriptive_changes.parquet')
    for table in (stats, changes):
        if not table.empty and not table['release_id'].eq(folder.name).all():
            raise PipelineError('Numeric report belongs to another release')
    return details, stats, changes


def _reference_books(folder, modes):
    folder = Path(folder)
    if set(modes) - set(REPORT_FILES):
        raise PipelineError('Unknown descriptive report mode')
    details, stats, changes = _reference_inputs(folder)
    if stats is None:
        return {mode: (folder/'4.1_descriptive'/REPORT_FILES[mode]).read_bytes() for mode in modes}
    books = {mode: workbook_bytes(stats[stats['mode'].eq(mode)], changes[changes['mode'].eq(mode)],
                                  mode, folder.name) for mode in modes}
    if any(file_hash(folder/name) != expected for name, expected in details['numeric_hashes'].items()):
        raise PipelineError('Numeric report inputs changed during rendering')
    return books


def reference_workbook(folder, mode):
    """Render canonical OFF/reference values; no scoring or publication write."""
    return _reference_books(folder, [mode])[mode]


def export_reference_reports(folder, destination):
    """Create a separate new directory; never add files to a frozen release."""
    folder, destination = Path(folder).resolve(), Path(destination).absolute()
    if destination.is_symlink() or destination.exists():
        raise PipelineError('Report destination must be a new directory')
    destination = destination.resolve()
    if destination == folder or folder in destination.parents or destination in folder.parents:
        raise PipelineError('Report destination must be separate from the frozen release')
    # Render/verify before creating any output. The destination is exclusive.
    books = {REPORT_FILES[mode]: data for mode, data in _reference_books(folder, REPORT_FILES).items()}
    destination.mkdir(mode=0o700)
    for name, data in books.items():
        with (destination/name).open('xb') as handle:
            (destination/name).chmod(0o600)
            handle.write(data)
    result = {'release_id': folder.name, 'membership': 'canonical_OFF_reference',
              'report_version': VERSION, 'files': {name: file_hash(destination/name) for name in books}}
    atomic_json(destination/'REPORT_EXPORT.json', result)
    return result


def render_reference_downloads(st, folder, summary, *, exclusions_on=False):
    """Prepare deferred downloads only on request, preserving OFF-only labeling."""
    if exclusions_on:
        st.info('Original OFF/reference workbooks are disabled while ON is requested. Use the current filtered analysis downloads, or select OFF.')
        return
    deferred = summary['descriptive'].get('formatted_reports', {}).get('status') == 'deferred'
    for mode, name in REPORT_FILES.items():
        if deferred:
            if not st.button('Prepare ' + MODES[mode], key='ds_prepare_' + mode):
                continue
            data = reference_workbook(folder, mode)
        else:
            data = (Path(folder)/'4.1_descriptive'/name).read_bytes()
        st.download_button(MODES[mode], data, 'OFF-reference-' + name,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', key='ds_report_' + mode)


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Render four OFF/reference workbooks without recalculation')
    parser.add_argument('--release', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(export_reference_reports(args.release, args.output), ensure_ascii=False, indent=2))
