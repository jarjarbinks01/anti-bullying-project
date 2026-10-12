from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd

from .common import contacts, dumps, plain, text, timestamp


def source_category(label):
    if label.startswith("paperform"):
        return "paperform"
    if label.startswith("separate_data"):
        return "school_batch"
    if label.startswith("socialworkdream"):
        return "socialworkdream"
    return "qualtrics" if "qualtrics" in label else "legacy_online"


def complete_score(answer: dict, role: str, domain: str, indices: list[int], mode="mean", shift=0):
    values = [answer.get(f"{role}_{domain}_{i:02d}") for i in indices]
    if any(v is None or text(v) == "" for v in values):
        return None
    values = [float(v) + shift for v in values]
    return sum(values) if mode == "sum" else sum(values) / len(values)


def metrics(answer: dict, role: str) -> dict:
    result = {
        "PHQ_anxiety": complete_score(answer, role, "PHQ", [1, 2], "sum"),
        "PHQ_depression": complete_score(answer, role, "PHQ", [3, 4], "sum"),
        "PHQ_total": complete_score(answer, role, "PHQ", list(range(1, 5)), "sum"),
        "FQOL_total": complete_score(answer, role, "FQOL", list(range(1, 22))),
        "FQOL_family_interaction": complete_score(answer, role, "FQOL", list(range(1, 7))),
        "FQOL_parenting": complete_score(answer, role, "FQOL", list(range(7, 13))),
        "FQOL_emotional_wellbeing": complete_score(answer, role, "FQOL", list(range(13, 17))),
        "FQOL_physical_wellbeing": complete_score(answer, role, "FQOL", list(range(17, 22))),
        "SEC_total": complete_score(answer, role, "SEC", list(range(1, 23)), shift=1),
        "SEC_responsible_decision_making": complete_score(answer, role, "SEC", [1, 5, 9, 13], shift=1),
        "SEC_social_awareness": complete_score(answer, role, "SEC", [2, 6, 10, 14], shift=1),
        "SEC_self_management": complete_score(answer, role, "SEC", [3, 7, 11, 15], shift=1),
        "SEC_relationship_skills": complete_score(answer, role, "SEC", [4, 8, 12, 16], shift=1),
        "SEC_self_awareness": complete_score(answer, role, "SEC", list(range(17, 23)), shift=1),
    }
    for domain, length in [("TBV", 7), ("TBP", 7), ("CBV", 6), ("CBP", 6)]:
        result[domain + "_mean"] = complete_score(answer, role, domain, list(range(1, length + 1)))
        values = [answer.get(f"{role}_{domain}_{i:02d}") for i in range(1, length + 1)]
        for threshold in [1, 2]:
            # Require complete domain for prevalence: missing items cannot be
            # interpreted as absence of bullying.
            result[f"{domain}_prevalence{threshold}"] = None if any(v is None or text(v) == "" for v in values) else int(any(float(v) >= threshold for v in values))
    return result


def excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    def safe(v):
        if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
            return "'" + v
        return v
    return frame.map(safe)


def write_excel(path: Path, sheets: dict[str, list | pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, rows in sheets.items():
            frame = rows if isinstance(rows, pd.DataFrame) else pd.DataFrame(rows)
            excel_safe(frame).to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(48, max(14, len(str(col[0].value or "")) + 2))
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="224B57")
    path.chmod(0o600)


def export_row(record: dict) -> dict:
    prefix = "STU" if record["role"] == "student" else "PAR"
    order = "01" if record["selected_wave"] == "T0" else "02"
    return {"same_person_entry_id": prefix + "_" + record["child_id"].split("_", 1)[1] + "_" + order,
            "same_person_match_rule": "v3_confirmed_identity", **record["answer"],
            "submission_id": record["submission_id"], "child_id": record["child_id"],
            "source_claimed_timepoint": record["claimed_wave"], "final_timepoint": record["selected_wave"],
            "wave_reason": record.get("wave_reason"),
            "adult_birth_year_status": record.get("adult_birth_year_status", ""),
            "adult_identity_exception_type": (record.get("adult_identity_exception") or {}).get("type", ""),
            "observed_at": record["observed_at"], "time_confidence": record["time_confidence"],
            "date_basis": record.get("date_basis"), "identity_rules": ";".join(record.get("identity_rules", []))}


def write_outputs(folder: Path, records: list[dict], cases: list[dict], school_queue: list[dict], schemas: dict,
                  adult_pairs: list[dict] | None = None) -> dict:
    schemas = {role: cols + ["school_type", "school_sex", "school_district", "school_religion"] for role, cols in schemas.items()}
    selected = [r for r in records if r["selected_wave"]]
    cells = {(r["child_id"], r["role"], r["selected_wave"]): r for r in selected}
    by_id = {r["submission_id"]: r for r in records}
    bd_by_child = {p["child_id"]: p for p in adult_pairs or []}
    bd_ids = {p[w] for p in adult_pairs or [] for w in ["t0", "t1"]}
    scored_ids = {r["submission_id"] for r in selected} | bd_ids
    scores = {sid: metrics(by_id[sid]["answer"], by_id[sid]["role"]) for sid in scored_ids}
    if len(cells) != len(selected):
        raise ValueError("Multiple selected submissions for one child/role/wave")
    summary = {"submissions": len(records), "status_counts": dict(Counter(r["status"] for r in records)),
               "case_status_counts": dict(Counter(c["status"] for c in cases)),
               "case_rule_counts": dict(Counter(c["rule"] for c in cases)),
               "school_names_pending": len(school_queue), "time_confidence_counts": dict(Counter(r["time_confidence"] for r in records))}
    summary["review_warning_cases"] = sum(bool(c.get("review_warnings")) for c in cases)
    summary["review_warning_code_counts"] = dict(Counter(code for c in cases for code in c.get("review_warning_codes", [])))
    summary["BD_general_contact_fallback"] = sum(p["rule"] == "same_adult_reported_birth_relationship_general_contact_fallback" for p in adult_pairs or [])
    names = {("student", "T0"): "A_T0-student_final", ("parent", "T0"): "B_T0-parent_final",
             ("student", "T1"): "C_T1-student_final", ("parent", "T1"): "D_T1-parent_final"}
    for (role, wave), name in names.items():
        rows = [export_row(r) for r in selected if r["role"] == role and r["selected_wave"] == wave]
        columns = ["same_person_entry_id", "same_person_match_rule"] + schemas[role] + ["submission_id", "child_id", "source_claimed_timepoint", "final_timepoint", "wave_reason", "observed_at", "time_confidence", "date_basis", "identity_rules", "adult_birth_year_status", "adult_identity_exception_type"]
        write_excel(folder / "3.1.1_ABCD_reconstruction" / (name + ".xlsx"), {"Sheet1": pd.DataFrame(rows, columns=columns)})
        summary[name[0]] = len(rows)
    longitudinal = set()
    dyads = set()
    dyad_rows = {}
    for role, label in [("student", "AC_T0T1-student"), ("parent", "BD_T0T1-parent")]:
        rows = []
        for child in sorted({r["child_id"] for r in selected}):
            a, b = cells.get((child, role, "T0")), cells.get((child, role, "T1"))
            if role == "parent" and adult_pairs is not None:
                pair = bd_by_child.get(child)
                a, b = (by_id[pair["t0"]], by_id[pair["t1"]]) if pair else (None, None)
            if a and b:
                longitudinal.add((child, role))
                def long_row(record, wave):
                    view = {**record, "selected_wave": wave}
                    if role == "parent" and adult_pairs is not None:
                        view["wave_reason"] = "baseline_caregiver_latest_same_adult_followup" if wave == "T1" else record.get("wave_reason")
                    return export_row(view)
                rows.append({"child_id": child,
                             **({"respondent_id": pair["respondent_id"], "adult_match_rule": pair["rule"],
                                 "T1_family_selected_wave": b["selected_wave"]} if role == "parent" and adult_pairs is not None else {}),
                             **{"T0_" + k: v for k, v in long_row(a, "T0").items()}, **{"T1_" + k: v for k, v in long_row(b, "T1").items()}})
        columns = ["child_id"] + [w + "_" + k for w in ["T0", "T1"] for k in (["same_person_entry_id", "same_person_match_rule"] + schemas[role] + ["submission_id", "child_id", "source_claimed_timepoint", "final_timepoint", "wave_reason", "observed_at", "time_confidence"])]
        if role == "parent" and adult_pairs is not None:
            columns = ["child_id", "respondent_id", "adult_match_rule", "T1_family_selected_wave"] + columns[1:]
        columns += [w + "_" + k for w in ["T0", "T1"] for k in ["date_basis", "identity_rules", "adult_birth_year_status", "adult_identity_exception_type"]]
        write_excel(folder / "3.1.1_ABCD_reconstruction" / (label + ".xlsx"), {"Sheet1": pd.DataFrame(rows, columns=columns)})
        summary[label[:2]] = len(rows)
    for wave, label in [("T0", "E_T0-student&parent-paired"), ("T1", "F_T1-student&parent-paired")]:
        paired, student_only, parent_only = [], [], []
        for child in sorted({r["child_id"] for r in selected}):
            s, p = cells.get((child, "student", wave)), cells.get((child, "parent", wave))
            if s and p:
                dyads.add((child, wave))
                row = {"child_id": child, "student_same_person_entry_id": export_row(s)["same_person_entry_id"],
                       "parent_same_person_entry_id": export_row(p)["same_person_entry_id"],
                       "student_submission_id": s["submission_id"], "parent_submission_id": p["submission_id"]}
                st,pt=timestamp(s.get("observed_at")),timestamp(p.get("observed_at"))
                row.update(student_observed_at=s.get("observed_at"),parent_observed_at=p.get("observed_at"),
                           student_time_confidence=s["time_confidence"],parent_time_confidence=p["time_confidence"],
                           student_parent_date_gap_days=abs((st-pt).total_seconds())/86400 if st and pt else None)
                # Preserve scale prefixes expected by legacy descriptive scripts.
                for role_record in [s, p]:
                    for k, v in role_record["answer"].items():
                        row[k if k.startswith(("student_", "parent_", "adult_")) else role_record["role"] + "_" + k] = v
                paired.append(row)
                dyad_rows[(child, wave)] = row
            elif s:
                student_only.append(export_row(s))
            elif p:
                parent_only.append(export_row(p))
        # Paired outputs live in corrected-stage folder; no administrative files touched.
        paired_cols = ["child_id", "student_same_person_entry_id", "parent_same_person_entry_id", "student_submission_id", "parent_submission_id"]
        if paired:
            paired_cols = list(paired[0])
        write_excel(folder / "3.1.3_manual_pairing" / (label + "_final.xlsx"), {
            "paired": pd.DataFrame(paired, columns=paired_cols),
            "student_only": pd.DataFrame(student_only, columns=["same_person_entry_id", "same_person_match_rule"] + schemas["student"] + ["submission_id", "child_id"]),
            "parent_only": pd.DataFrame(parent_only, columns=["same_person_entry_id", "same_person_match_rule"] + schemas["parent"] + ["submission_id", "child_id"])})
        summary[label[0]] = len(paired)
    ef = []
    adult_cases = {tuple(sorted([c["record_a"], c["record_b"]])): c
                   for c in cases if c["kind"] == "respondent" and c["record_b"]}
    for child in sorted({c for c, w in dyads}):
        a, b = dyad_rows.get((child, "T0")), dyad_rows.get((child, "T1"))
        if a and b:
            p0, p1 = cells[(child, "parent", "T0")], cells[(child, "parent", "T1")]
            relationship0 = text(p0["answer"].get("adult_demographics_06"))
            relationship1 = text(p1["answer"].get("adult_demographics_06"))
            birth0 = text(p0["answer"].get("adult_demographics_05"))
            birth1 = text(p1["answer"].get("adult_demographics_05"))
            caregiver = "reported_relationship_changed" if relationship0 and relationship1 and relationship0 != relationship1 else (
                "reported_birth_year_changed" if birth0 and birth1 and birth0 != birth1 else "identity_unconfirmed")
            if adult_pairs is not None:
                from .respondents import adult_match
                reviewed = adult_cases.get(tuple(sorted([p0["submission_id"], p1["submission_id"]])))
                status, rule = (reviewed["status"], reviewed["rule"]) if reviewed else adult_match(p0, p1)
                caregiver = ("same_adult_by_reviewed_evidence" if rule == "reviewer_confirmed_same_adult"
                             else "same_adult_by_accepted_evidence") if status == "confirmed" else rule
            ef.append({"child_id": child, "caregiver_status": caregiver,
                       **{"T0_" + k: v for k, v in a.items()}, **{"T1_" + k: v for k, v in b.items()}})
    write_excel(folder / "3.1.3_manual_pairing/EF_T0T1-student&parent-paired_final.xlsx", {"paired": pd.DataFrame(ef, columns=list(ef[0]) if ef else ["child_id", "caregiver_status"])})
    summary["EF"] = len(ef)
    analysis = []
    ef_children = {r["child_id"] for r in ef}
    child_longitudinal = {(child, role) for child, role, wave in cells
                          if (child, role, "T0") in cells and (child, role, "T1") in cells}
    def analysis_row(r, wave=None):
        role = r["role"]
        demo = "student_demographics_" if role == "student" else "adult_demographics_"
        dob, time = timestamp(r.get("dob")), timestamp(r.get("observed_at"))
        age = (time - dob).days / 365.2425 if dob and time and time >= dob and r["time_confidence"] in {"observed", "manual_verified"} else None
        school_type = text(r["answer"].get("school_type")).lower()
        mainstream = school_type in {"primary", "secondary", "primary & secondary", "primary school", "secondary school", "through-train school"}
        label = r.get("source_label", text(r["answer"].get("Survey")))
        return {"submission_id": r["submission_id"], "child_id": r["child_id"], "role": role,
                         "wave": wave or r["selected_wave"], "source": r["source_id"], "school": r.get("school"),
                         "source_label": label, "source_type": source_category(label),
                         "identity_rules": ";".join(r.get("identity_rules", [])),
                         "date_basis": r.get("date_basis"), "respondent_id": r.get("respondent_id"),
                         "school_type": school_type, "gender": r.get("gender"), "grade": r.get("grade"),
                         "age_at_observation": age, "time_confidence": r["time_confidence"],
                         "mainstream": mainstream, "longitudinal": r["submission_id"] in bd_ids if role == "parent" and adult_pairs is not None else (r["child_id"], role) in longitudinal,
                         "child_report_longitudinal": (r["child_id"], role) in child_longitudinal,
                         "dyad": (r["child_id"], r["selected_wave"]) in dyads,
                         "longitudinal_dyad": r["child_id"] in ef_children,
                         **scores[r["submission_id"]]}
    analysis = [analysis_row(r) for r in selected]
    summary["analysis_mainstream_rows"] = sum(r["mainstream"] for r in analysis)
    summary["analysis_other_or_unknown_rows"] = len(analysis) - summary["analysis_mainstream_rows"]
    analysis_columns = ["submission_id", "child_id", "role", "wave", "source", "source_type", "source_label", "school", "school_type", "gender", "grade", "age_at_observation", "time_confidence", "date_basis", "identity_rules", "respondent_id", "mainstream", "longitudinal", "child_report_longitudinal", "dyad", "longitudinal_dyad"] + list(metrics({}, "student"))
    pd.DataFrame(analysis, columns=analysis_columns).to_parquet(folder / "analysis.parquet", index=False)
    if adult_pairs is not None:
        bd_analysis = [analysis_row(by_id[p[w]], wave) for p in adult_pairs for w, wave in [("t0", "T0"), ("t1", "T1")]]
        pd.DataFrame(bd_analysis, columns=analysis_columns).to_parquet(folder / "parent_longitudinal.parquet", index=False)
        summary["BD_alternative_followups"] = sum(by_id[p["t1"]]["selected_wave"] != "T1" for p in adult_pairs)
        summary["parent_child_longitudinal"] = sum((child, "parent", "T0") in cells and (child, "parent", "T1") in cells for child in {r["child_id"] for r in selected})
        summary["adult_case_status_counts"] = dict(Counter(c["status"] for c in cases if c["kind"] == "respondent"))
    # Safe evidence view: no full birthday, phone, email, answer JSON or adult
    # birth year is exposed in the ordinary dashboard.
    from .matching import one_digit_phone
    from .respondents import adult_evidence
    evidence = []
    for c in cases:
        a, b = by_id.get(c["record_a"], {}), by_id.get(c["record_b"], {})
        ca, cb = contacts(a.get("answer", {})), contacts(b.get("answer", {}))
        ea, eb = adult_evidence(a), adult_evidence(b)
        row = {k: c[k] for k in ["case_id", "kind", "rule", "status", "reason", "record_a", "record_b"]}
        row.update(role_a=a.get("role"), role_b=b.get("role"), school_a=a.get("school"), school_b=b.get("school"),
                   grade_a=a.get("grade"), grade_b=b.get("grade"),
                   observed_at_a=a.get("observed_at"), observed_at_b=b.get("observed_at"),
                   date_basis_a=a.get("date_basis"), date_basis_b=b.get("date_basis"),
                   school_agrees=bool(a.get("school") and a.get("school") == b.get("school")),
                   dob_agrees=bool(a.get("dob") and a.get("dob") == b.get("dob")),
                   gender_agrees=bool(a.get("gender") and a.get("gender") == b.get("gender")),
                   exact_contact=bool(ca & cb), one_digit_phone=any(one_digit_phone(x, y) for x in ca for y in cb),
                   adult_birthyear_agrees=bool(ea["birth_year"] and ea["birth_year"] == eb["birth_year"]),
                   adult_relationship_agrees=bool(ea["relationship"] and ea["relationship"] == eb["relationship"]),
                   adult_personal_contact_agrees=bool(set(ea["personal_contacts"]) & set(eb["personal_contacts"])))
        row.update(review_warnings="；".join(c.get("review_warnings", [])),
                   review_warning_codes=";".join(c.get("review_warning_codes", [])))
        evidence.append(row)
    pd.DataFrame(evidence).to_parquet(folder / "candidate_evidence.parquet", index=False)
    summary["case_kind_status_counts"] = {kind: dict(Counter(c["status"] for c in cases if c["kind"] == kind)) for kind in sorted({c["kind"] for c in cases})}
    with (folder / "records.jsonl").open("w") as f:
        for r in records:
            f.write(dumps({k: v for k, v in r.items() if k != "contacts"}) + "\n")
    with (folder / "cases.jsonl").open("w") as f:
        for c in cases:
            f.write(dumps(c) + "\n")
    write_excel(folder / "audit.xlsx", {
        "counts": [{"metric": k, "value": dumps(v) if isinstance(v, dict) else v} for k, v in summary.items()],
        "school_review": pd.DataFrame(school_queue, columns=["normalized_key", "raw_name"]),
        "processing_decisions": [{"submission_id": r["submission_id"], "status": r["status"],
                                  "duplicate_of": r.get("duplicate_of"), "selected_wave": r["selected_wave"],
                                  "claimed_wave": r["claimed_wave"], "wave_reason": r.get("wave_reason"),
                                  "time_confidence": r["time_confidence"], "issues": ";".join(r["issues"])} for r in records]})
    return summary


def descriptive_summary(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    metrics_cols = [c for c in frame if c.startswith(("PHQ_", "FQOL_", "SEC_", "TBV_", "TBP_", "CBV_", "CBP_"))]
    for name in metrics_cols:
        values = frame[name].dropna()
        rows.append({"metric": name, "sample_n": len(frame), "valid_n": len(values),
                     "missing_n": len(frame) - len(values), "mean_or_proportion": values.mean() if len(values) else None,
                     "sd": values.std() if len(values) > 1 and "prevalence" not in name else None})
    return pd.DataFrame(rows)
