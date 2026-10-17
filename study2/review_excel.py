"""Native, macro-free XLSX adapter. No identity decisions are made here.

The template's styles are authored with Artifact Tool. Runtime population uses
the standard library's OpenXML writer; openpyxl is only used to read submissions.
"""
from __future__ import annotations

import copy
import json
import math
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from .common import PipelineError, digest, text
from .review_core import EVIDENCE_SOURCES, KINDS, VERDICTS, submit_reviews, workspace

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)
STATUS = {'confirmed':'已確認','pending':'待審','rejected':'已拒絕','unmatched':'未配對'}
RULES = {'exact_school_contact_child_demo':'同校、孩子生日及性別一致，聯絡資料精確相同',
         'child_contact_bridge_component':'直接規則連通整組、孩子生日及性別一致，自動確認',
         'child_group_contact_bridge':'群組聯絡連結核實，部分提交沒有直接共同聯絡',
         'child_group_contact_gap':'群組尚有未連起的子群，需要核實缺少證據的連結',
         'one_digit_phone_school_child_demo':'同校、孩子生日及性別一致，電話只差一位數字',
         'exact_contact_school_transfer_child_demo':'學校不同，孩子生日、性別及聯絡資料精確相同',
         'fuzzy_contact_child_demo':'電話近似，孩子身份證據需要核實',
         'school_child_demo':'同校及孩子基本資料一致，其他證據仍不足',
         'transfer_contact_child_demo':'跨校聯絡候選，孩子身份需要核實',
         'contact_missing_or_conflicting_child_demo':'共同聯絡，但孩子資料缺失或矛盾',
         'same_adult_reported_birth_relationship_personal_contact':'成人出生年份、關係及本人聯絡一致',
         'same_adult_reported_birth_relationship_general_contact_fallback':'本人聯絡缺失；成人年份、關係及其他精確聯絡一致',
         'same_adult_reported_birth_personal_contact':'成人出生年份及任一本人電話／電郵一致',
         'same_adult_reported_birth_any_exact_contact':'成人出生年份及任一聯絡欄精確電話／電郵一致',
         'no_exact_adult_contact':'沒有任何精確相同的電話／電郵',
         'different_child_identity':'兩份回應指向不同孩子',
         'adult_birth_year_missing_or_invalid':'成人出生年份缺失或不合法',
         'reported_birth_year_changed':'成人出生年份不同',
         'reported_relationship_changed':'與孩子關係不同',
         'adult_relationship_missing_or_invalid':'與孩子關係缺失或不合法',
         'adult_other_relationship_unspecified':'其他關係沒有說明',
         'reported_other_relationship_changed':'其他關係文字不同',
         'no_exact_personal_adult_contact':'沒有精確相同的成人本人聯絡',
         'manual_link':'已有人工孩子連結', 'reviewer_confirmed_same_adult':'已有人工成人連結',
         'reviewer_rejected_same_adult':'已有人工拒絕成人連結'}
def tag(name):
    return "{" + NS + "}" + name


class Sheet:
    def __init__(self, styles):
        self.styles = styles
        self.root = ET.Element(tag("worksheet"))
        views = ET.SubElement(self.root, tag("sheetViews"))
        view = ET.SubElement(views, tag("sheetView"), workbookViewId="0", showGridLines="0")
        ET.SubElement(view, tag("pane"), ySplit="3", topLeftCell="A4", activePane="bottomLeft", state="frozen")
        ET.SubElement(self.root, tag("sheetFormatPr"), defaultRowHeight="30")
        cols = ET.SubElement(self.root, tag("cols"))
        for n, width in enumerate((23, 44, 44, 22), 1):
            ET.SubElement(cols, tag("col"), min=str(n), max=str(n), width=str(width), customWidth="1")
        self.data = ET.SubElement(self.root, tag("sheetData"))
        self.rows, self.merges, self.validations, self.links = {}, [], [], []

    def row(self, number, height=30):
        if number not in self.rows:
            self.rows[number] = ET.SubElement(self.data, tag("row"), r=str(number), ht=str(height), customHeight="1")
        return self.rows[number]

    def cell(self, addr, value=None, style=4, formula=False, height=30):
        number = int("".join(c for c in addr if c.isdigit()))
        c = ET.SubElement(self.row(number, height), tag("c"), r=addr, s=str(self.styles[style]))
        if formula:
            ET.SubElement(c, tag("f")).text = value.lstrip("=")
            c.set("t", "str")
            ET.SubElement(c, tag("v")).text = "未填寫"  # Excel recalculates on open/edit.
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            ET.SubElement(c, tag("v")).text = str(value)
        elif value is not None:
            c.set("t", "inlineStr")
            t = ET.SubElement(ET.SubElement(c, tag("is")), tag("t"))
            t.text = str(value)  # Literal text, never formula injection.
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    def line(self, row, values, style=4, height=30):
        for col, value in zip("ABCD", values):
            self.cell(col + str(row), value, style, height=height)

    def merged(self, row, value, first="A", last="D", style=4, height=30, formula=False):
        for col in "ABCD"["ABCD".index(first):"ABCD".index(last) + 1]:
            self.cell(col + str(row), value if col == first else None, style, formula and col == first, height)
        self.merges.append(f"{first}{row}:{last}{row}")

    def hyperlink(self, addr, location):
        self.links.append((addr, location))

    def finish(self):
        # OOXML element order matters for Excel desktop.
        ET.SubElement(self.root, tag("sheetProtection"), sheet="1", objects="1", scenarios="1",
                      selectLockedCells="0", selectUnlockedCells="0", formatCells="1", insertRows="1", deleteRows="1")
        merges = ET.SubElement(self.root, tag("mergeCells"), count=str(len(self.merges)))
        for ref in self.merges:
            ET.SubElement(merges, tag("mergeCell"), ref=ref)
        if self.validations:
            dvs = ET.SubElement(self.root, tag("dataValidations"), count=str(len(self.validations)))
            for addr, values in self.validations:
                dv = ET.SubElement(dvs, tag("dataValidation"), type="list", sqref=addr, allowBlank="1",
                                   showErrorMessage="1", showInputMessage="1", errorStyle="stop",
                                   errorTitle="請選擇下拉選項", error="請勿自行輸入其他值", promptTitle="填寫答案",
                                   prompt="選擇一項；補充理由另填")
                ET.SubElement(dv, tag("formula1")).text = '"' + ",".join(values) + '"'
        if self.links:
            links = ET.SubElement(self.root, tag("hyperlinks"))
            for ref, location in self.links:
                ET.SubElement(links, tag("hyperlink"), ref=ref, location=location)
        ET.SubElement(self.root, tag("pageMargins"), left="0.25", right="0.25", top="0.4", bottom="0.4", header="0.2", footer="0.2")
        ET.SubElement(self.root, tag("pageSetup"), paperSize="9", orientation="landscape", fitToWidth="1", fitToHeight="0")
        return ET.tostring(self.root, encoding="utf-8", xml_declaration=True)


def shown(value):
    if value is None or value == "" or value == []:
        return "缺失"
    if isinstance(value, list):
        return "\n".join(map(str, value))
    if isinstance(value, dict):
        return "；".join(f"{k}: {v}" for k, v in value.items())
    return str(value)


def text_height(*values, minimum=30):
    lines=max(sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(c) in {'W','F'} else 1
                  for c in line)/42)) for line in str(v).split('\n')) for v in values)
    return min(409,max(minimum,lines*16+8))


def comparison(a, b):
    missing_a, missing_b = shown(a) == "缺失", shown(b) == "缺失"
    if missing_a or missing_b:
        return "兩方缺失" if missing_a and missing_b else "一方缺失"
    if a == b:
        return '相同'
    if isinstance(a, list) and isinstance(b, list) and set(a) & set(b):
        return '部分相同（有共同聯絡）'
    return '不同'


def fields(r, adult=False):
    e, own = r["child_evidence"], r["adult_evidence"]
    loc = r.get("location") or {}
    result = [("問卷對象", "學生" if r["role"] == "student" else "家長"), ("學校", e.get("school")),
              ("孩子生日", e.get("dob")), ("孩子性別（原編碼）", e.get("gender")), ("年級（清理值）", e.get("grade")),
              ("年級（原始值）", r.get('grade_raw')),
              ("現有班號／學號", e.get("roll")), ("聯絡資料", own.get('all_contacts', e.get('contacts')) if adult else e.get('contacts'))]
    if adult:
        result.extend([("成人出生年份", own["birth_year"]), ("與孩子關係", {"0":"父親 (0)", "1":"母親 (1)", "2":"其他 (2)"}.get(own["relationship"], own["relationship"])),
                       ("其他關係文字", own["other_relationship"]), ("本人聯絡資料", own["personal_contacts"])])
    result.extend([("來源檔案", loc.get("file") or loc.get("path")), ("來源行號", loc.get("row") or loc.get("original_row")),
                   ("來源類別", r.get("source_label") or r.get("source_id")), ("原始輪次", r.get("claimed_wave")),
                   ("目前選定輪次", r.get("selected_wave")), ("建立時間（原值）", r.get("answer", {}).get("建立時間")),
                   ("目前觀察時間", r.get("observed_at")), ("日期依據", r.get("date_basis")),
                   ("日期可信度", r.get("time_confidence")), ("永久提交 ID", r.get("submission_id"))])
    return result


def write_package(package, path):
    template = Path(__file__).parent / "assets/review_template.xlsx"
    if not template.exists():
        raise PipelineError("審核卡版面範本缺失")
    with zipfile.ZipFile(template) as z:
        files = {n: z.read(n) for n in z.namelist()}
    raw = ET.fromstring(files["xl/worksheets/sheet1.xml"])
    styles = {r: int(raw.find(f'.//{tag("c")}[@r="A{r}"]').get("s", "0")) for r in range(1, 11)}
    style_root = ET.fromstring(files["xl/styles.xml"])
    xfs = style_root.find(tag("cellXfs"))
    # Editable blue-on-yellow style, cloned without changing the authored base.
    editable = copy.deepcopy(xfs[styles[6]])
    editable.set("applyProtection", "1")
    ET.SubElement(editable, tag("protection"), locked="0")
    styles[11] = len(xfs)
    xfs.append(editable)
    xfs.set("count", str(len(xfs)))
    # Colour comparisons by text; both missing is never the same as equal evidence.
    for key, fill in ((12, "FCE7E5"), (13, "FFF1CC")):
        fills = style_root.find(tag("fills"))
        pattern = ET.SubElement(ET.SubElement(fills, tag("fill")), tag("patternFill"), patternType="solid")
        ET.SubElement(pattern, tag("fgColor"), rgb="FF" + fill)
        ET.SubElement(pattern, tag("bgColor"), indexed="64")
        fills.set("count", str(len(fills)))
        xf = copy.deepcopy(xfs[styles[4]])
        xf.set("fillId", str(len(fills)-1))
        xf.set("applyFill", "1")
        styles[key] = len(xfs)
        xfs.append(xf)
    xfs.set("count", str(len(xfs)))
    files["xl/styles.xml"] = ET.tostring(style_root, encoding="utf-8", xml_declaration=True)
    directory, cards, reference, system = [Sheet(styles) for _ in range(4)]
    purpose = package.get('purpose', 'review')
    title = {'review':'逐宗身份審核', 'validation':'身份品質抽查', 'adjudication':'Jarvis 指定裁定'}.get(purpose, '逐宗身份審核')
    has_groups = any(t['snapshot']['case'].get('group_review') for t in package['tasks'])
    directory.merged(1, "Study 2｜" + title, style=1, height=36)
    directory.merged(2, "審核者：" + package["reviewer"] + "　｜　共 " + str(len(package["tasks"])) + " 宗", style=2)
    directory.merged(3, "孩子群組先查看群組參考內全部 R 編號，再填三個黃色格。儲存並交回同一份 .xlsx。" if has_groups else "只填黃色格。填寫、儲存、交回同一份 .xlsx；毋須操作識別碼。", style=2)
    directory.merged(4, "三項必填：結論、主要證據來源、補充理由（包括證據位置）。", style=2)
    directory.merged(5, "資料不足不改目前配對。已填寫只表示答案齊備，仍待 Jarvis 批准。", style=2)
    directory.merged(6, "原始輪次及選定輪次分開列出。親子填卷日期差不作拒配條件。", style=2)
    directory.merged(7, "工作包：" + package["package_id"] + "　匯出時間：" + package["created_at"], style=10, height=44)
    directory.merged(8, "目前發布版本：" + package["release_id"], style=10)
    directory.cell("A9", "完成進度", 3)
    directory.merged(9, '=COUNTIF(D12:D' + str(11+len(package["tasks"])) + ',"已填寫")&" / ' + str(len(package["tasks"])) + ' 宗已填寫"', "B", "D", style=8, formula=True)
    directory.rows[9].find(f'{tag("c")}[@r="B9"]/{tag("v")}').text='0 / '+str(len(package['tasks']))+' 宗已填寫'
    directory.line(11, ["個案", "類型及問題", "審核者", "填寫狀態"], style=3)
    reference.merged(1, "群組參考｜必要的其他提交及程式原判斷", style=1, height=36)
    reference.merged(2, "群組卡須核對本宗全部 R 編號及連結表，再作一次整組判斷。獨立重審不顯示他人答案。" if has_groups else "先比較審核卡，再按需要查閱這裡。獨立重審不顯示他人答案。", style=2)
    reference.merged(3, "資料來源：" + package["release_id"] + " 的標準化提交及候選快照", style=10)
    system.line(1, ["格式版本", "工作包", "審核者", "原發布版本"], style=3)
    system.line(2, [package["format_version"], package["package_id"], package["reviewer"], package["release_id"]])
    system.line(4, ["任務 ID／修訂", "個案 ID", "證據快照", "輸入位置"], style=3)
    inputs, rr = {}, 5
    for index, task in enumerate(package["tasks"]):
        start = 1 + 40*index
        case, snapshot = task["snapshot"]["case"], task["snapshot"]
        a, b = (snapshot["records"][case[k]] for k in ("record_a", "record_b"))
        if case["kind"] == "family":
            a, b = sorted([a,b], key=lambda r: r["role"] != "student")
        else:
            a, b = sorted([a,b], key=lambda r: (r.get("observed_at") or "9999", r["submission_id"]))
        number = index+1
        question = "這兩份問卷是否由同一成人填寫？" if case["kind"] == "respondent" else "這兩份問卷是否指向同一孩子？"
        group = case.get('group_review')
        if group:
            question = f"群組參考內全部 {len(group['members'])} 份問卷，是否指向同一孩子？"
        cards.merged(start, f"個案 {number:02d}｜" + ('孩子身份群組' if group else KINDS[case['kind']]), style=1, height=36)
        cards.merged(start+1, question, style=2)
        cards.merged(start+2, "返回審核目錄", "A", "B", style=9)
        cards.hyperlink(f"A{start+2}", "'審核目錄'!A"+str(11+number))
        cards.merged(start+2, "下一宗" if number < len(package["tasks"]) else "返回目錄", "C", "D", style=9)
        cards.hyperlink(f"C{start+2}", "'審核卡'!A"+str(start+40) if number<len(package["tasks"]) else "'審核目錄'!A1")
        node_labels = {sid: f'R{i+1:03d}' for i, sid in enumerate(sorted(snapshot['records']))}
        cards.line(start+3, ["缺口代表比較" if group else "比較項目",
                            node_labels[a['submission_id']] if group else "學生" if case["kind"]=="family" else "較早提交／左方",
                            node_labels[b['submission_id']] if group else "家長" if case["kind"]=="family" else "較後提交／右方", "比較提示"], style=3)
        for offset, ((label, av), (_, bv)) in enumerate(zip(fields(a,case["kind"]=="respondent"), fields(b,case["kind"]=="respondent"))):
            r = start+4+offset
            comp = comparison(av,bv)
            height = text_height(shown(av),shown(bv),minimum=48 if label in
                                 {"來源檔案","聯絡資料","永久提交 ID","本人聯絡資料","其他關係文字"} else 30)
            cards.line(r, [label, shown(av), shown(bv)], height=height)
            identity_field=label in {'學校','孩子生日','孩子性別（原編碼）','年級（清理值）','年級（原始值）','現有班號／學號','聯絡資料',
                                    '成人出生年份','與孩子關係','其他關係文字','本人聯絡資料'}
            cards.cell("D"+str(r), comp, (5 if comp=="相同" else 12 if comp=="不同" else 13) if identity_field else 8, height=height)
        for blank in range(start+4+len(fields(a,case['kind']=='respondent')),start+27):
            if group and blank >= start+25:
                continue
            cards.row(blank,height=8)
        if group:
            message = (f"全組 {len(group['members'])} 份；{len(group['passed_pairs'])} 對已通過直接規則，毋須重審。\n"
                       + case['reason'] + '。先開啟群組參考；上方兩份只代表證據缺口。')
            if 'unresolved_pairs' in group:
                message += f"\n只核實子群之間的連結；{len(group['unresolved_pairs'])} 對尚未確認同一孩子。"
            cards.merged(start+25, message, style=2, height=text_height(message, minimum=66))
            cards.merged(start+26, '開啟全組提交及待判斷連結', style=9)
            cards.hyperlink('A'+str(start+26), "'群組參考'!A"+str(rr))
        cards.cell("A"+str(start+27), "審核結論", 3)
        cards.merged(start+27, None, "B", "D", style=11)
        cards.validations.append(("B"+str(start+27), list(VERDICTS)))
        cards.cell("A"+str(start+28), "主要證據來源", 3)
        cards.merged(start+28, None, "B", "D", style=11)
        cards.validations.append(("B"+str(start+28), EVIDENCE_SOURCES))
        cards.cell("A"+str(start+29), "補充理由（必填）", 3, height=108)
        cards.merged(start+29, None, "B", "D", style=11, height=108)
        v, e, reason = ("B"+str(start+n) for n in (27,28,29))
        formula = f'=IF(AND({v}<>"",{e}<>"",LEN(TRIM({reason}))>0),"已填寫",IF(AND({v}="",{e}="",LEN(TRIM({reason}))=0),"未填寫","尚欠必填項目"))'
        cards.cell("A"+str(start+31), "填寫狀態", 3)
        cards.merged(start+31, formula, "B", "D", style=8, formula=True)
        warnings = case.get("review_warnings", [])
        verification = (group or {}).get('verification_source') or snapshot.get('verification_source')
        if verification:
            warnings = list(warnings) + ['本次核實問題：'+verification['question'],
                '已登記來源：'+verification['source']+'；位置：'+verification['reference']]
        if warnings:
            message = "審核提示（不代表拒配）：\n" + "\n".join(warnings)
            cards.merged(start+32, message, style=13, height=text_height(message, minimum=48))
        cards.merged(start+33, "查閱相關提交、候選及程式原判斷", style=9)
        cards.hyperlink("A"+str(start+33), "'群組參考'!A"+str(rr))
        cards.merged(start+34, "資料有錯請在理由註明，交 Jarvis 使用既有修正工具處理。", style=10)
        cards.merged(start+35, "已填寫並不代表已批准。請勿移動、刪除或修改卡片。", style=10)
        inputs[task["task_id"]] = {"verdict": v, "evidence_source": e, "reason": reason}
        directory.line(11+number, [f"開啟個案 {number:02d}", '孩子身份群組' if group else KINDS[case["kind"]], package["reviewer"]], style=9)
        directory.cell("D"+str(11+number), f"='審核卡'!B{start+31}", style=8, formula=True)
        directory.hyperlink("A"+str(11+number), "'審核卡'!A"+str(start))
        system.line(5+index*2, [task["task_id"]+" / "+str(task['revision']), case["case_id"], snapshot["signature"], ", ".join((v,e,reason))],height=60)
        system.line(6+index*2, [case['record_a'],case['record_b'],case['version_a'],case['version_b']],height=60)
        reference.merged(rr, f"個案 {number:02d}｜{case['case_id']}", style=3, height=36)
        reference.merged(rr+1, f"程式原判斷：{STATUS.get(case['status'],case['status'])}　規則：{RULES.get(case['rule'],case['rule'])}\n規則代號：{case['rule']}", style=10, height=60)
        source_note = '程式理由：'+case.get('reason','')
        if verification:
            source_note += '\n核實問題：'+verification['question']+'\n核實來源：'+verification['source']+'；位置：'+verification['reference']
        reference.merged(rr+2, source_note, style=10, height=text_height(source_note,minimum=50))
        node_names = {sid: f'R{i+1:03d}' for i, sid in enumerate(sorted(snapshot['records']))}
        reference.line(rr+3, ["相關提交 ID", "身份資料", "輪次／來源位置", "群組／狀態"], style=3)
        rr += 4
        for sid, rec in snapshot["records"].items():
            loc=rec.get("location") or {}
            identity = ("學校："+shown(rec.get("school"))+"\n生日："+shown(rec.get("dob"))+"；性別："+shown(rec.get("gender"))
                        + ('\n'+('學生' if rec['role']=='student' else '家長')+'；年級：'+shown(rec.get('grade')) if group else '')
                        +"\n聯絡："+shown(rec["child_evidence"]["contacts"]))
            provenance = ("原始："+shown(rec.get("claimed_wave"))+"；選定："+shown(rec.get("selected_wave"))
                          + ('\n建立時間：'+shown(rec.get('observed_at'))+'\n日期依據：'+shown(rec.get('date_basis')) if group else '')
                          +"\n"+shown(loc))
            reference.line(rr, [(node_names[sid]+'\n' if group else '')+sid, identity, provenance,
                           shown(rec.get("child_id"))+"\n"+shown(rec.get("status"))],
                           height=text_height(identity,provenance,minimum=90) if group else 90)
            rr += 1
        if group:
            for i, component in enumerate(group.get('contact_components', []), 1):
                child_ids = {snapshot['records'][sid].get('child_id') for sid in component}
                status = ('孤立提交，沒有組內聯絡連結' if len(component) == 1 else
                          '已確認子群，毋須重審內部連結' if len(child_ids) == 1 and None not in child_ids else
                          '聯絡連通，須保留既有身份限制')
                reference.line(rr, [f'SG{i:03d}', '、'.join(node_names[sid] for sid in component),
                                    status, '只核實與其他子群的連結'], style=2, height=60)
                rr += 1
            reference.merged(rr, '本宗判斷範圍：全部 R 編號。確認不同請在理由列出分組，例如 R001、R003 一組；R002 另一組。', style=2, height=55)
            rr += 1
            reference.line(rr, ['提交一','提交二','連結證據','需要判斷甚麼'], style=3)
            rr += 1
            for pair in group['passed_pairs'] + group['missing_pairs']:
                passed = bool(pair['rule'])
                left, right = (snapshot['records'][pair[side]] for side in ('a','b'))
                confirmed = bool(left.get('child_id') and left['child_id'] == right.get('child_id'))
                reference.line(rr, [node_names[pair['a']],node_names[pair['b']],
                                    RULES.get(pair['rule'], pair['rule']) if passed else
                                    '已確認子群內的間接連結' if confirmed else '沒有直接符合聯絡規則的證據',
                                    '已確認，毋須逐對重審' if confirmed else '已通過直接規則，毋須逐對重審' if passed else
                                    '只核實跨子群連結；欠證據不等於已確認不同孩子'],
                               style=5 if passed or confirmed else 13,height=60)
                rr += 1
        reference.merged(rr, "相關候選（永久 ID 不因左右顯示次序改變）", style=3)
        rr += 1
        for c in snapshot["candidates"]:
            reference.line(rr, [c["case_id"], c["record_a"]+"\n"+c["record_b"], c["kind"]+" / "+c["rule"], c["status"]], height=60)
            rr += 1
        rr += 2
    for n, sheet in enumerate((directory,cards,reference,system),1):
        files[f"xl/worksheets/sheet{n}.xml"] = sheet.finish()
        # Remove template relationships when replacing its sheets.
        files.pop(f"xl/worksheets/_rels/sheet{n}.xml.rels",None)
    book=ET.fromstring(files["xl/workbook.xml"])
    for sheet in book.find(tag("sheets")):
        sheet.set("state", "veryHidden" if sheet.get("name")=="系統資料" else "visible")
    calc = book.find(tag("calcPr"))
    if calc is None:
        calc = ET.SubElement(book,tag("calcPr"))
    calc.set("fullCalcOnLoad","1")
    calc.set("forceFullCalc","1")
    calc.set("calcMode","auto")
    files["xl/workbook.xml"] = ET.tostring(book,encoding="utf-8",xml_declaration=True)
    path.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
        for name,data in files.items():
            z.writestr(name,data)
    path.chmod(0o600)
    return {"inputs": inputs, "immutable": workbook_contract(path,inputs)}


def workbook_contract(path, inputs):
    wb=load_workbook(path, data_only=False)
    try:
        allowed={cell for positions in inputs.values() for cell in positions.values()}
        result={"sheets":wb.sheetnames,"layout":{}}
        for sheet in wb:
            contents={}
            for row in sheet:
                for c in row:
                    if sheet.title=="審核卡" and c.coordinate in allowed:
                        if c.protection.locked:
                            raise PipelineError("答案格被鎖定："+c.coordinate)
                        continue
                    if c.value is not None or c.hyperlink:
                        contents[c.coordinate] = [c.value,c.data_type,
                            c.hyperlink.location if c.hyperlink else None,
                            c.hyperlink.target if c.hyperlink else None,c.protection.locked]
            result["layout"][sheet.title] = {"contents":contents,"merges":sorted(str(m) for m in sheet.merged_cells.ranges),
               "state":sheet.sheet_state,"protected":sheet.protection.sheet,
               "hidden_rows":[r for r,d in sheet.row_dimensions.items() if d.hidden],
               "hidden_columns":[c for c,d in sheet.column_dimensions.items() if d.hidden],
               "validation":sorted((str(v.sqref),v.type,v.formula1) for v in sheet.data_validations.dataValidation)}
        return result
    finally:
        wb.close()


def read_package(cfg,path):
    path=Path(path)
    try:
        wb=load_workbook(path,data_only=False)
        if "系統資料" not in wb.sheetnames:
            raise PipelineError("這不是新格式審核包，請使用原有平面工具")
        pid=text(wb["系統資料"]["B2"].value)
        work=workspace(cfg)
        try:
            row=work.execute("SELECT payload FROM packages WHERE id=?",(pid,)).fetchone()
        finally:
            work.close()
        if not row:
            raise PipelineError("未登記工作包或工作包 ID 被修改")
        package=json.loads(row[0])
        contract=package.get("contract")
        if not contract or digest(workbook_contract(path,contract["inputs"])) != digest(contract["immutable"]):
            raise PipelineError("工作包的識別碼、證據、公式、卡片結構或鎖定已被修改；請使用原包重填")
        answers=[]
        for tid,positions in contract["inputs"].items():
            answer={"task_id":tid}
            for key,addr in positions.items():
                c=wb["審核卡"][addr]
                if c.data_type=="f" or c.hyperlink:
                    raise PipelineError("答案不可填入公式或超連結："+addr)
                answer[key]=text(c.value)
            answers.append(answer)
        return pid,answers
    except (KeyError,ValueError,zipfile.BadZipFile) as error:
        raise PipelineError("工作包無法讀取："+str(error)) from error
    finally:
        if 'wb' in locals():
            wb.close()


def import_package(cfg,path):
    pid,answers=read_package(cfg,path)
    try:
        return submit_reviews(cfg,pid,answers,source=str(Path(path).resolve()))
    except PipelineError as error:
        work=workspace(cfg)
        try:
            package=json.loads(work.execute('SELECT payload FROM packages WHERE id=?',(pid,)).fetchone()[0])
        finally:
            work.close()
        message=str(error)
        for number,task in enumerate(package['tasks'],1):
            message=message.replace(task['task_id'],f'個案 {number:02d}')
        raise PipelineError(message) from error
