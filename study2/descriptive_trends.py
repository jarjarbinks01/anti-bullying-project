"""Calendar-response statistics over immutable, already-selected release rows.

No identity, wave selection or scoring is performed here. Date exclusions form
an ordered waterfall; aggregate outputs never contain submission-level rows.
"""
from datetime import date, datetime, timedelta
import io
import json

import pandas as pd

from .common import PipelineError
from .descriptive_spec import DICTIONARY_VERSION, VERSION as REPORT_VERSION, variables
from .descriptives import statistics

VERSION = 'study2_calendar_response_v1'
TIMEZONE = 'Asia/Hong_Kong'
BASIS = 'calendar_response'
REQUIRED = {'submission_id', 'version_id', 'wave', 'observed_at', 'date_basis',
            'time_confidence', 'report_mode', 'grouping_anchor'}
DATE_LABELS = {
    'system_creation': '系統記錄日期', 'manual_verified': '人工核實日期',
    'batch_assigned_creation': '批次指定日期', 'unknown_basis': '日期來源未能判定',
    'invalid_date': '缺失或無效日期',
}
COMPOSITION = {'school': '該輪學校', 'source_type': '該輪來源',
               'dem_child_gender': '該輪孩子性別', 'school_stage': '該輪學校階段'}


def dated_rows(rows):
    """Validate release capability, then derive HK dates on a private copy.

    Naive timestamps retain the pipeline's UTC convention. Numeric values are
    never guessed to be Unix timestamps. A SET_DATE manual verification may
    retain its original batch date_basis, so confidence takes precedence.
    """
    missing = sorted(REQUIRED - set(rows.columns))
    if missing:
        raise PipelineError('此發布未提供時段趨勢所需欄位：' + '、'.join(missing))
    if rows['submission_id'].isna().any() or rows['submission_id'].duplicated().any():
        raise PipelineError('時段趨勢的正式提交ID缺失或重複，不能計數。')
    if rows['version_id'].isna().any() or not rows['wave'].isin(['T0', 'T1']).all():
        raise PipelineError('時段趨勢的來源版本或正式輪次不完整。')
    if 'descriptive_dictionary_version' in rows and not rows['descriptive_dictionary_version'].eq(DICTIONARY_VERSION).all():
        raise PipelineError('此發布的描述統計字典不支援目前時段趨勢。')
    result = rows.copy().reset_index(drop=True)
    raw = result['observed_at'].map(
        lambda v: v if isinstance(v, (str, date, datetime, pd.Timestamp)) else None)
    times = pd.to_datetime(raw, format='mixed', utc=True, errors='coerce')
    result['_calendar_date'] = times.dt.tz_convert(TIMEZONE).dt.date
    quality = pd.Series('unknown_basis', index=result.index)
    quality.loc[result['time_confidence'].eq('observed') & result['date_basis'].eq('system_creation')] = 'system_creation'
    quality.loc[result['time_confidence'].eq('accepted_creation_time') & result['date_basis'].eq('batch_assigned_creation')] = 'batch_assigned_creation'
    quality.loc[result['time_confidence'].eq('manual_verified')] = 'manual_verified'
    quality.loc[times.isna()] = 'invalid_date'
    result['_date_quality'] = quality
    return result


def date_mask(rows, include_batch=False):
    allowed = ['system_creation', 'manual_verified']
    if include_batch:
        allowed.append('batch_assigned_creation')
    return rows['_date_quality'].isin(allowed)


def date_extent(rows, *, include_batch=False, waves=('T0', 'T1')):
    dated = dated_rows(rows)
    dates = dated.loc[date_mask(dated, include_batch) & dated['wave'].isin(waves), '_calendar_date']
    return (dates.min(), dates.max()) if len(dates) else None


def _quality_counts(dated, included):
    records = []
    for wave in ('T0', 'T1'):
        for code, label in DATE_LABELS.items():
            match = dated['wave'].eq(wave) & dated['_date_quality'].eq(code)
            records.append(dict(wave=wave, date_quality=code, date_quality_label=label,
                                input_n=int(match.sum()), included_n=int((match & included).sum())))
    return pd.DataFrame(records)


def date_diagnostics(rows):
    """Explain an undated cohort without inventing a diagnostic date window."""
    dated = dated_rows(rows)
    return _quality_counts(dated, pd.Series(False, index=dated.index)).drop(columns='included_n')


def periods(start, end, interval='month'):
    """All calendar bins, including empty and clipped boundary bins."""
    if not isinstance(start, date) or not isinstance(end, date) or start > end:
        raise ValueError('開始日期不可晚於結束日期。')
    if interval not in {'month', 'week'}:
        raise ValueError('時間間隔必須為每月或每週。')
    records = []
    for period in pd.period_range(start, end, freq='M' if interval == 'month' else 'W-SUN'):
        first, last = period.start_time.date(), period.end_time.date()
        lo, hi = max(first, start), min(last, end)
        records.append(dict(period_id=str(first)[:7] if interval == 'month' else first.isoformat(),
            period_start=first.isoformat(), period_end=last.isoformat(),
            covered_start=lo.isoformat(), covered_end=hi.isoformat(),
            partial_period=(lo != first or hi != last)))
    return records


def _period_id(value, interval):
    return value.isoformat()[:7] if interval == 'month' else (value - timedelta(days=value.weekday())).isoformat()


def _summarize(rows, role, release_id, chosen):
    # Reuse the formal report's exact definitions, availability and denominators.
    result = statistics(rows, role, release_id=release_id)
    return result[result['variable_id'].isin(chosen)].copy()


def calendar_trends(rows, role, release_id, *, start, end, interval='month',
                    include_batch=False, waves=('T0', 'T1'), variable_ids=None, metadata=None):
    """Return aggregate-only tables with a single reproducible analysis basis."""
    if role not in {'student', 'parent'} or not release_id:
        raise ValueError('需要填卷對象及發布版本。')
    if not set(waves) <= {'T0', 'T1'}:
        raise ValueError('輪次必須為T0或T1。')
    waves = tuple(w for w in ('T0', 'T1') if w in waves)
    if not waves:
        raise ValueError('至少選擇一個輪次。')
    defined = {v.key for v in variables(role)}
    chosen = sorted(defined if variable_ids is None else set(variable_ids))
    if not chosen or not set(chosen) <= defined:
        raise ValueError('指標不適用於目前填卷對象。')
    bins = periods(start, end, interval)
    dated = dated_rows(rows)
    if 'role' in dated and not dated['role'].eq(role).all():
        raise PipelineError('時段趨勢不可混合學生及家長回應。')
    eligible = date_mask(dated, include_batch)
    in_range = dated['_calendar_date'].map(lambda d: pd.notna(d) and start <= d <= end)
    mask = dated['wave'].isin(waves) & eligible & in_range
    selected = dated.loc[mask].copy()
    selected['_period_id'] = selected['_calendar_date'].map(lambda d: _period_id(d, interval))
    audit = []
    stages = [
        ('現有篩選後', pd.Series(True, index=dated.index)),
        ('輪次選擇後', dated['wave'].isin(waves)),
        ('有效日期及已知來源後', dated['wave'].isin(waves) & ~dated['_date_quality'].isin(['invalid_date', 'unknown_basis'])),
        ('日期來源選擇後', dated['wave'].isin(waves) & eligible),
        ('指定期間內', mask),
    ]
    for wave in ('T0', 'T1'):
        previous = int(dated['wave'].eq(wave).sum())
        for stage, keep in stages:
            n = int((dated['wave'].eq(wave) & keep).sum())
            audit.append(dict(wave=wave, stage=stage, remaining_n=n, excluded_n=previous-n))
            previous = n
    statistics_parts, count_rows, composition = [], [], []
    for period in bins:
        block = selected[selected['_period_id'].eq(period['period_id'])]
        table = _summarize(block, role, release_id, chosen)
        table = table[table['wave'].isin(waves)]
        for key, value in period.items():
            table[key] = value
        statistics_parts.append(table)
        for wave in waves:
            sample = block[block['wave'].eq(wave)]
            count_rows.append(dict(**period, wave=wave, sample_n=len(sample),
                batch_date_n=int(sample['_date_quality'].eq('batch_assigned_creation').sum()),
                manual_date_n=int(sample['_date_quality'].eq('manual_verified').sum())))
            for field, label in COMPOSITION.items():
                values = sample[field] if field in sample else pd.Series(None, index=sample.index, dtype=object)
                codes = values.fillna('').astype(str).replace('', '未知／缺失')
                for code, n in codes.value_counts().sort_index().items():
                    composition.append(dict(**period, wave=wave, dimension=field, dimension_label=label,
                        category=code, count=int(n), denominator=len(sample), proportion=int(n)/len(sample)))
    whole = _summarize(selected, role, release_id, chosen)
    whole = whole[whole['wave'].isin(waves)]
    mode = rows['report_mode'].iloc[0] if len(rows) else rows.attrs.get('mode', '')
    anchor = rows['grouping_anchor'].iloc[0] if len(rows) else ''
    settings = dict(metadata or {})
    settings.update(basis=BASIS, trend_version=VERSION, report_version=REPORT_VERSION,
        dictionary_version=DICTIONARY_VERSION, release_id=release_id, role=role, mode=mode,
        grouping_anchor=anchor, timezone=TIMEZONE, start=start.isoformat(), end=end.isoformat(),
        interval=interval, include_batch_dates=bool(include_batch), waves=list(waves),
        variable_ids=chosen, date_policy='manual_verified_overrides_basis;observed_system;optional_accepted_batch',
        interpretation='各期各自有效回應；填卷者可以不同；不作逐人變化或因果推論')
    outputs = dict(statistics=pd.concat(statistics_parts, ignore_index=True), whole_period=whole,
        counts=pd.DataFrame(count_rows), exclusions=pd.DataFrame(audit), date_quality=_quality_counts(dated, mask),
        composition=pd.DataFrame(composition, columns=[*bins[0], 'wave', 'dimension', 'dimension_label',
                                                      'category', 'count', 'denominator', 'proportion']))
    for frame in outputs.values():
        frame['release_id'] = release_id
        frame['basis'] = BASIS
        frame['trend_version'] = VERSION
    outputs['metadata'] = settings
    return outputs


TABLE_NAMES = {'statistics': '每期指標統計', 'whole_period': '全期間指標統計',
               'counts': '每期回應數', 'exclusions': '日期排除數',
               'date_quality': '日期來源核對', 'composition': '每期樣本組成'}


def download_table(result, key):
    """Attach complete settings to aggregate CSV, without mutating the result."""
    frame = result[key].copy()
    frame['settings_json'] = json.dumps(result['metadata'], ensure_ascii=False, sort_keys=True)
    from .outputs import excel_safe
    return excel_safe(frame)


def workbook_bytes(result):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .descriptive_reports import _write_frame
    book = Workbook()
    book.remove(book.active)
    for key, title in TABLE_NAMES.items():
        _write_frame(book.create_sheet(title), result[key])
    settings = pd.DataFrame([dict(setting=k, value=json.dumps(v, ensure_ascii=False))
                             for k, v in result['metadata'].items()])
    _write_frame(book.create_sheet('設定及版本'), settings)
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 25
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='224B57')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[1].height = 44
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000'
            sheet.row_dimensions[row[0].row].height = 32
        sheet.print_title_rows = '1:1'
    book['設定及版本'].column_dimensions['B'].width = 100
    for row in book['設定及版本'].iter_rows(min_row=2):
        book['設定及版本'].row_dimensions[row[0].row].height = max(32, 16 * (1 + len(str(row[1].value)) // 85))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
