"""Research-facing comparisons over one immutable published release."""
import io
import json
from pathlib import Path

import pandas as pd

from .dashboard_views import SOURCE_NAMES
from .descriptive_spec import GROUPS, MODES, REPORT_FILES, SECTIONS, STATUS_LABELS, VERSION
from .descriptives import (cohort, display_table, filter_cohort, paired_changes,
                           statistics, wave_counts)


def filtered_workbook(stats, changes, metadata, exploration=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .descriptive_reports import _write_frame
    book = Workbook(); book.remove(book.active)
    frames = [] if stats is None else [('目前篩選', display_table(stats)), ('統計明細', stats)]
    if exploration:
        frames += [(title,exploration[key]) for key,title in [('distribution_summary','分布摘要'),
            ('histograms','共同分箱'),('pair_summary','完整配對摘要'),('pair_transitions','二元四格轉移'),
            ('categories','分類分布')] if not exploration[key].empty]
    frames.append(('篩選及版本', pd.DataFrame([{'項目': k, '值': json.dumps(v, ensure_ascii=False)} for k,v in metadata.items()])))
    for name, frame in frames:
        # Complete settings live once in the metadata sheet. Repeating the JSON
        # in every row would make every histogram row hundreds of pixels tall.
        ws = book.create_sheet(name); _write_frame(ws, frame.drop(columns='filters_json',errors='ignore'))
    if not changes.empty:
        _write_frame(book.create_sheet('配對差'), changes)
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        for col in range(1, sheet.max_column+1):
            sheet.column_dimensions[get_column_letter(col)].width = 35 if col == 1 else 24
        if sheet.title == '篩選及版本':
            sheet.column_dimensions['B'].width = 100
        for cell in sheet[1]:
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='224B57')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=48
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value,float):cell.number_format='0.0000'
            sheet.row_dimensions[row[0].row].height = 42
        if sheet.title == '目前篩選':
            sheet.column_dimensions['A'].width = 60
            sheet.column_dimensions['B'].width = 45
        for i,column in enumerate([c.value for c in sheet[1]],1):
            if column in {'variable','note','missing_detail','filters_json','t0_missing_detail','t1_missing_detail'}:
                sheet.column_dimensions[get_column_letter(i)].width=65
            for cells in sheet.iter_rows(min_row=2,min_col=i,max_col=i):
                if str(column).endswith('_n') or column in {'count','numerator','denominator','bin_index'}:
                    cells[0].number_format='#,##0'
                elif column in {'proportion','t0_proportion','t1_proportion'}:
                    cells[0].number_format='0.0%'
        # Long filter and method descriptions must be readable after Excel
        # export, rather than relying on its unavailable automatic row height.
        import math
        for row in sheet.iter_rows(min_row=2):
            lines=max(math.ceil(len(str(c.value or ''))/max(10,sheet.column_dimensions[c.column_letter].width-2)) for c in row)
            sheet.row_dimensions[row[0].row].height=max(42,15*(lines+1))
        sheet.print_title_rows='1:1'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=0;sheet.page_setup.fitToHeight=0
        sheet.page_setup.scale=100
    output = io.BytesIO(); book.save(output)
    return output.getvalue()


def render_results(rows, role, release_id, *, prefix, section='scales', groups=('all',), metadata=None):
    """Compatibility entry point for the shared research overview."""
    from .descriptive_workbench import render_overview
    return render_overview(rows,role,release_id,prefix=prefix,groups=groups,metadata=metadata)


def _reset(prefix, homepage=False):
    import streamlit as st
    for name in ('schools','sources','rules','types'):
        st.session_state[prefix+'_'+name]=[]
    st.session_state[prefix+'_group']='all'
    st.session_state[prefix+'_mainstream']=not homepage


def _capability(folder,summary):
    import streamlit as st
    from .descriptive_exploration_outputs import available
    from .common import file_hash
    @st.cache_data(show_spinner=False)
    def checked(release_id,path,manifest_hash,details):
        return available(Path(path),details)
    path=Path(folder)/'descriptive_exploration_manifest.json'
    return checked(summary['release_id'],str(folder),file_hash(path) if path.exists() else '',summary)


def _legacy(folder,summary,frame,parents):
    import streamlit as st
    st.info('此歷史版本未包含新版 4.1 人口特徵及並排報表，未提供時段趨勢。以下保留舊量表統計。')
    from .dashboard_views import analysis_cohort
    from .outputs import descriptive_summary
    role=st.selectbox('填卷類別',['student','parent'],key='legacy_ds_role')
    wave=st.selectbox('輪次',['T0','T1'],key='legacy_ds_wave')
    rows=analysis_cohort(frame,parents,role,wave,'全部保留樣本')
    st.metric('此輪全部學校 N',len(rows))
    st.dataframe(descriptive_summary(rows),hide_index=True,width='stretch')


def render_descriptive(folder,summary,frame,parents,rule_names):
    import streamlit as st
    from .descriptive_workbench import render_overview,render_workbench
    folder=Path(folder)
    if (not summary.get('descriptive') or not (folder/'descriptive_measurements.parquet').exists()
            or summary['descriptive'].get('version')!=VERSION):
        return _legacy(folder,summary,frame,parents)
    @st.cache_data(show_spinner=False)
    def load_measurements(release_id,path):
        return pd.read_parquet(Path(path)/'descriptive_measurements.parquet')
    measurements=load_measurements(summary['release_id'],str(folder))
    reading=st.radio('閱讀方式',['研究總覽','跨期比較','時段趨勢'],horizontal=True,key='ds_reading')
    paired=reading=='跨期比較';prefix='dx' if paired else 'ds'
    controls=st.columns(3)
    modes=['acbd','efmatched'] if paired else list(MODES)
    mode=controls[0].selectbox('跨期分析樣本' if paired else '分析樣本',modes,format_func=MODES.get,key=prefix+'_mode')
    role=controls[1].selectbox('填卷對象',['student','parent'],format_func=lambda x:'學生' if x=='student' else '家長',key=prefix+'_role')
    main=controls[2].selectbox('學校範圍',[True,False],format_func=lambda x:'只限主流學校（研究分析）' if x else '全部學校（首頁口徑）',key=prefix+'_mainstream')
    context=(summary['release_id'],mode,role)
    context_key=prefix+'_context'
    rows=cohort(frame,parents,measurements,mode,role)
    options=sorted({r for text in rows['anchor_identity_rules'].dropna() for r in text.split(';') if r})
    if st.session_state.get(context_key)!=context:
        valid={'schools':set(rows['anchor_school'].dropna()),'sources':set(rows['anchor_source_type'].dropna()),'rules':set(options)}
        removed=False
        for name,allowed in valid.items():
            key=prefix+'_'+name
            previous=st.session_state.get(key,[])
            retained=[v for v in previous if v in allowed]
            if previous!=retained:
                st.session_state[key]=retained;removed=True
        if removed:
            st.caption('樣本、角色或版本已改變，已清除不再適用的篩選；其餘保留。')
        st.session_state[context_key]=context
    with st.expander('孩子性別／學校階段分組'):
        group=st.selectbox('目前孩子分組',list(GROUPS),format_func=GROUPS.get,key=prefix+'_group')
        st.caption('跨期以基線固定分組。九組互有重疊，不可相加。')
    with st.expander('進階篩選（留空代表不限）'):
        c=st.columns(3)
        schools=c[0].multiselect('學校',sorted(rows['anchor_school'].dropna().unique()),key=prefix+'_schools')
        sources=c[1].multiselect('資料來源',sorted(rows['anchor_source_type'].dropna().unique()),format_func=lambda v:SOURCE_NAMES.get(v,v),key=prefix+'_sources')
        rules=c[2].multiselect('孩子身份配對規則',options,format_func=lambda v:rule_names.get(v,v),key=prefix+'_rules')
        st.caption('身份規則選項只篩選目前正式樣本；多選符合任一項。清空代表不限，不會重新做身份配對。')
    filters=dict(mainstream=main,schools=schools,sources=sources,rules=rules,group=group)
    selected,counts=filter_cohort(rows,**filters)
    n=wave_counts(selected)
    st.metric('日期篩選前 T0／T1 N' if reading=='時段趨勢' else '目前 T0／T1 N',f"{n['T0']:,} / {n['T1']:,}")
    anchor=rows['grouping_anchor'].iloc[0] if len(rows) else '此模式既定錨點'
    st.caption(f"{MODES[mode]}；{'學生' if role=='student' else '家長'}；{'主流學校' if main else '全部學校'}；{GROUPS[group]}。分組依據：{anchor}。")
    active=[]
    for label,value in [('學校',schools),('來源',sources),('身份規則',rules)]:
        if value:active.append(label+'：'+('、'.join(value if label=='學校' else [SOURCE_NAMES.get(v,v) if label=='來源' else rule_names.get(v,v) for v in value])))
    st.caption('目前篩選：'+('；'.join(active) if active else '無進階篩選')+'。')
    if paired:
        st.caption('這是明確選定的正式跨期樣本。與總覽的各輪 A–D／E–F 人數可能不同；指標缺失另扣。BD 使用其實際選定 T1。')
    elif reading=='時段趨勢':
        st.caption('以下按填卷日期篩選正式選定回應，輪次不重新選取；每期各自有效回答，完整設定及下載標示 calendar_response。')
    else:
        st.caption('兩輪各自有效回答。A–D／E–F 可以不同人；逐人變化請明確選取「跨期比較」的樣本及新 N。完整設定及下載標示 wave_valid。')
    with st.expander('核對口徑、逐步排除數及資料版本'):
        st.write({'發布版本':summary['release_id'],'報表版本':VERSION,'分組依據':anchor,'完整篩選':filters})
        detail=counts.copy()
        for w in ['T0','T1']:detail[w+'本步排除']=(-detail[w].diff()).fillna(0).astype(int)
        st.dataframe(detail.rename(columns={'stage':'步驟'}),hide_index=True,width='stretch')
        st.caption(f"未知學段：T0 {selected.loc[selected['wave'].eq('T0'),'anchor_school_stage'].isna().sum():,}；T1 {selected.loc[selected['wave'].eq('T1'),'anchor_school_stage'].isna().sum():,}，保留在全部組。")
    buttons=st.columns(2)
    buttons[0].button('清除篩選',on_click=_reset,args=(prefix,),key=prefix+'_reset')
    buttons[1].button('首頁口徑：全部學校、取消篩選',on_click=_reset,args=(prefix,True),key=prefix+'_home')
    metadata={'mode':mode,'grouping_anchor':anchor,**filters}
    if paired:
        if not _capability(folder,summary):
            st.info('此發布未提供跨期探索能力；研究總覽及原有報表仍可閱讀。發布新版後可使用工作台。')
        else:
            render_workbench(selected,role,summary['release_id'],prefix=prefix,metadata=metadata)
    elif reading=='時段趨勢':
        from .descriptive_trend_dashboard import render_trends
        render_trends(selected,role,summary['release_id'],prefix='ds_trend',metadata=metadata)
    else:
        render_overview(selected,role,summary['release_id'],prefix=prefix,groups=(group,),metadata=metadata)
    if role=='parent' and mode=='efmatched':
        with st.expander('兩輪親子樣本的填卷成人身份'):
            caregivers=pd.read_parquet(folder/'descriptive_caregivers.parquet')
            caregivers=caregivers[caregivers['child_id'].isin(selected['child_id'])]
            status=caregivers['caregiver_status'].map(lambda s:'已確認同一成人' if str(s).startswith('same_adult_by_') else '已確認不同成人' if s=='reviewer_rejected_same_adult' else '成人身份尚未確認')
            st.dataframe(status.value_counts().rename_axis('填卷人狀態').reset_index(name='家庭組數'),hide_index=True,width='stretch')
    if role=='parent' and 'student_parent_calendar_gap_days' in selected:
        with st.expander('核對家長與學生日期差'):
            gaps=selected.reindex(columns=['submission_id','wave','student_parent_calendar_gap_days',
                'student_parent_elapsed_gap_days','parent_response_interval_days']).rename(columns={
                'submission_id':'家長提交ID','wave':'選定輪次','student_parent_calendar_gap_days':'家長減學生香港日期差（日）',
                'student_parent_elapsed_gap_days':'家長減學生實際時間差（日）','parent_response_interval_days':'距家長T0實際日數'})
            st.caption('負數表示家長較早，正數表示較遲。學生錨定家長樣本沒有自身90日門檻；獨立家長分支另計。')
            st.dataframe(gaps,hide_index=True,width='stretch')
            st.download_button('下載目前家長日期差',gaps.to_csv(index=False).encode('utf-8-sig'),
                'study2-parent-date-gaps.csv','text/csv',key=prefix+'_date_gaps')
    with st.expander('下載本版四份完整 4.1 報表（主流學校、九個分組）'):
        st.caption('本版固定報表使用每輪各自有效的定義，不受目前篩選影響。完整配對新統計另列。')
        for report_mode,name in REPORT_FILES.items():
            st.download_button(MODES[report_mode],(folder/'4.1_descriptive'/name).read_bytes(),name,
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='ds_report_'+report_mode)
